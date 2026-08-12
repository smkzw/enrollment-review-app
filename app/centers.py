"""Project center roster helpers."""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Iterable

from app.config import PROJECTS_DIR


def normalize_center_code(value: object) -> str:
    text = str(value or "").strip()
    text = re.sub(r"\s+", "", text)
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        text = str(int(float(text)))
    if text.isdigit() and len(text) < 2:
        text = text.zfill(2)
    return text


def normalize_center_name(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def center_label(center: dict) -> str:
    code = normalize_center_code(center.get("center_code"))
    name = normalize_center_name(center.get("center_name"))
    if code and name:
        return f"{code}｜{name}"
    return code or name


def centers_path(project_path: Path | None = None) -> Path:
    system_dir = PROJECTS_DIR / "_system"
    system_dir.mkdir(parents=True, exist_ok=True)
    return system_dir / "centers.json"


def normalize_center_record(raw: dict | None) -> dict | None:
    raw = raw or {}
    code = normalize_center_code(raw.get("center_code") or raw.get("code") or raw.get("中心代号") or raw.get("中心编号"))
    name = normalize_center_name(raw.get("center_name") or raw.get("name") or raw.get("中心名称") or raw.get("机构名称"))
    if not code and not name:
        return None
    return {
        "center_code": code,
        "center_name": name,
        "label": f"{code}｜{name}" if code and name else code or name,
        "owner_username": str(raw.get("owner_username") or raw.get("created_by") or "").strip(),
        "updated_at": str(raw.get("updated_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    }


def _dedupe_centers(items: Iterable[dict]) -> list[dict]:
    by_key: dict[str, dict] = {}
    ordered: list[str] = []
    for item in items:
        rec = normalize_center_record(item)
        if not rec:
            continue
        key = rec["center_code"] or rec["center_name"]
        if key not in by_key:
            ordered.append(key)
            by_key[key] = rec
            continue
        if rec["center_name"]:
            by_key[key]["center_name"] = rec["center_name"]
            by_key[key]["label"] = center_label(by_key[key])
        if rec.get("owner_username") and not by_key[key].get("owner_username"):
            by_key[key]["owner_username"] = rec["owner_username"]
        by_key[key]["updated_at"] = rec["updated_at"]
    return sorted((by_key[k] for k in ordered), key=lambda c: (c.get("center_code") or "~~~~", c.get("center_name") or ""))


def load_centers(project_path: Path | None = None, include_subjects: bool = True) -> list[dict]:
    items: list[dict] = []
    path = centers_path(project_path)
    if path.exists():
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(data, list):
                items.extend([x for x in data if isinstance(x, dict)])
        except Exception:
            pass
    if include_subjects:
        for info_path in sorted(PROJECTS_DIR.glob("*/subjects/*/info.json")):
            try:
                info = json.loads(info_path.read_text(encoding="utf-8"))
            except Exception:
                continue
            items.append({
                "center_code": info.get("center_code", ""),
                "center_name": info.get("center_name", ""),
                "owner_username": info.get("owner_username", ""),
            })
    return _dedupe_centers(items)


def save_centers(project_path: Path | None, centers: Iterable[dict]) -> list[dict]:
    normalized = _dedupe_centers(centers)
    centers_path(project_path).write_text(
        json.dumps(normalized, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return normalized


def normalize_centers(centers: Iterable[dict]) -> list[dict]:
    return _dedupe_centers(centers)


def upsert_center(project_path: Path | None, center_code: str = "", center_name: str = "", owner_username: str = "") -> list[dict]:
    rec = normalize_center_record({"center_code": center_code, "center_name": center_name, "owner_username": owner_username})
    centers = load_centers(project_path, include_subjects=False)
    if rec:
        centers.append(rec)
    return save_centers(project_path, centers)


def parse_centers_from_text(text: str) -> list[dict]:
    rows: list[dict] = []
    for raw_line in (text or "").splitlines():
        line = raw_line.strip().strip("|")
        if not line or set(line.replace("|", "").strip()) <= {"-", ":"}:
            continue
        if any(key in line for key in ("中心代号", "中心编号", "中心名称", "机构名称")):
            continue
        parts = [p.strip() for p in re.split(r"\s*[|｜,\t]\s*", line) if p.strip()]
        if len(parts) >= 2:
            rows.append({"center_code": parts[0], "center_name": parts[1]})
            continue
        m = re.match(r"^([A-Za-z0-9_-]{1,12})\s+(.+)$", line)
        if m:
            rows.append({"center_code": m.group(1), "center_name": m.group(2)})
    return _dedupe_centers(rows)


def parse_centers_from_excel_bytes(content: bytes, filename: str = "") -> list[dict]:
    suffix = Path(filename or "").suffix.lower()
    if suffix in {".csv", ".txt"}:
        text = content.decode("utf-8-sig", errors="ignore")
        return _parse_csv_or_text(text)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows: list[dict] = []
    for ws in wb.worksheets:
        values = list(ws.iter_rows(values_only=True))
        if not values:
            continue
        header_idx = _find_header_row(values)
        headers = [str(v or "").strip() for v in values[header_idx]]
        code_idx = _find_col(headers, ("中心代号", "中心编号", "中心代码", "机构编号", "site code", "center code"))
        name_idx = _find_col(headers, ("中心名称", "机构名称", "研究中心", "site name", "center name"))
        if code_idx is None and name_idx is None:
            continue
        for row in values[header_idx + 1:]:
            code = row[code_idx] if code_idx is not None and code_idx < len(row) else ""
            name = row[name_idx] if name_idx is not None and name_idx < len(row) else ""
            rows.append({"center_code": code, "center_name": name})
    return _dedupe_centers(rows)


def _parse_csv_or_text(text: str) -> list[dict]:
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample)
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        rows = []
        for row in reader:
            rows.append({
                "center_code": row.get("中心代号") or row.get("中心编号") or row.get("center_code") or row.get("site code") or "",
                "center_name": row.get("中心名称") or row.get("机构名称") or row.get("center_name") or row.get("site name") or "",
            })
        parsed = _dedupe_centers(rows)
        if parsed:
            return parsed
    except Exception:
        pass
    return parse_centers_from_text(text)


def _find_header_row(values: list[tuple]) -> int:
    for idx, row in enumerate(values[:20]):
        text = " ".join(str(v or "") for v in row)
        if any(key in text for key in ("中心代号", "中心编号", "中心名称", "机构名称", "site code", "site name")):
            return idx
    return 0


def _find_col(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    lowered = [h.lower().replace(" ", "") for h in headers]
    for i, h in enumerate(lowered):
        for cand in candidates:
            if cand.lower().replace(" ", "") in h:
                return i
    return None


def centers_to_markdown(centers: Iterable[dict]) -> str:
    lines = ["| 中心代号 | 中心名称 |", "|---|---|"]
    for center in _dedupe_centers(centers):
        lines.append(f"| {center.get('center_code', '')} | {center.get('center_name', '')} |")
    return "\n".join(lines)
