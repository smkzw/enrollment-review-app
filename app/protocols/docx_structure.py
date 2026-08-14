"""DOCX 结构提取：段落、表格（含嵌套）、编号、页眉页脚、section、修订标记。

产出两部分：

1. 规范化块集 :class:`StructureBlock`（完整内容工件），按 ``source_ref`` 稳定
   定位，``source_ref`` 编码嵌套表格路径（如 ``body.t3.r1.c2.t0.r0.c0.p1``），
   满足「来源定位必须表示嵌套表格路径」的要求；
2. 不可变的 :class:`ProtocolExtractionSnapshot`，其 ``content_sha256`` 为块集
   规范化序列化的 SHA-256，``content_storage_ref`` 指向块集落盘位置——计数本身
   不足以证明摘录属于该文件哈希，块集哈希才能逐字核验。

页眉页脚按「default / first / even」三类分别捕获，按 OOXML 部件身份（partname）
去重继承/共享部件，并保留部件适用的全部 section。编号定义从 ``word/numbering.xml``
解析抽象编号身份、级别格式、级别文本模板、起始值与具体覆盖，写入不可变块集，
使后续目录构建无需重新打开源文件即可重建官方编号。

本模块用 python-docx 读取 section/页眉页脚/编号部件，用受控 OOXML（``w:body``
子元素顺序遍历）补齐文档顺序、嵌套表格、编号定义与修订标记。源文件只读，派生
块集写入调用方给定的输出目录。
"""
from __future__ import annotations

import hashlib
import json
import os
from dataclasses import dataclass
from datetime import datetime
from enum import Enum
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pydantic import BaseModel, ConfigDict, Field

from app.domain.contracts.enums import DocumentPart, ExtractionStatus
from app.domain.contracts.protocol_ingestion import (
    ExtractionAnomaly,
    ExtractionCoverage,
    ProtocolExtractionSnapshot,
    ProtocolSourceArtifact,
)
from .ingestion import (
    ProtocolFileKind,
    SourceIngestionError,
    compute_sha256,
    detect_format,
    utc_now,
)

PARSER_NAME = "docx-ooxml"
PARSER_VERSION = "1.2.0"

# OOXML 限定名
_W_P = qn("w:p")
_W_TBL = qn("w:tbl")
_W_TR = qn("w:tr")
_W_TC = qn("w:tc")
_W_T = qn("w:t")
_W_TAB = qn("w:tab")
_W_BR = qn("w:br")
_W_CR = qn("w:cr")
_W_NUMPR = qn("w:numPr")
_W_ILVL = qn("w:ilvl")
_W_NUMID = qn("w:numId")
_W_SECTPR = qn("w:sectPr")
_W_PPR = qn("w:pPr")
_W_PSTYLE = qn("w:pStyle")
_W_VAL = qn("w:val")
_W_INS = qn("w:ins")
_W_DEL = qn("w:del")
_W_TCPR = qn("w:tcPr")
_W_GRIDSPAN = qn("w:gridSpan")
_W_STYLE = qn("w:style")
_W_STYLE_ID = qn("w:styleId")
_W_BASED_ON = qn("w:basedOn")
_W_SDT = qn("w:sdt")
_W_SDT_CONTENT = qn("w:sdtContent")
_W_CUSTOM_XML = qn("w:customXml")
_W_ALT_CHUNK = qn("w:altChunk")

# numbering.xml
_W_ABSTRACT_NUM = qn("w:abstractNum")
_W_ABSTRACT_NUM_ID = qn("w:abstractNumId")
_W_NUM = qn("w:num")
_W_LVL = qn("w:lvl")
_W_LVL_OVERRIDE = qn("w:lvlOverride")
_W_START_OVERRIDE = qn("w:startOverride")
_W_NUMFMT = qn("w:numFmt")
_W_LVLTEXT = qn("w:lvlText")
_W_START = qn("w:start")

_BLOCKS_STORAGE_PREFIX = "blobs/protocol_blocks"


class BlockKind(str, Enum):
    PARAGRAPH = "paragraph"
    TABLE = "table"


class HeaderFooterKind(str, Enum):
    """页眉/页脚部件种类（OOXML ``w:headerReference``/``w:footerReference`` 的 type）。"""

    DEFAULT = "default"
    FIRST = "first"
    EVEN = "even"


class NumberingRef(BaseModel):
    """段落编号引用：抽象编号身份、级别格式、级别文本模板、起始值与覆盖标志。

    这些数据写入不可变块集，使后续官方编号目录构建无需重新打开源文件。
    """

    model_config = ConfigDict(extra="forbid", frozen=True)

    num_id: int
    level: int
    abstract_num_id: int | None = None
    num_fmt: str | None = None
    lvl_text: str | None = None
    start: int | None = None
    has_level_override: bool = False


class StructureBlock(BaseModel):
    """结构块：结构通道的完整内容工件单元，``source_ref`` 唯一稳定。"""

    model_config = ConfigDict(extra="forbid", frozen=True)

    source_ref: str
    document_part: DocumentPart
    section_index: int | None = None
    block_order: int = Field(ge=0)
    kind: BlockKind
    text: str
    style: str | None = None
    numbering: NumberingRef | None = None
    # 段落/嵌套表所在单元格的完整祖先 (row, col, ...) 链；body 顶层内容为 None。
    table_path: tuple[int, ...] | None = None
    table_rows: int | None = None
    table_cols: int | None = None
    tracked_change: bool = False
    # 页眉/页脚部件种类；正文块为 None。
    part_kind: HeaderFooterKind | None = None
    # 页眉/页脚部件适用的全部 section 序号；正文块为 None。
    section_indexes: tuple[int, ...] | None = None


@dataclass(frozen=True)
class StructureExtraction:
    """一次结构提取的完整结果：块集 + 不可变快照。"""

    blocks: tuple[StructureBlock, ...]
    snapshot: ProtocolExtractionSnapshot


class StructureExtractionError(SourceIngestionError):
    """DOCX 结构提取失败（格式错误或零内容）。"""


def _int_attr(element, attr_qname: str) -> int | None:
    if element is None:
        return None
    value = element.get(attr_qname)
    if value is None:
        return None
    try:
        return int(value)
    except ValueError:
        return None


def _para_text(p_el) -> str:
    """按阅读顺序提取段落文本：``w:t`` 文本 + 制表符 + 换行。"""
    parts: list[str] = []
    for node in p_el.iter():
        if node.tag == _W_T:
            parts.append(node.text or "")
        elif node.tag == _W_TAB:
            parts.append("\t")
        elif node.tag in (_W_BR, _W_CR):
            parts.append("\n")
    return "".join(parts)


def _para_style(p_el) -> str | None:
    ppr = p_el.find(_W_PPR)
    if ppr is None:
        return None
    pstyle = ppr.find(_W_PSTYLE)
    return pstyle.get(_W_VAL) if pstyle is not None else None


def _para_tracked_change(p_el) -> bool:
    for node in p_el.iter():
        if node.tag in (_W_INS, _W_DEL):
            return True
    return False


def _grid_span(tc) -> int:
    tcpr = tc.find(_W_TCPR)
    if tcpr is not None:
        gs = tcpr.find(_W_GRIDSPAN)
        if gs is not None:
            span = _int_attr(gs, _W_VAL)
            if span is not None and span > 0:
                return span
    return 1


def _cells_with_col(row_el):
    """按 gridSpan 展开列序号，避免合并单元格文本重复。

    每个单元格产出 (tc, 起始列, 跨列数)。
    """
    col = 0
    for tc in row_el.findall(_W_TC):
        span = _grid_span(tc)
        yield tc, col, span
        col += span


def _iter_block_children(container_el):
    """按顺序产出段落/表格，并展开常见透明 OOXML 内容容器。"""
    for child in container_el:
        if child.tag in (_W_P, _W_TBL):
            yield child
        elif child.tag == _W_SDT:
            content = child.find(_W_SDT_CONTENT)
            if content is not None:
                yield from _iter_block_children(content)
        elif child.tag == _W_CUSTOM_XML:
            yield from _iter_block_children(child)


def _unhandled_content_tags(container_el) -> set[str]:
    """找出可能承载正文却未被结构遍历支持的直接容器。"""
    tags: set[str] = set()
    for child in container_el:
        if child.tag in (_W_P, _W_TBL):
            continue
        if child.tag == _W_SDT:
            content = child.find(_W_SDT_CONTENT)
            if content is not None:
                tags.update(_unhandled_content_tags(content))
            continue
        if child.tag == _W_CUSTOM_XML:
            tags.update(_unhandled_content_tags(child))
            continue
        if child.tag == _W_ALT_CHUNK or any(
            node.tag in (_W_P, _W_TBL, _W_T) for node in child.iter()
        ):
            tags.add(child.tag.rsplit("}", 1)[-1])
    return tags


def _table_size(tbl_el) -> tuple[int, int]:
    rows = tbl_el.findall(_W_TR)
    row_count = len(rows)
    col_count = 0
    for row in rows:
        for _tc, col, span in _cells_with_col(row):
            col_count = max(col_count, col + span)
    return row_count, col_count


def _level_definition(lvl_el) -> dict:
    fmt = lvl_el.find(_W_NUMFMT)
    text = lvl_el.find(_W_LVLTEXT)
    start = lvl_el.find(_W_START)
    return {
        "num_fmt": fmt.get(_W_VAL) if fmt is not None else None,
        "lvl_text": text.get(_W_VAL) if text is not None else None,
        "start": _int_attr(start, _W_VAL) if start is not None else None,
    }


def parse_numbering(root):
    """解析 ``w:numbering`` 根元素，返回 (abstracts, nums)。

    ``abstracts``: abstractNumId -> {level -> LevelDef}；
    ``nums``: numId -> (abstractNumId, {level -> lvlOverride 元素})。
    """
    if root is None:
        return {}, {}
    abstracts: dict[int, dict[int, dict]] = {}
    for ab_el in root.findall(_W_ABSTRACT_NUM):
        aid = _int_attr(ab_el, _W_ABSTRACT_NUM_ID)
        if aid is None:
            continue
        levels: dict[int, dict] = {}
        for lvl_el in ab_el.findall(_W_LVL):
            ilvl = _int_attr(lvl_el, _W_ILVL)
            if ilvl is None:
                continue
            levels[ilvl] = _level_definition(lvl_el)
        abstracts[aid] = levels
    nums: dict[int, tuple[int | None, dict[int, object]]] = {}
    for num_el in root.findall(_W_NUM):
        nid = _int_attr(num_el, _W_NUMID)
        if nid is None:
            continue
        abstract_el = num_el.find(_W_ABSTRACT_NUM_ID)
        aid = _int_attr(abstract_el, _W_VAL) if abstract_el is not None else None
        overrides: dict[int, object] = {}
        for lo_el in num_el.findall(_W_LVL_OVERRIDE):
            ilvl = _int_attr(lo_el, _W_ILVL)
            if ilvl is not None:
                overrides[ilvl] = lo_el
        nums[nid] = (aid, overrides)
    return abstracts, nums


def _numbering_definitions(document):
    """从 ``word/numbering.xml`` 读取编号定义。"""
    part = document.part.numbering_part
    root = part.element if part is not None else None
    return parse_numbering(root)


def _styles_element(document):
    """返回 ``w:styles`` 根元素；缺失时返回 None。"""
    try:
        return document.styles.element
    except Exception:
        return None


def _style_numbering_definitions(document) -> dict[str, tuple[int, int]]:
    """从 ``styles.xml`` 解析段落样式的编号引用，含 ``w:basedOn`` 样式链继承。

    返回 ``style_id -> (num_id, level)``；仅包含能解析出有效 ``numId``（非 0）
    的样式。Word 常把自动编号定义在段落样式上，段落本身不带 ``w:numPr``，
    必须沿样式继承链才能还原官方编号（章节标题编号等）。
    """
    root = _styles_element(document)
    if root is None:
        return {}

    styles: dict[str, dict] = {}
    for style_el in root.findall(_W_STYLE):
        style_id = style_el.get(_W_STYLE_ID)
        if style_id is None:
            continue
        ppr = style_el.find(_W_PPR)
        based_on = None
        num_id = None
        level = 0
        based_on_el = style_el.find(_W_BASED_ON)
        if based_on_el is not None:
            based_on = based_on_el.get(_W_VAL)
        if ppr is not None:
            numpr = ppr.find(_W_NUMPR)
            if numpr is not None:
                numid_el = numpr.find(_W_NUMID)
                ilvl_el = numpr.find(_W_ILVL)
                num_id = _int_attr(numid_el, _W_VAL) if numid_el is not None else None
                level = _int_attr(ilvl_el, _W_VAL) if ilvl_el is not None else 0
        styles[style_id] = {"based_on": based_on, "num_id": num_id, "level": level}

    resolved: dict[str, tuple[int, int] | None] = {}

    def resolve(style_id: str, seen: frozenset[str]) -> tuple[int, int] | None:
        if style_id in resolved:
            return resolved[style_id]
        if style_id in seen:
            return None
        info = styles.get(style_id)
        if info is None:
            resolved[style_id] = None
            return None
        if info["num_id"] is not None and info["num_id"] != 0:
            result: tuple[int, int] = (info["num_id"], info["level"])
            resolved[style_id] = result
            return result
        based_on = info["based_on"]
        if based_on is not None:
            inherited = resolve(based_on, seen | {style_id})
            resolved[style_id] = inherited
            return inherited
        resolved[style_id] = None
        return None

    for style_id in styles:
        resolve(style_id, frozenset())
    return {
        style_id: result
        for style_id, result in resolved.items()
        if result is not None
    }


def _effective_numbering(num_id, level, abstracts, nums) -> NumberingRef:
    aid, overrides = nums.get(num_id, (None, {}))
    base = abstracts.get(aid, {}).get(level, {}) if aid is not None else {}
    num_fmt = base.get("num_fmt")
    lvl_text = base.get("lvl_text")
    start = base.get("start")
    override_el = overrides.get(level)
    has_override = override_el is not None
    if override_el is not None:
        start_override = override_el.find(_W_START_OVERRIDE)
        if start_override is not None:
            start = _int_attr(start_override, _W_VAL)
        # lvlOverride 可携带完整 <w:lvl> 覆盖级别定义
        lvl_el = override_el.find(_W_LVL)
        if lvl_el is not None:
            over_def = _level_definition(lvl_el)
            if over_def.get("num_fmt") is not None:
                num_fmt = over_def["num_fmt"]
            if over_def.get("lvl_text") is not None:
                lvl_text = over_def["lvl_text"]
            if over_def.get("start") is not None:
                start = over_def["start"]
    return NumberingRef(
        num_id=num_id,
        level=level,
        abstract_num_id=aid,
        num_fmt=num_fmt,
        lvl_text=lvl_text,
        start=start,
        has_level_override=has_override,
    )


class _Extractor:
    """一次提取的可变状态：全局块序、计数、编号上下文与异常。"""

    def __init__(self, abstracts, nums, style_numbering: dict[str, tuple[int, int]] | None = None) -> None:
        self.blocks: list[StructureBlock] = []
        self.paragraph_count = 0
        self.table_count = 0
        self.nested_table_count = 0
        self.header_detected = False
        self.footer_detected = False
        self.abstracts = abstracts
        self.nums = nums
        self.style_numbering = style_numbering or {}
        self.anomalies: list[ExtractionAnomaly] = []

    def _next_order(self) -> int:
        return len(self.blocks)

    def _record_unresolved_numbering(self, source_ref: str, message: str) -> None:
        self.anomalies.append(
            ExtractionAnomaly(
                anomaly_id=f"unresolved_numbering:{source_ref}",
                kind="unresolved_numbering",
                message=message,
                scope_ref=source_ref,
            )
        )

    def _para_numbering(self, p_el, source_ref: str) -> NumberingRef | None:
        """解析段落编号：直接 ``w:numPr`` 优先，缺省时沿段落样式继承链查找。

        - ``numId=0`` 表示显式取消编号（去除继承自样式的编号），非异常；
        - 能确定 ``numId`` 却无法映射到有效抽象编号定义/级别格式时，记录
          ``unresolved_numbering`` 异常，绝不静默丢弃编号上下文。
        """
        ppr = p_el.find(_W_PPR)
        if ppr is None:
            return None
        numpr = ppr.find(_W_NUMPR)
        num_id: int | None = None
        level = 0
        if numpr is not None:
            numid_el = numpr.find(_W_NUMID)
            if numid_el is None:
                self._record_unresolved_numbering(source_ref, "段落含 numPr 但缺少 numId")
                return None
            num_id = _int_attr(numid_el, _W_VAL)
            if num_id is None:
                self._record_unresolved_numbering(source_ref, "numId 无法解析为整数")
                return None
            ilvl_el = numpr.find(_W_ILVL)
            level = _int_attr(ilvl_el, _W_VAL) if ilvl_el is not None else 0
            if num_id == 0:
                return None  # 显式取消编号
        else:
            style = _para_style(p_el)
            if style:
                style_ref = self.style_numbering.get(style)
                if style_ref:
                    num_id, level = style_ref
        if num_id is None:
            return None
        resolved = _effective_numbering(num_id, level, self.abstracts, self.nums)
        if resolved.abstract_num_id is None or resolved.lvl_text is None:
            self._record_unresolved_numbering(
                source_ref,
                f"numId={num_id} 无法解析为有效编号定义"
                f"（abstract_num_id={resolved.abstract_num_id}，lvl_text={resolved.lvl_text}）",
            )
        return resolved

    def add_paragraph(
        self,
        *,
        source_ref: str,
        document_part: DocumentPart,
        section_index: int | None,
        p_el,
        table_path: tuple[int, ...] | None = None,
        part_kind: HeaderFooterKind | None = None,
        section_indexes: tuple[int, ...] | None = None,
    ) -> None:
        text = _para_text(p_el)
        self.blocks.append(
            StructureBlock(
                source_ref=source_ref,
                document_part=document_part,
                section_index=section_index,
                block_order=self._next_order(),
                kind=BlockKind.PARAGRAPH,
                text=text,
                style=_para_style(p_el),
                numbering=self._para_numbering(p_el, source_ref),
                table_path=table_path,
                tracked_change=_para_tracked_change(p_el),
                part_kind=part_kind,
                section_indexes=section_indexes,
            )
        )
        self.paragraph_count += 1
        if document_part == DocumentPart.HEADER and text.strip():
            self.header_detected = True
        if document_part == DocumentPart.FOOTER and text.strip():
            self.footer_detected = True

    def add_table(
        self,
        *,
        source_ref: str,
        document_part: DocumentPart,
        section_index: int | None,
        tbl_el,
        table_path: tuple[int, ...] | None = None,
        part_kind: HeaderFooterKind | None = None,
        section_indexes: tuple[int, ...] | None = None,
    ) -> None:
        rows, cols = _table_size(tbl_el)
        self.blocks.append(
            StructureBlock(
                source_ref=source_ref,
                document_part=document_part,
                section_index=section_index,
                block_order=self._next_order(),
                kind=BlockKind.TABLE,
                text="",
                table_path=table_path,
                table_rows=rows,
                table_cols=cols,
                part_kind=part_kind,
                section_indexes=section_indexes,
            )
        )
        self.table_count += 1

    def walk_paragraphs_and_tables(
        self,
        container_el,
        *,
        document_part,
        section_index,
        prefix,
        part_kind: HeaderFooterKind | None = None,
        section_indexes: tuple[int, ...] | None = None,
    ) -> None:
        """按文档顺序遍历容器（body/header/footer），产出段落与表格块。"""
        para_seq = 0
        table_seq = 0
        for child in _iter_block_children(container_el):
            if child.tag == _W_P:
                self.add_paragraph(
                    source_ref=f"{prefix}.p{para_seq}",
                    document_part=document_part,
                    section_index=section_index,
                    p_el=child,
                    part_kind=part_kind,
                    section_indexes=section_indexes,
                )
                para_seq += 1
            elif child.tag == _W_TBL:
                table_ref = f"{prefix}.t{table_seq}"
                self.add_table(
                    source_ref=table_ref,
                    document_part=document_part,
                    section_index=section_index,
                    tbl_el=child,
                    part_kind=part_kind,
                    section_indexes=section_indexes,
                )
                table_seq += 1
                self._walk_table(
                    child, table_ref, document_part, section_index, part_kind,
                    section_indexes, parent_path=None,
                )

    def _walk_table(
        self,
        tbl_el,
        table_ref,
        document_part,
        section_index,
        part_kind: HeaderFooterKind | None,
        section_indexes: tuple[int, ...] | None,
        parent_path: tuple[int, ...] | None,
    ) -> None:
        """遍历表格行/单元格，含嵌套表格；段落与嵌套表均以嵌套路径为 source_ref。

        ``parent_path`` 为本表所在单元格的完整祖先链（顶层表为 None）；单元格内
        段落/嵌套表的 ``table_path`` 追加当前 (row, col) 形成完整链。
        """
        for row_idx, row in enumerate(tbl_el.findall(_W_TR)):
            for tc, col_idx, _span in _cells_with_col(row):
                cell_path = (
                    parent_path + (row_idx, col_idx)
                    if parent_path is not None
                    else (row_idx, col_idx)
                )
                cell_prefix = f"{table_ref}.r{row_idx}.c{col_idx}"
                para_seq = 0
                nested_table_seq = 0
                for child in _iter_block_children(tc):
                    if child.tag == _W_P:
                        self.add_paragraph(
                            source_ref=f"{cell_prefix}.p{para_seq}",
                            document_part=document_part,
                            section_index=section_index,
                            p_el=child,
                            table_path=cell_path,
                            part_kind=part_kind,
                            section_indexes=section_indexes,
                        )
                        para_seq += 1
                    elif child.tag == _W_TBL:
                        nested_ref = f"{cell_prefix}.t{nested_table_seq}"
                        self.add_table(
                            source_ref=nested_ref,
                            document_part=document_part,
                            section_index=section_index,
                            tbl_el=child,
                            table_path=cell_path,
                            part_kind=part_kind,
                            section_indexes=section_indexes,
                        )
                        nested_table_seq += 1
                        self.nested_table_count += 1
                        self._walk_table(
                            child, nested_ref, document_part, section_index,
                            part_kind, section_indexes, parent_path=cell_path,
                        )


def _walk_body(document, extractor: _Extractor) -> None:
    """按文档顺序遍历正文，保留 section 归属与段落/表格交错顺序。"""
    body = document.element.body
    para_seq = 0
    table_seq = 0
    current_section = 0
    for child in _iter_block_children(body):
        if child.tag == _W_P:
            extractor.add_paragraph(
                source_ref=f"body.p{para_seq}",
                document_part=DocumentPart.BODY,
                section_index=current_section,
                p_el=child,
            )
            para_seq += 1
            ppr = child.find(_W_PPR)
            if ppr is not None and ppr.find(_W_SECTPR) is not None:
                current_section += 1
        elif child.tag == _W_TBL:
            table_ref = f"body.t{table_seq}"
            extractor.add_table(
                source_ref=table_ref,
                document_part=DocumentPart.BODY,
                section_index=current_section,
                tbl_el=child,
            )
            table_seq += 1
            extractor._walk_table(
                child, table_ref, DocumentPart.BODY, current_section, None, None,
                parent_path=None,
            )


_HEADER_ATTRS = {
    HeaderFooterKind.DEFAULT: "header",
    HeaderFooterKind.FIRST: "first_page_header",
    HeaderFooterKind.EVEN: "even_page_header",
}
_FOOTER_ATTRS = {
    HeaderFooterKind.DEFAULT: "footer",
    HeaderFooterKind.FIRST: "first_page_footer",
    HeaderFooterKind.EVEN: "even_page_footer",
}


def _collect_parts(document, header: bool) -> list[dict]:
    """按 ``(OOXML 部件身份 partname, 角色 kind)`` 去重收集页眉/页脚部件。

    返回 ``[{"kind", "element", "sections"}]``；``sections`` 为该部件在该角色下
    适用的全部 section 序号（含继承）。

    去重键必须包含角色：同一部件被不同 section 以 default/first/even 不同角色
    引用时（Word 常见于继承或显式共享关系），每种角色都要独立保留，不得因
    按 partname 去重而把 first/even 角色折叠进 default。第一页页眉是权威元
    信息来源，同样不得被静默省略。
    """
    attrs = _HEADER_ATTRS if header else _FOOTER_ATTRS
    by_key: dict[tuple[str, HeaderFooterKind], dict] = {}
    for section_index, section in enumerate(document.sections):
        for kind, attr in attrs.items():
            part = getattr(section, attr)
            if part is None:
                continue
            name = str(part.part.partname)
            key = (name, kind)
            entry = by_key.get(key)
            if entry is None:
                entry = {"kind": kind, "element": part._element, "sections": set()}
                by_key[key] = entry
            entry["sections"].add(section_index)
    return [
        {
            "kind": entry["kind"],
            "element": entry["element"],
            "sections": tuple(sorted(entry["sections"])),
        }
        for entry in by_key.values()
    ]


def _part_has_content(element) -> bool:
    for node in element.iter():
        if node.tag == _W_T and (node.text or "").strip():
            return True
    return element.find(_W_TBL) is not None


def _walk_headers_footers(document, extractor: _Extractor) -> list[object]:
    """捕获 default/first/even 页眉页脚，去重继承部件并保留 section 归属。"""
    part_elements: list[object] = []
    for part in _collect_parts(document, header=True):
        if not _part_has_content(part["element"]):
            continue
        kind = part["kind"]
        sections = part["sections"]
        owner = sections[0]
        prefix = f"header.{kind.value}.s{owner}"
        extractor.walk_paragraphs_and_tables(
            part["element"],
            document_part=DocumentPart.HEADER,
            section_index=owner,
            prefix=prefix,
            part_kind=kind,
            section_indexes=sections,
        )
        part_elements.append(part["element"])
    for part in _collect_parts(document, header=False):
        if not _part_has_content(part["element"]):
            continue
        kind = part["kind"]
        sections = part["sections"]
        owner = sections[0]
        prefix = f"footer.{kind.value}.s{owner}"
        extractor.walk_paragraphs_and_tables(
            part["element"],
            document_part=DocumentPart.FOOTER,
            section_index=owner,
            prefix=prefix,
            part_kind=kind,
            section_indexes=sections,
        )
        part_elements.append(part["element"])
    return part_elements


def _count_tracked_changes(elements) -> int:
    total = 0
    for el in elements:
        for node in el.iter():
            if node.tag in (_W_INS, _W_DEL):
                total += 1
    return total


def block_set_hash(blocks: list[StructureBlock] | tuple[StructureBlock, ...]) -> str:
    """块集规范化序列化的 SHA-256，与 :func:`serialize_blocks` 一致。"""
    text = serialize_blocks(blocks)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def serialize_blocks(blocks: list[StructureBlock] | tuple[StructureBlock, ...]) -> str:
    """块集规范化 JSON 文本（sort_keys + 紧凑分隔符，与 canonical_hash 同构）。"""
    payload = [block.model_dump(mode="json") for block in blocks]
    return json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _anomaly(anomaly_id: str, kind: str, message: str) -> ExtractionAnomaly:
    return ExtractionAnomaly(anomaly_id=anomaly_id, kind=kind, message=message)


def _persist_block_blob(
    output_dir: Path, content_storage_ref: str, content_sha256: str, blocks: list[StructureBlock]
) -> Path:
    """原子落盘规范块集工件并写后校验 SHA-256。

    - 既有 blob 若哈希一致则直接复用；若不一致（损坏/占位）则原子修复；
    - 写入走临时文件 + ``os.replace`` 原子落位，写后回读校验哈希，
      不一致抛 :class:`StructureExtractionError`，绝不留下指向损坏工件
      或不存在工件的成功快照。
    """
    blob_path = Path(output_dir) / content_storage_ref
    blob_path.parent.mkdir(parents=True, exist_ok=True)
    if blob_path.is_file():
        existing_sha = hashlib.sha256(blob_path.read_bytes()).hexdigest()
        if existing_sha == content_sha256:
            return blob_path
    text = serialize_blocks(blocks)
    tmp_path = blob_path.with_name(blob_path.name + ".tmp")
    try:
        tmp_path.write_text(text, encoding="utf-8")
        os.replace(tmp_path, blob_path)
    finally:
        tmp_path.unlink(missing_ok=True)
    written_sha = hashlib.sha256(blob_path.read_bytes()).hexdigest()
    if written_sha != content_sha256:
        raise StructureExtractionError(
            f"规范块集工件写入后哈希不一致：期望 {content_sha256}，实际 {written_sha}"
        )
    return blob_path


def extract_docx_structure(
    docx_path: str | Path,
    *,
    snapshot_id: str,
    source_artifact: ProtocolSourceArtifact,
    output_dir: str | Path,
    render_artifact_ids: list[str] | None = None,
    parser_name: str = PARSER_NAME,
    parser_version: str = PARSER_VERSION,
    created_at: datetime | None = None,
) -> StructureExtraction:
    """提取 DOCX 结构并产出块集与不可变快照。

    - 源文件只读；块集原子落盘到 ``output_dir``（``content_storage_ref`` 指向的
      规范内容工件）并写后校验 SHA-256，成功快照绝不指向不存在或不一致的工件。
    - 零段落且零表格视为失败；任意部件存在未接受修订标记或无法解析的自动编号
      进入「需要核对」。
    """
    file_path = Path(docx_path)
    _mime, kind = detect_format(file_path)
    if kind != ProtocolFileKind.DOCX:
        raise StructureExtractionError(
            f"仅支持 DOCX 结构提取，检测到 {kind.value}：{file_path}"
        )
    actual_sha256 = compute_sha256(file_path)
    if actual_sha256 != source_artifact.sha256:
        raise StructureExtractionError(
            "待提取方案与登记文件哈希不一致，已停止解析；请重新登记原始方案。"
        )

    try:
        document = Document(str(file_path))
    except Exception as exc:
        raise StructureExtractionError(
            f"DOCX 解析失败（格式损坏或部件缺失）：{file_path}（{exc}）"
        ) from exc

    abstracts, nums = _numbering_definitions(document)
    style_numbering = _style_numbering_definitions(document)
    extractor = _Extractor(abstracts, nums, style_numbering)
    _walk_body(document, extractor)
    header_footer_parts = _walk_headers_footers(document, extractor)

    unhandled_tags = _unhandled_content_tags(document.element.body)
    for part in header_footer_parts:
        unhandled_tags.update(_unhandled_content_tags(part))
    for tag in sorted(unhandled_tags):
        extractor.anomalies.append(
            ExtractionAnomaly(
                anomaly_id=f"unhandled_ooxml_container:{tag}",
                kind="unhandled_ooxml_container",
                message=f"发现尚未支持的方案内容容器 {tag}，需核对是否存在漏提内容",
                scope_ref=tag,
            )
        )

    tracked_change_count = _count_tracked_changes(
        [document.element.body, *header_footer_parts]
    )

    coverage = ExtractionCoverage(
        paragraph_count=extractor.paragraph_count,
        table_count=extractor.table_count,
        section_count=len(document.sections),
        nested_table_count=extractor.nested_table_count,
        header_detected=extractor.header_detected,
        footer_detected=extractor.footer_detected,
        tracked_change_count=tracked_change_count,
    )

    anomalies: list[ExtractionAnomaly] = list(extractor.anomalies)
    status = ExtractionStatus.COMPLETED
    if coverage.paragraph_count == 0 and coverage.table_count == 0:
        status = ExtractionStatus.FAILED
        anomalies.append(
            _anomaly("empty_document", "empty_document", "未提取到任何段落或表格")
        )
    if coverage.tracked_change_count > 0:
        status = ExtractionStatus.NEEDS_REVIEW
        anomalies.append(
            _anomaly(
                "unaccepted_tracked_changes",
                "unaccepted_tracked_changes",
                f"存在 {coverage.tracked_change_count} 处未接受的修订标记",
            )
        )
    if any(a.kind == "unresolved_numbering" for a in extractor.anomalies):
        status = ExtractionStatus.NEEDS_REVIEW
    if any(a.kind == "unhandled_ooxml_container" for a in extractor.anomalies):
        status = ExtractionStatus.NEEDS_REVIEW

    blocks = extractor.blocks
    content_sha256 = block_set_hash(blocks)
    content_storage_ref = f"{_BLOCKS_STORAGE_PREFIX}/{content_sha256}.json"
    _persist_block_blob(
        Path(output_dir), content_storage_ref, content_sha256, blocks
    )

    snapshot = ProtocolExtractionSnapshot(
        snapshot_id=snapshot_id,
        source_artifact_id=source_artifact.source_artifact_id,
        source_sha256=source_artifact.sha256,
        parser_name=parser_name,
        parser_version=parser_version,
        status=status,
        coverage=coverage,
        anomalies=anomalies,
        render_artifact_ids=render_artifact_ids or [],
        content_sha256=content_sha256,
        content_storage_ref=content_storage_ref,
        created_at=created_at or utc_now(),
    )
    return StructureExtraction(blocks=tuple(blocks), snapshot=snapshot)
