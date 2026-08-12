"""Protocol deconstruction — extract eligibility criteria from DOCX/PDF/TXT.

Contains the LLM prompt for deconstructing clinical trial protocols and
the DOCX number-aware text extractor.
"""

from __future__ import annotations

import re
from pathlib import Path

DECONSTRUCT_PROMPT = """你是一位资深的临床试验医学监查专家，擅长从研究方案中解构入选排除标准。

请仔细阅读以下研究方案文档内容，完成以下任务：

## 核心原则：严格保持原文结构

**必须严格按照方案原文的编号体系提取标准，不得拆分、合并或重新编号。**
**不得省略任何入选或排除标准。** 如果输入中已经提供了“入选标准：N条、排除标准：M条”的预提取信息，你的输出必须覆盖对应数量的父级条款。输出前必须自检：最后一个排除标准也必须被输出，不能只输出前几条排除标准。
**父级IN/EX数量、顺序、编号必须与方案原文在本次选择期别下的父级条款完全一致。** 例如方案本期排除标准有16条，输出必须且只能有EX-01至EX-16；不得把子项、检查项目、流程表项目升级成EX-17或新的父级条目。
**如果方案同时包含Ⅱ期和Ⅲ期，用户必须先选择本次解构Ⅱ期还是Ⅲ期。** 一旦用户选择期别，本次输出应把该期别视为一个独立项目规则集：只读取并使用所选期别对应的IN/EX父级条款、编号和范围；另一期别只用于判断是否需要用户选择，不得进入正式规则，不得用于比较，不得输出一致/差异说明。
**不要把“本次选择某期别”写成规则适用性标签。** 正式IN/EX条目内不得出现“仅Ⅱ期/仅Ⅲ期”“适用期别：Ⅱ期/Ⅲ期”“按本次所选期别编号”“与另一阶段对应条款一致/差异”等由选择动作产生的说明；选定期别后的规则应像单独方案一样直接呈现。

方案原文通常用数字编号（如1）、2）、3）...）标记每条标准。有些标准包含子项（如"符合以下任一项"后列出a. b. c.），这些子项**必须保留在同一条标准内**，不得拆成独立条目。

常见需要保持合并的结构：
- "符合以下任一项结核筛查标准：" → 这是**1条**标准，其下子项是判断条件
- "存在以下任何一种实验室检查异常：" → 这是**1条**标准
- "筛选访视时存在下列任一感染者：" → 这是**1条**标准
- "规定的时间内接受了以下任何一种治疗者：" → 这是**1条**标准（洗脱期）

允许的复杂条款拆解方式：
- 父级条目仍为`IN-04`或`EX-06`；
- 父级条目内可以列`IN-04a`、`IN-04b`或`IN-04-GROUP`作为组件/情形组；
- 必须写清组件之间的逻辑关系（all/any/阈值/时间窗/研究者判断）；
- 子项不得变成新的父级`IN-05`或`EX-07`，除非方案原文父级编号本来如此。

## 逻辑关系保真（非常重要）

方案原文中的逻辑词必须原样保真到“通过/不通过”条件中：
- `任一`、`任何一种`、`或` 通常表示 OR，只要满足任一并达到完整触发条件即可；
- `且`、`并且`、`同时`、`经研究者评估`、`如果/将可能` 通常表示 AND，必须同时满足全部组件，不得弱化成 OR；
- 如果一个父级条款下某个子项本身包含 AND 条件，父级“不通过”摘要必须保留该子项内部的 AND 条件，不能写成“存在任一异常即不通过”。

典型例子：`任何其它实验室检查结果异常且有临床意义，经研究者评估如果参与研究将可能对参与者构成不可接受的风险`
必须解构为三项同时满足：1）其它实验室检查异常；2）该异常有临床意义；3）经研究者评估参与研究将可能构成不可接受风险。
尿糖1+、潜血阳性、GGT升高等单个异常本身不是该子项的完整排除条件；缺少研究者不可接受风险评估时，只能写为需研究者判定/证据不足，不能写为不通过。

典型例子：`存在重大或不稳定系统性疾病，研究者明确判断不具备临床研究条件`
必须解构为两项同时满足：1）存在方案定义范围内的重大或不稳定系统性疾病；2）研究者明确判断不具备临床研究条件。
只有疾病、异常、既往用药或病史存在，或者只写“未明确判断具备条件”“必要时专科就诊”“疗效不详”，不得写为不通过；应写为需研究者判定/证据不足。

典型例子：`梅毒特异性抗体试验阳性（梅毒非特异性抗体阴性且研究者判断为既往感染已治愈的除外）`
必须解构为：TPPA/TP-Ab/梅毒特异性抗体阳性是触发项；只有同时具备“TRUST/RPR/梅毒非特异性抗体阴性 + 研究者明确判断既往感染已治愈”才满足例外。
目前无不适、未予治疗、非活动性感染、研究者笼统写可以入组，均不能替代既往感染已治愈判断。

## 任务

1. **提取所有入选标准**：逐条列出，保持原文编号
2. **提取所有排除标准**：逐条列出，保持原文编号
3. **对每条标准进行解构**，明确：
   - **判断点**：需要在入排审核资料中核实的具体内容
   - **通过条件**：明确的判定标准
   - **不通过条件**：排除/不入选的条件
   - **证据不足条件**（如适用）：什么情况下判定为证据不足
   - **子项列表**（如适用）：将子项作为该条标准的判断条件列出
4. **识别情形组**（如适用）：如需要多个时间点均达标的情况

## 输出格式要求

严格按以下Markdown格式输出：

```
---
项目代号: [从方案中提取]
方案编码: [从方案中提取]
---

# [项目代号] 入排审核规则

**排除条（EX-*）**：**通过** = 无该排除情形；**不通过** = 存在排除。**锚点**：见证据包「锚点日期」。

---

## 一、入选标准（Inclusion）

#### IN-01 [标准名称]

- **判断点**：[具体需要核实的内容]
- **通过**：[通过条件]
- **不通过**：[不通过条件]
- **证据不足**（如适用）：[证据不足条件]

[继续其他入选标准...]

---

## 二、排除标准（Exclusion）

#### EX-01 [标准名称]

- **判断点**：[具体需要核实的内容]
- **子项**（如适用）：
  - a. [子项1]
  - b. [子项2]
- **通过**：[通过条件]
- **不通过**：[不通过条件]

[继续其他排除标准...]
```

## 重要注意事项

1. **严格按原文编号**：方案标了多少条就是多少条，不得拆分子项为独立条目
2. **子项保留在父条目内**：如"实验室检查异常"是一条，其下的WBC、Hb等是子项
3. **逻辑关系不得弱化**：AND条件必须全部满足，OR条件才可任一满足；不要把“异常且有临床意义且经研究者评估不可接受风险”改写成“存在异常即不通过”
4. 时间窗必须明确标注（如"首次给药前3个月"、"筛选期"等）
5. 数值阈值必须精确（如"PASI≥12分"、"WBC<3×10⁹/L"）
6. 例外情况必须明确列出
7. 对于需要研究者判断的条款，标注"需研究者判定"
8. 保持与方案原文的一致性，不要自行添加或修改标准内容
9. 如果方案同时包含Ⅱ期和Ⅲ期，只输出用户选择期别的父级IN/EX标准；不得把Ⅱ期/Ⅲ期合并后改写成新的项目编号、方案编号或重复父级编号；不得在正式规则中比较或提及另一期别。
10. 对于研究流程表中筛选期、基线/随机前及基线以前必须完成的检查、检验、评分、知情同意、日志卡、合并用药/治疗、随机前限制等，应在入排标准之后追加“基线及以前方案流程核查”章节。
11. 若无法完整覆盖所有父级条款，应明确标注“证据不足：方案条文提取不完整”，不得输出看似完整但实际缺条的规则。
12. “基线及以前方案流程核查”不是正式IN/EX父级条款，不得改变IN/EX数量。
"""


def extract_protocol_text(doc_path: str) -> str:
    """Extract readable text from a protocol file without calling an LLM."""
    path = Path(doc_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        import fitz

        doc = fitz.open(str(path))
        try:
            return "\n".join(page.get_text() for page in doc)
        finally:
            doc.close()

    if suffix in (".docx", ".doc"):
        from docx import Document

        doc = Document(str(path))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    raise ValueError(f"不支持的文件格式: {suffix}")


def extract_protocol_front_matter_text(doc_path: str) -> str:
    """Extract metadata-priority text from headers, footers, and first pages."""
    path = Path(doc_path)
    suffix = path.suffix.lower()
    parts = []

    if suffix == ".pdf":
        import fitz

        doc = fitz.open(str(path))
        try:
            for page in list(doc)[:2]:
                parts.append(page.get_text())
        finally:
            doc.close()
    elif suffix in (".docx", ".doc"):
        from docx import Document

        doc = Document(str(path))
        for section in doc.sections:
            for block in (
                section.first_page_header,
                section.header,
                section.first_page_footer,
                section.footer,
            ):
                for para in block.paragraphs:
                    if para.text.strip():
                        parts.append(para.text.strip())
                for table in block.tables:
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            parts.append(" | ".join(cells))

        for para in doc.paragraphs[:120]:
            if para.text.strip():
                parts.append(para.text.strip())
        for table in doc.tables[:5]:
            for row in table.rows[:60]:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
    elif suffix == ".txt":
        parts.append(path.read_text(encoding="utf-8", errors="replace")[:12000])
    else:
        raise ValueError(f"不支持的文件格式: {suffix}")

    seen = set()
    deduped = []
    for item in parts:
        item = re.sub(r"\s+", " ", item or "").strip()
        if item and item not in seen:
            deduped.append(item)
            seen.add(item)
    return "\n".join(deduped)


def _normalize_protocol_date(value: str) -> str:
    value = (value or "").strip()
    compact = re.search(r"(20\d{2})(\d{2})(\d{2})", value)
    if compact:
        year, month, day = compact.groups()
        return f"{year}-{month}-{day}"
    zh = re.search(r"(20\d{2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", value)
    if zh:
        year, month, day = zh.groups()
        return f"{year}-{int(month):02d}-{int(day):02d}"
    dashed = re.search(r"(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})", value)
    if dashed:
        year, month, day = dashed.groups()
        return f"{year}-{int(month):02d}-{int(day):02d}"
    return ""


def _version_candidates(source: str) -> list[str]:
    values = []
    for hit in re.finditer(r"(?:版本号|版本|Version)\s*[：:]?\s*(V?\s*\d+(?:\.\d+)*)", source, flags=re.I):
        context = source[max(0, hit.start() - 80):hit.end() + 80]
        if _is_non_protocol_version_context(context):
            continue
        values.append(hit.group(1))
    for hit in re.finditer(r"(?<![A-Za-z0-9])V\s*\d+(?:\.\d+)*(?![\d.])", source, flags=re.I):
        context = source[max(0, hit.start() - 80):hit.end() + 80]
        if _is_non_protocol_version_context(context):
            continue
        values.append(hit.group(0))
    return [re.sub(r"\s+", "", v).upper() for v in values if v]


def _date_candidates(source: str) -> list[str]:
    values = []
    for date_hit in re.findall(
        r"20\d{6}|20\d{2}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日|20\d{2}[-/.]\d{1,2}[-/.]\d{1,2}",
        source,
    ):
        normalized = _normalize_protocol_date(date_hit)
        if normalized:
            values.append(normalized)
    return values


def _format_protocol_version(value: str) -> str:
    version = re.sub(r"\s+", "", value or "").upper().strip("：:|/／ ")
    if version and not version.startswith("V"):
        version = f"V{version}"
    return version


def _is_non_protocol_version_context(context: str) -> bool:
    """Avoid treating template/database/form versions as protocol versions."""
    normalized = re.sub(r"\s+", "", context or "").lower()
    protocol_tokens = (
        "方案编号",
        "方案编码",
        "方案版本",
        "方案版本号",
        "方案版本日期",
        "研究方案",
        "clinicalstudyprotocol",
        "protocol",
    )
    if any(token in normalized for token in protocol_tokens):
        return False
    noise_tokens = (
        "文件编号",
        "cmss-sop",
        "sop",
        "数据库版本",
        "crfver",
        "表单版本",
        "edc版本",
        "系统版本",
        "页面版本",
        "版本控制",
        "修订记录",
    )
    return any(token in normalized for token in noise_tokens)


def _version_date_from_labeled_text(source: str) -> tuple[str, str]:
    date_pat = r"(?:20\d{6}|20\d{2}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日|20\d{2}[-/.]\d{1,2}[-/.]\d{1,2})"
    version_pat = r"(?:[Vv]\s*)?\d+(?:\.\d+)*"
    labels = r"(?:方案)?版本(?:号)?\s*/\s*(?:版本)?日期|版本号\s*/\s*日期|版本\s*/\s*日期"
    patterns = [
        rf"(?:{labels})\s*[：:\|]?\s*({version_pat})\s*(?:[/／]+|\|)?\s*({date_pat})",
        rf"(?:方案版本号|方案版本|版本号|版本)\s*[：:\|]\s*({version_pat}).{{0,80}}?(?:方案版本日期|版本日期|日期)\s*[：:\|]\s*({date_pat})",
        rf"(?:方案版本日期|版本日期)\s*[：:\|]\s*({date_pat}).{{0,80}}?(?:方案版本号|方案版本|版本号|版本)\s*[：:\|]\s*({version_pat})",
        rf"(?:方案编号|方案编码|方案号|Protocol\s*(?:No\.?|Number|ID|Code)?).{{0,120}}?(?:方案版本号|方案版本|版本号|版本)\s*[：:\|]\s*({version_pat}).{{0,80}}?(?:方案版本日期|版本日期|日期)\s*[：:\|]\s*({date_pat})",
    ]
    for idx, pattern in enumerate(patterns):
        found = re.search(pattern, source, flags=re.I | re.S)
        if not found:
            continue
        context = source[max(0, found.start() - 80):found.end() + 80]
        if _is_non_protocol_version_context(context):
            continue
        if idx == 2:
            date_value, version_value = found.group(1), found.group(2)
        else:
            version_value, date_value = found.group(1), found.group(2)
        return _format_protocol_version(version_value), _normalize_protocol_date(date_value)
    return "", ""


def _best_version_candidate(source: str) -> str:
    protocol_like = []
    for hit in re.finditer(r"(?:方案版本号|方案版本|版本号|版本)\s*[：:]?\s*(V?\s*\d+(?:\.\d+)*)", source, flags=re.I):
        context = source[max(0, hit.start() - 80):hit.end() + 80]
        if _is_non_protocol_version_context(context):
            continue
        value = _format_protocol_version(hit.group(1))
        if value:
            protocol_like.append(value)
    candidates = protocol_like or _version_candidates(source)
    if not candidates:
        return ""
    return sorted(set(candidates), key=lambda v: ("." in v, len(v), v), reverse=True)[0]


def _best_protocol_date_candidate(source: str) -> str:
    date_pat = r"20\d{6}|20\d{2}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日|20\d{2}[-/.]\d{1,2}[-/.]\d{1,2}"
    for hit in re.finditer(date_pat, source):
        context = source[max(0, hit.start() - 120):hit.end() + 120]
        compact = re.sub(r"\s+", "", context)
        if ("版本日期" in compact or "方案日期" in compact or "方案版本" in compact) and not _is_non_protocol_version_context(context):
            normalized = _normalize_protocol_date(hit.group(0))
            if normalized:
                return normalized
    dates = _date_candidates(source)
    return dates[0] if dates else ""


def _clean_metadata_name(
    stem: str,
    protocol_id: str,
    project_code: str,
    version: str,
    protocol_date: str,
) -> str:
    name = stem.replace("_", " ").replace("-", " ")
    for token in (protocol_id, project_code, version, protocol_date, protocol_date.replace("-", "")):
        if token:
            name = name.replace(token, " ")
            name = name.replace(token.replace("-", " "), " ")
    name = re.sub(r"\bclean\b|清洁版|clean版|定稿|正式版", " ", name, flags=re.I)
    name = re.sub(r"\s+", " ", name).strip(" _-")
    if not name or len(name) < 3:
        name = f"{project_code or protocol_id} 临床研究方案".strip()
    return name


def _is_generic_protocol_title(value: str) -> bool:
    cleaned = re.sub(r"\s+", "", value or "").strip("：:")
    return cleaned.lower() in {
        "临床研究方案",
        "研究方案",
        "方案",
        "clinicalstudyprotocol",
        "protocol",
    }


def extract_protocol_metadata(doc_path: str) -> dict:
    """Infer project metadata from a protocol filename and document body.

    This deliberately uses deterministic parsing rather than an LLM so project
    creation stays fast and reproducible.
    """
    path = Path(doc_path)
    stem = path.stem.strip()
    front_text = extract_protocol_front_matter_text(doc_path)
    text = extract_protocol_text(doc_path)
    priority_probe = f"{front_text}\n{stem}"
    probe = f"{priority_probe}\n{text[:12000]}"

    protocol_id = ""
    labeled_patterns = [
        r"(?:方案编号|方案编码|方案代码|方案号)\s*[：:]\s*([A-Za-z0-9][A-Za-z0-9_.-]{2,})",
        r"(?:Protocol\s*(?:No\.?|Number|ID|Code)?)\s*[：:]\s*([A-Za-z0-9][A-Za-z0-9_.-]{2,})",
    ]
    for source in (front_text, stem, probe):
        for pattern in labeled_patterns:
            found = re.search(pattern, source, flags=re.I)
            if found:
                protocol_id = found.group(1).strip(" ._-/")
                break
        if protocol_id:
            break

    if not protocol_id:
        candidates = re.findall(
            r"(?<![A-Za-z0-9])([A-Z][A-Z0-9]*(?:-[A-Z0-9]+){1,})(?![A-Za-z0-9-])",
            probe,
            flags=re.I,
        )
        if candidates:
            candidates = sorted({c.strip(" ._-/") for c in candidates}, key=lambda x: (-x.count("-"), -len(x)))
            protocol_id = candidates[0]

    version, protocol_date = _version_date_from_labeled_text(front_text)
    if not version:
        version = _best_version_candidate(front_text) or _best_version_candidate(stem) or _best_version_candidate(probe)
    if version:
        version = _format_protocol_version(version)

    if not protocol_date:
        for source in (front_text, stem, probe):
            protocol_date = _best_protocol_date_candidate(source)
            if protocol_date:
                break

    project_code = protocol_id
    if protocol_id:
        project_code = re.sub(r"[-_]\d{2,4}$", "", protocol_id).strip("-_") or protocol_id

    title = ""
    title_hit = re.search(r"(?:方案标题|研究题目|Protocol\s*Title)\s*[：:\|]?\s*(.+)", front_text, flags=re.I)
    if title_hit:
        title = re.sub(r"\s+", " ", title_hit.group(1)).strip(" |")
    for line in text.splitlines()[:300]:
        if title:
            break
        line = re.sub(r"\s+", " ", line).strip()
        if (
            4 <= len(line) <= 160
            and not _is_generic_protocol_title(line)
            and ("研究方案" in line or "临床试验方案" in line or "Clinical Study Protocol" in line)
        ):
            title = line
            break

    name = title or _clean_metadata_name(stem, protocol_id, project_code, version, protocol_date)
    if name in {protocol_id, project_code} or _is_generic_protocol_title(name):
        name = f"{project_code} 临床研究方案"

    return {
        "project_code": project_code or stem,
        "protocol_id": protocol_id or stem,
        "name": name,
        "protocol_version": version,
        "protocol_date": protocol_date,
        "protocol_source_filename": path.name,
    }


def _normalize_study_stage_label(value: str) -> str:
    text = re.sub(r"\s+", "", value or "")
    if text in {"2期", "二期", "II期", "Ⅱ期", "II", "Ⅱ"}:
        return "Ⅱ期"
    if text in {"3期", "三期", "III期", "Ⅲ期", "III", "Ⅲ"}:
        return "Ⅲ期"
    return value.strip()


def sanitize_selected_stage_scope_language(text: str, selected_stage: str, workflow: dict | None = None) -> str:
    """Remove mixed-stage wording from generated rules after a phase is chosen."""
    stage = _normalize_study_stage_label(selected_stage)
    if not text or not stage:
        return text
    stages = (workflow or {}).get("study_stages") or []
    if len(stages) <= 1:
        return text

    if stage == "Ⅲ期":
        stage_alias = r"(?:Ⅲ期|III期|III\s*期|3期|三期)"
    elif stage == "Ⅱ期":
        stage_alias = r"(?:Ⅱ期|II期|II\s*期|2期|二期)"
    else:
        stage_alias = re.escape(stage)

    per_rule_stage = re.compile(
        rf"^\s*[-*]?\s*(?:\*\*)?(?:适用期别|适用阶段|适用范围)(?:\*\*)?\s*[：:].*{stage_alias}.*$",
        flags=re.I,
    )
    note_line = re.compile(r"^\s*[-*]?\s*(?:\*\*)?(?:说明|备注|注)(?:\*\*)?\s*[：:]", flags=re.I)
    any_stage = re.compile(r"(?:Ⅱ期|II\s*期|2期|二期|Ⅲ期|III\s*期|3期|三期)", flags=re.I)
    mixed_terms = ("对应", "另一阶段", "另一期别", "实质一致", "一致", "差异", "按本次所选")
    selected_stage_note = re.compile(
        rf"^\s*[-*]?\s*(?:\*\*)?(?:说明|备注|注)(?:\*\*)?\s*[：:].*(?:仅\s*{stage_alias}|按本次所选\s*{stage_alias}|{stage_alias}\s*编号).*$",
        flags=re.I,
    )
    inline_only = re.compile(rf"仅\s*{stage_alias}", flags=re.I)
    inline_selected_numbering = re.compile(rf"按本次所选\s*{stage_alias}\s*编号", flags=re.I)

    sanitized_lines: list[str] = []
    for line in (text or "").splitlines():
        if per_rule_stage.search(line) or selected_stage_note.search(line):
            continue
        if note_line.search(line) and any_stage.search(line) and any(term in line for term in mixed_terms):
            continue
        if "本次解构期别" not in line and inline_only.search(line):
            line = inline_only.sub(stage, line)
        if "本次解构期别" not in line and inline_selected_numbering.search(line):
            line = inline_selected_numbering.sub(stage, line)
        sanitized_lines.append(line)
    return "\n".join(sanitized_lines)


def strip_outer_markdown_fence(text: str) -> str:
    """Remove a single whole-document Markdown code fence from LLM output."""
    stripped = (text or "").strip()
    lines = stripped.splitlines()
    if len(lines) >= 2 and lines[0].strip().startswith("```") and lines[-1].strip() == "```":
        return "\n".join(lines[1:-1]).strip()
    return text


def _replace_or_append_line(block: str, label: str, replacement: str) -> str:
    pattern = rf"- \*\*{re.escape(label)}\*\*：.*"
    if re.search(pattern, block):
        return re.sub(pattern, replacement, block, count=1)
    return block.rstrip() + "\n" + replacement + "\n"


def strengthen_compound_condition_summaries(text: str) -> str:
    """Keep generated rule summaries from flattening subitem-internal AND logic.

    LLM deconstruction can preserve the literal subitem text but then summarize
    the parent "不通过" condition as if every mentioned abnormality were an OR
    trigger. This deterministic pass repairs common compound patterns where a
    factual component must remain tied to an explicit investigator judgment or
    complete protocol exception.
    """
    source = text or ""

    def strengthen_lab_block(match: re.Match[str]) -> str:
        block = match.group(0)
        if not re.search(r"(任何其它|任何其他|其它|其他).{0,12}实验室检查结果异常", block):
            return block
        if "有临床意义" not in block or "不可接受" not in block:
            return block
        block = _replace_or_append_line(
            block,
            "通过",
            "- **通过**：筛选和基线时未见a-g任一明确阈值异常；h项未见同时满足“其它实验室检查异常 + 有临床意义 + 经研究者评估参与研究将可能构成不可接受风险”。",
        )
        block = _replace_or_append_line(
            block,
            "不通过",
            "- **不通过**：筛选或基线时存在a-g任一明确阈值异常；或h项同时满足“其它实验室检查异常 + 有临床意义 + 经研究者评估参与研究将可能构成不可接受风险”。h项缺少任一组件时不得直接判不通过。",
        )
        return block

    def strengthen_systemic_disease_block(match: re.Match[str]) -> str:
        block = match.group(0)
        if not re.search(r"(重大|不稳定).{0,24}系统性疾病", block):
            return block
        block = _replace_or_append_line(
            block,
            "通过",
            "- **通过**：未见首次给药前6个月内重大或不稳定的系统性疾病；或虽有相关疾病但研究者明确判断具备临床研究条件。",
        )
        block = _replace_or_append_line(
            block,
            "不通过",
            "- **不通过**：同时满足“存在相关系统性疾病 + 研究者明确判断不具备临床研究条件”。仅有疾病存在、仅嘱必要时专科就诊或未明确判断具备/不具备时不得直接判不通过。",
        )
        block = _replace_or_append_line(
            block,
            "证据不足",
            "- **证据不足**：存在相关疾病但缺少稳定性、时间窗或研究者临床研究条件判断；或病史资料不足。",
        )
        return block

    def strengthen_vitals_exam_block(match: re.Match[str]) -> str:
        block = match.group(0)
        if not re.search(r"(生命体征|体格检查|心电图|ECG|CT)", block):
            return block
        if "不可接受" not in block and "有临床意义" not in block:
            return block
        block = _replace_or_append_line(
            block,
            "通过",
            "- **通过**：筛选/基线生命体征、体格检查、12导联心电图、胸部CT未见异常；或虽有异常但无临床意义，且未见研究者评估参与研究将构成不可接受风险。",
        )
        block = _replace_or_append_line(
            block,
            "不通过",
            "- **不通过**：同时满足“上述检查异常 + 异常有临床意义 + 经研究者评估参与研究将可能构成不可接受风险”。缺少临床意义或不可接受风险判断时不得直接判不通过。",
        )
        block = _replace_or_append_line(
            block,
            "证据不足",
            "- **证据不足**：检查结果缺失；或存在异常但缺少研究者对临床意义及参与研究风险的明确判断。",
        )
        return block

    def strengthen_syphilis_block(match: re.Match[str]) -> str:
        block = match.group(0)
        if "梅毒特异" not in block or "非特异" not in block:
            return block
        block = _replace_or_append_line(
            block,
            "通过",
            "- **通过**：HBsAg阴性，且未见“HBcAb阳性 + HBV-DNA阳性”；HCV抗体阴性或HCV-RNA阴性；HIV Ab阴性；梅毒筛查阴性，或TPPA/TP-Ab/梅毒特异性抗体阳性但同时满足“非特异性抗体阴性 + 研究者明确判断既往感染已治愈”的完整例外。",
        )
        block = _replace_or_append_line(
            block,
            "不通过",
            "- **不通过**：HBsAg阳性；或HBcAb阳性且HBV-DNA阳性；或HCV抗体阳性且HCV-RNA阳性；或HIV Ab阳性；或TPPA/TP-Ab/梅毒特异性抗体阳性且未完整满足“非特异性抗体阴性 + 研究者明确判断既往感染已治愈”的例外。缺少治愈判断时不得按例外通过。",
        )
        block = _replace_or_append_line(
            block,
            "证据不足",
            "- **证据不足**：缺少规定感染筛查结果；或梅毒特异性抗体阳性但缺少非特异性抗体结果/研究者既往感染已治愈判断，需补充后复核。",
        )
        return block

    source = re.sub(
        r"(?ms)^####\s+EX-\d+[^\n]*实验室检查异常.*?(?=^####\s+(?:IN|EX)-\d+|\Z)",
        strengthen_lab_block,
        source,
    )
    source = re.sub(
        r"(?ms)^####\s+EX-\d+[^\n]*(?:重大|不稳定)[^\n]*系统性疾病.*?(?=^####\s+(?:IN|EX)-\d+|\Z)",
        strengthen_systemic_disease_block,
        source,
    )
    source = re.sub(
        r"(?ms)^####\s+EX-\d+[^\n]*(?:生命体征|体格检查|ECG|心电图|CT).*?(?=^####\s+(?:IN|EX)-\d+|\Z)",
        strengthen_vitals_exam_block,
        source,
    )
    source = re.sub(
        r"(?ms)^####\s+EX-\d+[^\n]*(?:感染|梅毒).*?(?=^####\s+(?:IN|EX)-\d+|\Z)",
        strengthen_syphilis_block,
        source,
    )
    return source


def extract_docx_criteria(doc_path: str, study_stage: str = "") -> dict:
    """Extract inclusion/exclusion criteria from DOCX using paragraph numbering.
    
    Returns dict with keys: 'inclusion_text', 'exclusion_text',
    'inclusion_count', 'exclusion_count'.
    Uses the body headings and Word's numId to distinguish parent criteria from
    sub-items. Table-of-contents entries are ignored because they can otherwise
    cause schedule rows to be counted as exclusion criteria.
    """
    from docx import Document
    
    doc = Document(doc_path)

    def is_heading(para, text: str, label: str) -> bool:
        cleaned = text.rstrip("：:")
        style = para.style.name or ""
        style_lower = style.lower()
        return cleaned == label and (
            style_lower.startswith("heading")
            or "标题" in style
            or style_lower.startswith("title")
        )

    def is_section_heading(para) -> bool:
        style = para.style.name or ""
        style_lower = style.lower()
        return (
            style_lower.startswith("heading")
            or "标题" in style
            or style_lower.startswith("title")
        )

    def para_num_id(para) -> str | None:
        numPr = para._element.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr')
        if numPr is None:
            return None
        numIdEl = numPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numId')
        if numIdEl is None:
            return None
        return numIdEl.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')

    def collect_group(start: int, end: int, parent_num_id: str | None = None) -> tuple[list[dict], str | None]:
        items: list[dict] = []
        active_parent_num_id = parent_num_id
        for para in doc.paragraphs[start:end]:
            text = para.text.strip()
            if not text:
                continue
            num_id = para_num_id(para)
            if not num_id:
                if items and not is_section_heading(para):
                    items[-1]["sub_items"].append(text)
                continue
            if active_parent_num_id is None:
                active_parent_num_id = num_id
            if num_id == active_parent_num_id:
                items.append({"text": text, "sub_items": []})
            elif items:
                items[-1]["sub_items"].append(text)
        return items, active_parent_num_id

    inclusion_heading = None
    exclusion_heading = None
    lifestyle_heading = None
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        if inclusion_heading is None and is_heading(para, text, "入选标准"):
            inclusion_heading = i
            continue
        if inclusion_heading is not None and exclusion_heading is None and is_heading(para, text, "排除标准"):
            exclusion_heading = i
            continue
        if exclusion_heading is not None and is_section_heading(para) and i > exclusion_heading:
            lifestyle_heading = i
            break

    if inclusion_heading is None or exclusion_heading is None:
        raise ValueError("未能在正文中定位入选标准/排除标准标题")

    stage = _normalize_study_stage_label(study_stage)
    inclusion_start = inclusion_heading + 1
    inclusion_end = exclusion_heading
    preamble = []
    if stage:
        stage_marker = None
        next_stage_marker = None
        all_stage_markers = []
        for i in range(inclusion_start, inclusion_end):
            para = doc.paragraphs[i]
            text = para.text.strip()
            if not text:
                continue
            if stage_marker is None and not para_num_id(para) and ("签署ICF" in text or "知情同意" in text):
                preamble.append(text)
            if (
                ("Ⅱ期" in text or "Ⅲ期" in text or "II期" in text or "III期" in text)
                and "符合下列所有标准" in text
            ):
                all_stage_markers.append(i)
            if stage in text and "符合下列所有标准" in text:
                stage_marker = i
                continue
            if stage_marker is not None and (
                ("Ⅱ期" in text or "Ⅲ期" in text)
                and stage not in text
                and "符合下列所有标准" in text
            ):
                next_stage_marker = i
                break
        if all_stage_markers and stage_marker is None:
            raise ValueError(f"未能定位{stage}入选标准")
        if stage_marker is not None:
            inclusion_start = stage_marker + 1
            inclusion_end = next_stage_marker or exclusion_heading

    inclusion_items, _ = collect_group(inclusion_start, inclusion_end)
    exclusion_items, _ = collect_group(exclusion_heading + 1, lifestyle_heading or len(doc.paragraphs))

    inclusion_lines = []
    for note in preamble:
        inclusion_lines.append(f"注：{note}")
    for i, item in enumerate(inclusion_items, 1):
        inclusion_lines.append(f"{i}）{item['text']}")
        for sub in item['sub_items']:
            inclusion_lines.append(f"  - {sub}")
    
    exclusion_lines = []
    for i, item in enumerate(exclusion_items, 1):
        exclusion_lines.append(f"{i}）{item['text']}")
        for sub in item['sub_items']:
            exclusion_lines.append(f"  - {sub}")
    
    return {
        'inclusion_text': '\n'.join(inclusion_lines),
        'exclusion_text': '\n'.join(exclusion_lines),
        'inclusion_count': len(inclusion_items),
        'exclusion_count': len(exclusion_items),
        'study_stage': stage,
    }


def _clean_cell_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def _clean_schedule_label(text: str) -> str:
    text = _clean_cell_text(text)
    return re.sub(r"(?<=[\u4e00-\u9fffA-Za-z）)])\d+$", "", text).strip()


def _phase_id_from_visit(stage: str, visit: str, day_window: str) -> str:
    joined = f"{stage} {visit} {day_window}"
    visit_window = f"{visit} {day_window}"
    if "V2" in visit_window or "基线" in visit_window or "D1" in visit_window or "随机" in visit_window:
        return "baseline_randomization"
    if "V1" in joined or "筛选" in joined or "导入" in joined:
        return "screening_run_in"
    slug = re.sub(r"[^A-Za-z0-9]+", "_", joined).strip("_").lower()
    return slug or "review_phase"


def _phase_name_from_visit(stage: str, visit: str, day_window: str) -> str:
    joined = f"{stage} {visit} {day_window}"
    visit_window = f"{visit} {day_window}"
    if "V2" in visit_window or "基线" in visit_window or "D1" in visit_window or "随机" in visit_window:
        return "基线/随机前"
    if "V1" in joined or "筛选" in joined or "导入" in joined:
        return "筛选/导入期"
    return visit or stage or "阶段审核"


def _included_document_phases(phase_id: str) -> list[str]:
    if phase_id == "screening_run_in":
        return ["prior", "screening", "communication", "unknown"]
    if phase_id == "baseline_randomization":
        return ["prior", "screening", "baseline", "communication", "unknown"]
    return ["prior", "screening", "baseline", "postbaseline", "communication", "unknown"]


def _merge_review_phase(phases: list[dict], phase: dict) -> None:
    for existing in phases:
        if (
            existing.get("phase_id") != phase.get("phase_id")
            or (existing.get("study_stage") or "") != (phase.get("study_stage") or "")
        ):
            continue
        for field in ("stage", "visit", "study_week", "day_window"):
            old_value = existing.get(field) or ""
            new_value = phase.get(field) or ""
            if new_value and new_value not in old_value.split("；"):
                existing[field] = f"{old_value}；{new_value}" if old_value else new_value
        old_items = existing.setdefault("required_items", [])
        for item in phase.get("required_items") or []:
            if item not in old_items:
                old_items.append(item)
        return
    phases.append(phase)


def _infer_schedule_table_stage(flat_text: str) -> str:
    """Infer a schedule table's study stage only from explicit stage evidence."""
    compact = re.sub(r"\s+", "", flat_text or "")
    has_phase2 = bool(re.search(r"(Ⅱ期|II期|2期|二期|探索试验阶段|探索试验)", compact, flags=re.I))
    has_phase3 = bool(re.search(r"(Ⅲ期|III期|3期|三期|确证试验阶段|确证试验)", compact, flags=re.I))
    if has_phase2 and has_phase3:
        return ""
    if has_phase3 or "基础期" in compact or "扩展期" in compact:
        return "Ⅲ期"
    if has_phase2:
        return "Ⅱ期"
    return ""


def extract_docx_protocol_workflow(doc_path: str) -> dict:
    """Extract staged review workflow from a protocol DOCX schedule table.

    The function looks for the study schedule table and identifies visit columns
    where "审核入选/排除标准" is marked. Those columns become selectable review
    phases. It also captures all schedule rows marked X in the same columns, so
    baseline-or-earlier required operations are available to the review prompt.
    """
    from docx import Document

    doc = Document(doc_path)
    all_text = "\n".join(_clean_cell_text(p.text) for p in doc.paragraphs)
    study_stages = []
    if any(token in all_text for token in ("Ⅱ/Ⅲ", "II期/Ⅲ期")) or ("Ⅱ期" in all_text and "Ⅲ期" in all_text):
        study_stages = ["Ⅱ期", "Ⅲ期"]
    else:
        if "Ⅱ期" in all_text or "II期" in all_text:
            study_stages.append("Ⅱ期")
        if "Ⅲ期" in all_text or "III期" in all_text:
            study_stages.append("Ⅲ期")

    schedule_tables: list[tuple[list[list[str]], str]] = []
    for table in doc.tables:
        rows = [[_clean_cell_text(cell.text) for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        first_col = [row[0] for row in rows if row]
        has_review_row = any(
            ("审核入选/排除标准" in c or "入排标准审核" in c)
            for c in first_col
        )
        has_stage_or_visit_row = any(
            ("试验阶段" in c or c == "项目" or "访视" in c or "试验周数" in c or "试验天数" in c)
            for c in first_col
        )
        if not (has_review_row and has_stage_or_visit_row):
            continue
        flat = " ".join(" ".join(row) for row in rows[:8])
        table_stage = _infer_schedule_table_stage(flat)
        schedule_tables.append((rows, table_stage))

    if (
        len(schedule_tables) > 1
        and "Ⅱ期" in study_stages
        and "Ⅲ期" in study_stages
        and any(stage == "Ⅲ期" for _, stage in schedule_tables)
        and not any(stage == "Ⅱ期" for _, stage in schedule_tables)
    ):
        schedule_tables = [
            (rows, "Ⅱ期" if not stage else stage)
            for rows, stage in schedule_tables
        ]

    review_phases = []
    for schedule_table, table_stage in schedule_tables:
        stage_row = next(
            (
                row for row in schedule_table
                if row and ("试验阶段" in row[0] or (row[0] == "项目" and any("期" in cell for cell in row[1:])))
            ),
            [],
        )
        visit_row = next((row for row in schedule_table if row and row[0] == "访视"), [])
        if not visit_row:
            visit_row = next(
                (
                    row for row in schedule_table
                    if row
                    and (
                        row[0] == "项目"
                        or "访视" in row[0]
                    )
                    and any((cell.strip() == "筛选" or "基线" in cell or "随机" in cell) for cell in row[1:])
                ),
                [],
            )
        week_row = next(
            (
                row for row in schedule_table
                if row and ("试验周数" in row[0] or ("访视(周)" in row[0] and any(re.search(r"\bW-?\d|\bW\d", cell) for cell in row[1:])))
            ),
            [],
        )
        day_row = next(
            (
                row for row in schedule_table
                if row and ("试验天数" in row[0] or "访视(天)" in row[0])
            ),
            [],
        )
        review_row = next(
            (
                row for row in schedule_table
                if row and ("审核入选/排除标准" in row[0] or "入排标准审核" in row[0])
            ),
            [],
        )

        for col_idx in range(1, len(review_row)):
            if "X" not in review_row[col_idx]:
                continue
            stage = stage_row[col_idx] if col_idx < len(stage_row) else ""
            if table_stage and table_stage not in stage:
                stage = f"{table_stage}{stage}".strip()
            visit = visit_row[col_idx] if col_idx < len(visit_row) else ""
            week = week_row[col_idx] if col_idx < len(week_row) else ""
            day_window = day_row[col_idx] if col_idx < len(day_row) else ""
            phase_id = _phase_id_from_visit(stage, visit, day_window)
            required_items = []
            for row in schedule_table:
                if not row or col_idx >= len(row):
                    continue
                label = _clean_schedule_label(row[0])
                mark = row[col_idx]
                if not label or label == mark:
                    continue
                if "X" in mark and label not in required_items:
                    required_items.append(label)
            _merge_review_phase(review_phases, {
                "phase_id": phase_id,
                "name": _phase_name_from_visit(stage, visit, day_window),
                "stage": stage.replace("*", ""),
                "study_stage": table_stage,
                "visit": visit,
                "study_week": week,
                "day_window": day_window,
                "description": f"{visit or stage}（{day_window or week}）入排审核时点",
                "required_items": required_items,
                "included_document_phases": _included_document_phases(phase_id),
            })

    return {
        "source": "protocol_schedule_table",
        "study_stages": study_stages,
        "requires_study_stage_selection": len(study_stages) > 1,
        "review_phases": review_phases,
    }


def render_protocol_workflow_markdown(workflow: dict, study_stage: str = "") -> str:
    """Render extracted workflow metadata as a compact rules appendix."""
    phases = workflow.get("review_phases") or []
    selected_stage = _normalize_study_stage_label(study_stage)
    if selected_stage:
        stage_specific = [
            phase for phase in phases
            if phase.get("study_stage") == selected_stage
        ]
        if stage_specific:
            phases = stage_specific
    if not phases:
        return ""
    stage_text = "、".join(workflow.get("study_stages") or []) or "未识别"
    lines = [
        "## 三、基线及以前方案流程核查（Protocol Workflow Checks）",
        "",
        f"- **方案包含研究分期**：{stage_text}",
        f"- **本次解构期别**：{selected_stage or '未指定'}",
        f"- **多研究阶段确认**：{'需要用户选择' if workflow.get('requires_study_stage_selection') else '不需要'}",
        "- **阶段性审核原则**：按研究流程表中“审核入选/排除标准”标记的访视逐阶段审核；未到达的后续阶段要求不得提前判为证据不足。",
        "",
        "| 阶段ID | 阶段 | 访视 | 时间窗 | 本阶段应完成/核查项目 |",
        "|---|---|---|---|---|",
    ]
    for phase in phases:
        items = "；".join(phase.get("required_items") or [])
        lines.append(
            f"| {phase.get('phase_id', '')} | {phase.get('name', '')} | "
            f"{phase.get('visit', '')} | {phase.get('day_window', '')} | {items} |"
        )
    return "\n".join(lines)
