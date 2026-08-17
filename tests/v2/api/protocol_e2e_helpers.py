"""E2E helpers for protocol deconstruction API tests."""
from __future__ import annotations

import uuid
from pathlib import Path

from docx import Document

from app.domain.contracts.agent_io import (
    ProtocolSemanticDeconstructionCandidate,
    SemanticEvidenceRequirement,
    SemanticRule,
    SemanticRuleComponent,
)
from app.domain.contracts.enums import CatalogItemKind, Comparator, ReviewStage
from app.domain.contracts.rules import AtomicExpression, AtomicPredicate
from app.protocols.deconstruction_service import ProtocolDeconstructionInputPackage
from app.protocols.docx_structure import StructureBlock


def build_pipeline_e2e_docx(path: Path) -> None:
    """Synthetic DOCX with metadata, II 期 markers, IN/EX lists, and visit table."""
    doc = Document()
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    header = section.header
    header.paragraphs[0].text = "方案编号：E2E-001"
    if len(header.paragraphs) == 1:
        header.add_paragraph("版本号：V1.0")
        header.add_paragraph("日期：2026-08-17")

    doc.add_paragraph("E2E 临床研究方案")
    doc.add_paragraph("II期临床试验")
    doc.add_paragraph("入选标准", style="Heading 2")
    doc.add_paragraph("年龄≥18岁", style="List Number")
    doc.add_paragraph("排除标准", style="Heading 2")
    doc.add_paragraph("活动性感染", style="List Number")
    doc.add_paragraph("II期研究流程表", style="Heading 2")

    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "检查项目"
    table.cell(0, 1).text = "筛选期"
    table.cell(0, 2).text = "基线期"
    table.cell(1, 0).text = "访视"
    table.cell(1, 1).text = "V1"
    table.cell(1, 2).text = "V2"
    table.cell(2, 0).text = "血生化检查"
    table.cell(2, 1).text = "X"
    table.cell(2, 2).text = "X"
    doc.save(str(path))


def page_texts_from_blocks(blocks: tuple[StructureBlock, ...]) -> list[str]:
    parts = [block.text for block in blocks if block.text]
    return ["\n".join(parts)]


def build_passing_draft_json(package: ProtocolDeconstructionInputPackage) -> str:
    """Build semantic candidate JSON aligned to the assembled input package."""
    source_input = package.source_input
    materials = {
        material.source_span_id: material.text
        for material in source_input.source_materials
    }
    proposed_rules: list[SemanticRule] = []
    for item in source_input.parent_rule_catalog.items:
        if item.kind != CatalogItemKind.PARENT_RULE or not item.official_code:
            continue
        span_id = item.source_span_ids[0]
        excerpt = materials.get(span_id, item.label)
        proposed_rules.append(
            SemanticRule(
                official_code=item.official_code,
                components=[
                    SemanticRuleComponent(
                        title=item.label,
                        expression=AtomicExpression(
                            predicate=AtomicPredicate(
                                predicate_id=f"predicate-{item.official_code.lower()}",
                                subject="受试者",
                                attribute="方案条件",
                                source_term=item.label,
                                source_clause=excerpt,
                                comparator=Comparator.EQ,
                                value=True,
                                unit=None,
                            )
                        ),
                        evidence_requirements=[
                            SemanticEvidenceRequirement(
                                fact_type="方案要求事实",
                                required_source_types=["clinical_record"],
                                due_stage=ReviewStage.SCREENING,
                                description="核对正式原始资料和研究者记录",
                            )
                        ],
                        source_span_ids=[span_id],
                        source_excerpts=[excerpt],
                    )
                ],
            )
        )

    candidate = ProtocolSemanticDeconstructionCandidate(
        candidate_id=uuid.uuid4().hex,
        proposed_rules=proposed_rules,
        created_by_agent_call_id="protocol-e2e-test",
    )
    return candidate.model_dump_json()
