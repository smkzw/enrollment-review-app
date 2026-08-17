"""V2 方案解构 API 集成测试。"""
from __future__ import annotations

import io
import tempfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.api.v2.protocols import _metadata_candidate_dto
from app.domain.contracts.enums import StudyPhase
from app.services.protocol_deconstruction_executor import (
    ProtocolDeconstructionExecutorConfig,
    create_protocol_deconstruction_executor,
)
from app.services.protocol_workbench_service import (
    PROTOCOL_DECONSTRUCTION_JOB_TYPE,
    PROTOCOL_DECONSTRUCTION_STEPS,
    ProtocolWorkbenchService,
)
from app.workflow.runner import JobRunner
from app.workflow.jobstore import JobStore
from tests.v2.api.protocol_e2e_helpers import (
    build_passing_draft_json,
    build_pipeline_e2e_docx,
    page_texts_from_blocks,
)
from tests.v2.protocols.slice4_helpers import confirmed_fixture


def _minimal_docx_bytes() -> bytes:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as handle:
        path = Path(handle.name)
    doc = Document()
    doc.add_paragraph("临床研究方案")
    doc.add_paragraph("方案编号：TEST-001")
    doc.save(str(path))
    payload = path.read_bytes()
    path.unlink(missing_ok=True)
    return payload


def _pipeline_docx_bytes() -> bytes:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as handle:
        path = Path(handle.name)
    build_pipeline_e2e_docx(path)
    payload = path.read_bytes()
    path.unlink(missing_ok=True)
    return payload


def _create_protocol_job(
    client: TestClient,
    *,
    key: str = "proto-key-1",
    docx_bytes: bytes | None = None,
) -> str:
    payload = docx_bytes if docx_bytes is not None else _minimal_docx_bytes()
    files = {
        "file": (
            "test-protocol.docx",
            io.BytesIO(payload),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    data = {"idempotency_key": key, "actor": "测试用户"}
    response = client.post("/api/v2/protocol/deconstructions", files=files, data=data)
    assert response.status_code == 201, response.text
    return response.json()["job_id"]


def _test_executor(app) -> JobRunner:
    executor = create_protocol_deconstruction_executor(
        ProtocolDeconstructionExecutorConfig(
            data_paths=app.state.data_paths,
            session_factory=app.state.session_factory,
            page_texts_builder=page_texts_from_blocks,
            draft_response_builder=build_passing_draft_json,
        )
    )
    return JobRunner(
        app.state.session_factory,
        {PROTOCOL_DECONSTRUCTION_JOB_TYPE: executor},
        worker_id="test-runner",
        poll_interval=0.01,
    )


def _run_until(client: TestClient, app, job_id: str, *, awaiting_user: str, limit: int = 40) -> None:
    runner = _test_executor(app)
    for _ in range(limit):
        session = client.get(f"/api/v2/protocol/deconstructions/{job_id}").json()
        if session.get("awaiting_user") == awaiting_user:
            return
        runner.run_job(job_id)
    session = client.get(f"/api/v2/protocol/deconstructions/{job_id}").json()
    pytest.fail(
        f"任务未在 {limit} 次推进后到达 awaiting_user={awaiting_user!r}；"
        f"当前状态={session.get('state')!r} awaiting_user={session.get('awaiting_user')!r}"
    )


def _seed_review_job(app, job_id: str, *, wait_at: str | None = None) -> tuple:
    source_input, draft, spans = confirmed_fixture()
    service: ProtocolWorkbenchService = app.state.protocol_workbench_service
    kwargs: dict = {
        "source_input": source_input,
        "draft": draft,
        "source_spans": spans,
    }
    if wait_at is not None:
        kwargs["wait_at"] = wait_at
    service.seed_review_session(job_id, **kwargs)
    return source_input, draft


def test_upload_registers_source_and_creates_protocol_job(client) -> None:
    job_id = _create_protocol_job(client, key="upload-1")
    session = client.get(f"/api/v2/protocol/deconstructions/{job_id}")
    assert session.status_code == 200
    body = session.json()
    assert body["job_type"] == PROTOCOL_DECONSTRUCTION_JOB_TYPE
    assert body["file_name"] == "test-protocol.docx"
    assert body["source_artifact_id"]
    assert body["progress_total"] == len(PROTOCOL_DECONSTRUCTION_STEPS)
    assert body["progress_completed"] >= 1


def test_upload_registration_survives_api_claim_race(client, monkeypatch) -> None:
    original_claim_job = JobStore.claim_job
    api_claims = 0

    def claim_job_once_for_runner_race(self, job_id: str, worker_id: str):
        nonlocal api_claims
        if worker_id == ProtocolWorkbenchService.WORKER_ID and api_claims == 0:
            api_claims += 1
            return None
        return original_claim_job(self, job_id, worker_id)

    monkeypatch.setattr(JobStore, "claim_job", claim_job_once_for_runner_race)

    job_id = _create_protocol_job(client, key="upload-claim-race")
    session = client.get(f"/api/v2/protocol/deconstructions/{job_id}")
    assert session.status_code == 200
    assert session.json()["source_artifact_id"]
    assert session.json()["progress_completed"] == 0

    runner = _test_executor(client.app)
    runner.run_job(job_id)
    status = client.get(f"/api/v2/jobs/{job_id}")
    assert status.status_code == 200
    assert status.json()["error_code"] != "REGISTER_INCOMPLETE"


def test_upload_idempotent_same_key(client) -> None:
    files = {
        "file": (
            "test-protocol.docx",
            io.BytesIO(_minimal_docx_bytes()),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    data = {"idempotency_key": "upload-dup", "actor": "测试用户"}
    first = client.post("/api/v2/protocol/deconstructions", files=files, data=data)
    second = client.post("/api/v2/protocol/deconstructions", files=files, data=data)
    assert first.status_code == 201
    assert second.status_code == 200
    assert second.json()["job_id"] == first.json()["job_id"]
    assert second.json()["created"] is False


def test_missing_job_returns_chinese_envelope(client) -> None:
    response = client.get("/api/v2/protocol/deconstructions/does-not-exist")
    assert response.status_code == 404
    error = response.json()["error"]
    assert error["code"] == "NOT_FOUND"
    assert error["title"] == "任务不存在"
    assert "correlation_id" in error


def test_draft_integrity_and_sources_after_seed(client, build_app) -> None:
    app = build_app()
    with TestClient(app) as test_client:
        job_id = _create_protocol_job(test_client, key="seed-review-1")
        _seed_review_job(app, job_id)

        draft = test_client.get(f"/api/v2/protocol/deconstructions/{job_id}/draft")
        assert draft.status_code == 200
        draft_body = draft.json()
        assert draft_body["revision_number"] == 1
        assert draft_body["rule_count"] == 2
        assert draft_body["status_label"] == "已保存"
        assert "content" in draft_body

        integrity = test_client.get(f"/api/v2/protocol/deconstructions/{job_id}/integrity")
        assert integrity.status_code == 200
        integrity_body = integrity.json()
        assert integrity_body["publishable"] is True
        assert integrity_body["blocking_count"] == 0
        assert integrity_body["summary"]

        sources = test_client.get(f"/api/v2/protocol/deconstructions/{job_id}/sources")
        assert sources.status_code == 200
        sources_body = sources.json()
        assert sources_body["selected_phase"] == "phase_ii"
        assert sources_body["selected_phase_label"] == "II 期"
        assert sources_body["source_materials"]

        session = test_client.get(f"/api/v2/protocol/deconstructions/{job_id}")
        assert session.status_code == 200
        assert session.json()["draft_revision_number"] == 1
        assert session.json()["awaiting_user"] == "review"


def test_save_draft_is_idempotent(client, build_app) -> None:
    app = build_app()
    with TestClient(app) as test_client:
        job_id = _create_protocol_job(test_client, key="save-draft-1")
        _seed_review_job(app, job_id)
        draft = test_client.get(f"/api/v2/protocol/deconstructions/{job_id}/draft").json()
        response = test_client.post(
            f"/api/v2/protocol/deconstructions/{job_id}/draft/save",
            json={"expected_revision_id": draft["revision_id"]},
        )
        assert response.status_code == 200
        assert response.json()["status_label"] == "已保存"


def test_publish_first_project(client, build_app) -> None:
    app = build_app()
    with TestClient(app) as test_client:
        job_id = _create_protocol_job(test_client, key="publish-1")
        _seed_review_job(app, job_id, wait_at="publish")
        draft = test_client.get(f"/api/v2/protocol/deconstructions/{job_id}/draft").json()
        response = test_client.post(
            f"/api/v2/protocol/deconstructions/{job_id}/publish",
            json={"idempotency_key": "pub-key-1", "actor": "医学监查员"},
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["project_id"]
        assert body["rule_set_revision"] == 1
        assert body["replay"] is False

        replay = test_client.post(
            f"/api/v2/protocol/deconstructions/{job_id}/publish",
            json={"idempotency_key": "pub-key-1", "actor": "医学监查员"},
        )
        assert replay.status_code == 200
        assert replay.json()["replay"] is True


def test_identity_not_ready_before_pipeline(client) -> None:
    job_id = _create_protocol_job(client, key="identity-wait-1")
    response = client.get(f"/api/v2/protocol/deconstructions/{job_id}/identity")
    assert response.status_code == 409
    error = response.json()["error"]
    assert error["code"] == "IDENTITY_NOT_READY"
    assert error["recovery_action"]


def test_draft_not_ready_before_review(client) -> None:
    job_id = _create_protocol_job(client, key="draft-wait-1")
    response = client.get(f"/api/v2/protocol/deconstructions/{job_id}/draft")
    assert response.status_code == 409
    error = response.json()["error"]
    assert error["code"] in {"DRAFT_NOT_READY", "STATE_CONFLICT", "STEP_STATE_CONFLICT"}
    assert error["title"]
    assert error["recovery_action"]


@pytest.mark.parametrize(
    ("official_date_value", "official_date_precision"),
    [("2026", "year"), ("2026-08", "month"), ("2026-08-17", "day")],
)
def test_upload_pipeline_reaches_identity_and_review_without_seed(
    build_app,
    official_date_value: str,
    official_date_precision: str,
) -> None:
    app = build_app(run_runner=False)
    with TestClient(app) as client:
        job_id = _create_protocol_job(
            client,
            key=f"pipeline-e2e-{official_date_precision}",
            docx_bytes=_pipeline_docx_bytes(),
        )
        _run_until(client, app, job_id, awaiting_user="identity")

        identity_session = client.get(f"/api/v2/jobs/{job_id}").json()
        identity_wait = next(
            step
            for step in identity_session["steps"]
            if step["step_id"] == "await_identity_confirm"
        )
        assert identity_wait["state"] == "waiting_user"
        assert identity_wait["attempt"] == 0
        identity_events = [item["event_type"] for item in identity_session["events"]]
        assert "waiting_user" in identity_events
        assert "step_failed" not in identity_events
        assert "retry_scheduled" not in identity_events

        identity = client.get(f"/api/v2/protocol/deconstructions/{job_id}/identity")
        assert identity.status_code == 200, identity.text
        identity_body = identity.json()
        assert identity_body["snapshot_id"]
        assert identity_body["phase_candidates"]
        assert len(identity_body["phase_candidates"]) == 1
        phase_candidate = identity_body["phase_candidates"][0]
        assert phase_candidate["phase_label"] == "II 期"
        assert "方案原文" in phase_candidate["rationale"]
        assert phase_candidate["source_excerpt"]
        assert identity_body["identity"]["status_label"] == "需要确认"
        assert all(
            item["field_label"] and item["source_label"]
            for item in identity_body["metadata_candidates"]
        )

        confirm = client.post(
            f"/api/v2/protocol/deconstructions/{job_id}/identity/confirm",
            json={
                "protocol_code": "E2E-001",
                "project_name": "E2E 测试研究",
                "official_version": "V1.0",
                "official_date_value": official_date_value,
                "official_date_precision": official_date_precision,
                "study_phase": StudyPhase.PHASE_II.value,
                "actor": "测试用户",
            },
        )
        assert confirm.status_code == 200, confirm.text

        confirmed_identity = client.get(
            f"/api/v2/protocol/deconstructions/{job_id}/identity"
        )
        assert confirmed_identity.status_code == 200, confirmed_identity.text
        assert confirmed_identity.json()["identity"]["official_date_value"] == official_date_value
        assert (
            confirmed_identity.json()["identity"]["official_date_precision"]
            == official_date_precision
        )

        _run_until(client, app, job_id, awaiting_user="review", limit=60)

        session = client.get(f"/api/v2/protocol/deconstructions/{job_id}").json()
        assert session["awaiting_user"] == "review"
        assert session["draft_revision_number"] == 1
        review_job = client.get(f"/api/v2/jobs/{job_id}").json()
        review_wait = next(
            step for step in review_job["steps"] if step["step_id"] == "await_review"
        )
        assert review_wait["state"] == "waiting_user"
        assert review_wait["attempt"] == 0

        draft = client.get(f"/api/v2/protocol/deconstructions/{job_id}/draft")
        assert draft.status_code == 200, draft.text
        draft_body = draft.json()
        assert draft_body["revision_number"] == 1
        assert draft_body["rule_count"] >= 1
        assert "content" in draft_body

        sources = client.get(f"/api/v2/protocol/deconstructions/{job_id}/sources")
        assert sources.status_code == 200, sources.text
        assert sources.json()["selected_phase_label"] == "II 期"


def test_session_projects_persisted_publish_wait(build_app) -> None:
    app = build_app()
    with TestClient(app) as client:
        job_id = _create_protocol_job(client, key="publish-wait-projection")
        _seed_review_job(app, job_id, wait_at="publish")

        session = client.get(f"/api/v2/protocol/deconstructions/{job_id}")
        assert session.status_code == 200, session.text
        body = session.json()
        assert body["state"] == "waiting_user"
        assert body["awaiting_user"] == "publish"
        assert body["awaiting_user_label"] == "等待发布确认"
        assert "发布正式项目" in body["next_action"]


def test_identity_date_candidate_projection_is_form_ready() -> None:
    candidate = _metadata_candidate_dto(
        {
            "candidate_id": "date-candidate",
            "field_category": "protocol_date",
            "candidate_value": "2026年8月17日",
            "normalized_value": "2026-08-17",
            "source_kind": "first_page",
            "excerpt": "版本日期：2026年8月17日",
        }
    )
    assert candidate.value == "2026-08-17"
    assert candidate.source_excerpt == "版本日期：2026年8月17日"
