"""DOCX 结构提取的结构不变量测试：嵌套表格路径、gridSpan、页眉页脚、编号、修订。"""
from __future__ import annotations

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.domain.contracts.enums import ExtractionStatus
from app.protocols.docx_structure import (
    HeaderFooterKind,
    _effective_numbering,
    extract_docx_structure,
    parse_numbering,
)
from app.protocols.ingestion import register_source_artifact
from .helpers import (
    build_altchunk_body_docx,
    build_gridspan_docx,
    build_header_docx,
    build_nested_table_docx,
    build_tracked_change_header_docx,
    build_sdt_body_docx,
)


def _extract(tmp_path, name, builder):
    path = tmp_path / f"{name}.docx"
    builder(path)
    artifact = register_source_artifact(
        path, source_artifact_id=name, storage_root=tmp_path
    )
    return extract_docx_structure(
        path, snapshot_id=f"{name}-snap", source_artifact=artifact, output_dir=tmp_path
    )


def _blocks_by_ref(blocks, ref):
    return [b for b in blocks if b.source_ref == ref]


def test_nested_table_full_ancestor_path(tmp_path):
    ext = _extract(tmp_path, "nested", build_nested_table_docx)
    # 嵌套表块自身定位到外层单元格 (0,0)
    nested_tables = [
        b for b in ext.blocks if b.kind.value == "table" and b.table_path is not None
    ]
    assert len(nested_tables) == 1
    assert nested_tables[0].source_ref == "body.t0.r0.c0.t0"
    assert nested_tables[0].table_path == (0, 0)
    # 嵌套单元格段落保留完整祖先链 (0,0,0,0)
    inner_para = _blocks_by_ref(ext.blocks, "body.t0.r0.c0.t0.r0.c0.p0")
    assert inner_para, "嵌套表格内段落未被捕获"
    assert inner_para[0].table_path == (0, 0, 0, 0)
    assert inner_para[0].text == "inner00"


def test_gridspan_table_size(tmp_path):
    ext = _extract(tmp_path, "gridspan", build_gridspan_docx)
    table = _blocks_by_ref(ext.blocks, "body.t0")
    assert table and table[0].kind.value == "table"
    # 合并单元格右边缘为 col + span = 3，而非 col + 1
    assert table[0].table_cols == 3
    assert table[0].table_rows == 1


def test_superscript_and_subscript_are_preserved_as_semantic_text(tmp_path):
    def build(path):
        document = Document()
        paragraph = document.add_paragraph("ANC<1.2×10")
        superscript = paragraph.add_run("9")
        superscript.font.superscript = True
        paragraph.add_run("/L，H")
        subscript = paragraph.add_run("2")
        subscript.font.subscript = True
        paragraph.add_run("O")
        footnote = document.add_paragraph("胸片（正侧位）")
        for digit in "14":
            run = footnote.add_run(digit)
            run.font.superscript = True
        document.save(path)

    ext = _extract(tmp_path, "vertical-align", build)

    assert any(block.text == "ANC<1.2×10^9/L，H_2O" for block in ext.blocks)
    assert any(block.text == "胸片（正侧位）^14" for block in ext.blocks)


def test_headers_capture_default_first_even_and_dedup(tmp_path):
    ext = _extract(tmp_path, "hdr", build_header_docx)
    headers = [b for b in ext.blocks if b.document_part.value == "header"]
    by_kind: dict[str, list] = {}
    for b in headers:
        by_kind.setdefault(b.part_kind.value, []).append(b)

    assert set(by_kind) == {"default", "first", "even"}, by_kind.keys()
    assert "FIRST_H" in "".join(b.text for b in by_kind["first"])
    assert "DEFAULT_H" in "".join(b.text for b in by_kind["default"])
    assert "EVEN_H" in "".join(b.text for b in by_kind["even"])

    # 默认页眉被两个 section 共享，去重后仅一份，section_indexes 保留全部适用 section
    default_blocks = by_kind["default"]
    default_text = "".join(b.text for b in default_blocks)
    assert default_text.count("DEFAULT_H") == 1
    assert default_blocks[0].section_indexes == (0, 1)
    assert default_blocks[0].part_kind == HeaderFooterKind.DEFAULT

    # 第一页页脚同样被捕获
    footers = [b for b in ext.blocks if b.document_part.value == "footer"]
    footer_text = "".join(b.text for b in footers)
    assert "FIRST_F" in footer_text
    assert "DEFAULT_F" in footer_text


def test_numbering_definition_preserved_with_override():
    numbering_xml = (
        f'<w:numbering {nsdecls("w")}>'
        '<w:abstractNum w:abstractNumId="5">'
        '<w:lvl w:ilvl="0"><w:numFmt w:val="decimal"/>'
        '<w:lvlText w:val="%1、"/><w:start w:val="1"/></w:lvl>'
        "</w:abstractNum>"
        '<w:num w:numId="7"><w:abstractNumId w:val="5"/>'
        '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="4"/></w:lvlOverride>'
        "</w:num>"
        "</w:numbering>"
    )
    abstracts, nums = parse_numbering(parse_xml(numbering_xml))
    ref = _effective_numbering(7, 0, abstracts, nums)
    assert ref.abstract_num_id == 5
    assert ref.num_fmt == "decimal"
    assert ref.lvl_text == "%1、"
    assert ref.start == 4  # startOverride 生效
    assert ref.has_level_override is True


def test_numbering_definition_without_override():
    numbering_xml = (
        f'<w:numbering {nsdecls("w")}>'
        '<w:abstractNum w:abstractNumId="3">'
        '<w:lvl w:ilvl="1"><w:numFmt w:val="lowerLetter"/>'
        '<w:lvlText w:val="%2)"/><w:start w:val="1"/></w:lvl>'
        "</w:abstractNum>"
        '<w:num w:numId="9"><w:abstractNumId w:val="3"/></w:num>'
        "</w:numbering>"
    )
    abstracts, nums = parse_numbering(parse_xml(numbering_xml))
    ref = _effective_numbering(9, 1, abstracts, nums)
    assert ref.abstract_num_id == 3
    assert ref.num_fmt == "lowerLetter"
    assert ref.lvl_text == "%2)"
    assert ref.start == 1
    assert ref.has_level_override is False


def test_tracked_change_in_header_marks_needs_review(tmp_path):
    ext = _extract(tmp_path, "tc", build_tracked_change_header_docx)
    assert ext.snapshot.coverage.tracked_change_count >= 1
    assert ext.snapshot.status == ExtractionStatus.NEEDS_REVIEW
    assert any(a.kind == "unaccepted_tracked_changes" for a in ext.snapshot.anomalies)
    # 页眉中的插入文本被计入块集原文
    header_text = "".join(
        b.text for b in ext.blocks if b.document_part.value == "header"
    )
    assert "inserted" in header_text


def test_content_control_body_is_unwrapped_in_document_order(tmp_path):
    ext = _extract(tmp_path, "sdt", build_sdt_body_docx)
    assert any(
        block.text == "内容控件中的基线前必做检查" for block in ext.blocks
    )


def test_unhandled_altchunk_is_explicit_anomaly(tmp_path):
    ext = _extract(tmp_path, "altchunk", build_altchunk_body_docx)
    assert ext.snapshot.status == ExtractionStatus.NEEDS_REVIEW
    assert any(
        anomaly.kind == "unhandled_ooxml_container"
        and anomaly.scope_ref == "altChunk"
        for anomaly in ext.snapshot.anomalies
    )
