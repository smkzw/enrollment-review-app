"""受控渲染护栏：陈旧目录污染不可选错文件、成功渲染至少 1 页、失败保留结构化原因。"""
from __future__ import annotations

import hashlib

import pytest
from docx import Document

from app.domain.contracts.enums import RenderStatus
from app.protocols.ingestion import register_source_artifact
from app.protocols.rendering import locate_libreoffice, render_to_pdf

try:
    locate_libreoffice()
    _HAS_LO = True
except Exception:
    _HAS_LO = False

pytestmark = pytest.mark.skipif(not _HAS_LO, reason="LibreOffice 不可用")


def _docx(tmp_path, name="sample"):
    path = tmp_path / f"{name}.docx"
    doc = Document()
    doc.add_paragraph("入排标准示例正文")
    doc.save(str(path))
    return path


def test_render_never_selects_unrelated_stale_pdf(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    # 预置无关/陈旧 PDF，模拟污染目录
    stale = out / "unrelated.pdf"
    stale.write_bytes(b"%PDF-1.4 stale content that is not a real render")
    stale_before = stale.read_bytes()

    source = _docx(tmp_path, "protocol-a")
    artifact = register_source_artifact(
        source, source_artifact_id="protocol-a", storage_root=tmp_path
    )
    result = render_to_pdf(source, out, source_artifact=artifact)

    assert result.status in (RenderStatus.SUCCEEDED, RenderStatus.DEGRADED)
    assert result.pdf_path is not None
    assert result.pdf_path.name == f"{result.pdf_sha256}.pdf"
    # 无关陈旧 PDF 保持原样，绝不作为本次产物被选中/覆盖
    assert stale.read_bytes() == stale_before
    # 成功渲染至少 1 页
    assert result.page_count is not None and result.page_count >= 1


def test_render_keeps_stale_same_name_pdf_and_writes_content_addressed_artifact(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    source = _docx(tmp_path, "protocol-b")
    artifact = register_source_artifact(
        source, source_artifact_id="protocol-b", storage_root=tmp_path
    )
    stale_same = out / f"{source.stem}.pdf"
    stale_same.write_bytes(b"%PDF-1.4 bogus")
    stale_before = stale_same.read_bytes()

    result = render_to_pdf(source, out, source_artifact=artifact)

    assert result.status in (RenderStatus.SUCCEEDED, RenderStatus.DEGRADED)
    assert result.pdf_path != stale_same
    assert stale_same.read_bytes() == stale_before, "内容寻址渲染不应覆盖旧同名文件"
    assert result.page_count is not None and result.page_count >= 1
    assert result.pdf_path is not None
    assert hashlib.sha256(result.pdf_path.read_bytes()).hexdigest() == result.pdf_sha256


def test_render_refuses_file_changed_after_registration(tmp_path):
    source = _docx(tmp_path, "protocol-changed")
    artifact = register_source_artifact(
        source, source_artifact_id="protocol-changed", storage_root=tmp_path
    )
    Document().save(str(source))

    with pytest.raises(Exception, match="哈希不一致"):
        render_to_pdf(source, tmp_path / "out", source_artifact=artifact)
