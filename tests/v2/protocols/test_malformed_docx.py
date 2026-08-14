"""聚焦畸形/空 DOCX 用例：格式误判、损坏部件与空正文都必须显式失败，不伪装成功。"""
from __future__ import annotations

import zipfile

import pytest

from app.domain.contracts.enums import ExtractionStatus
from app.protocols.docx_structure import (
    StructureExtractionError,
    extract_docx_structure,
)
from app.protocols.ingestion import register_source_artifact
from .helpers import build_corrupt_docx, build_empty_body_docx


def _artifact(path, name):
    return register_source_artifact(
        path, source_artifact_id=name, storage_root=path.parent / "store"
    )


def test_empty_file_rejected(tmp_path):
    path = tmp_path / "empty.docx"
    path.write_bytes(b"")
    artifact = _artifact(path, "empty")  # 登记层只记录 MIME（text/plain），不拒绝
    with pytest.raises(StructureExtractionError):
        extract_docx_structure(
            path, snapshot_id="empty-snap", source_artifact=artifact, output_dir=tmp_path
        )


def test_zip_without_document_part_rejected(tmp_path):
    path = tmp_path / "notdocx.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("readme.txt", "just a zip")
    artifact = _artifact(path, "zip")  # 无 word/document.xml -> UNKNOWN
    with pytest.raises(StructureExtractionError):
        extract_docx_structure(
            path, snapshot_id="zip-snap", source_artifact=artifact, output_dir=tmp_path
        )


def test_corrupt_document_xml_rejected(tmp_path):
    path = tmp_path / "corrupt.docx"
    build_corrupt_docx(path)
    artifact = _artifact(path, "corrupt")
    with pytest.raises(StructureExtractionError):
        extract_docx_structure(
            path, snapshot_id="corrupt-snap", source_artifact=artifact, output_dir=tmp_path
        )


def test_empty_body_is_failed_not_success(tmp_path):
    path = tmp_path / "empty-body.docx"
    build_empty_body_docx(path)
    artifact = _artifact(path, "empty-body")
    ext = extract_docx_structure(
        path, snapshot_id="empty-body-snap", source_artifact=artifact, output_dir=tmp_path
    )
    assert ext.snapshot.status == ExtractionStatus.FAILED
    assert any(a.kind == "empty_document" for a in ext.snapshot.anomalies)
    assert ext.snapshot.coverage.paragraph_count == 0
    assert ext.snapshot.coverage.table_count == 0


def test_extraction_refuses_file_changed_after_registration(tmp_path):
    path = tmp_path / "changed.docx"
    from docx import Document

    doc = Document()
    doc.add_paragraph("入选标准")
    doc.save(str(path))
    artifact = _artifact(path, "changed")
    doc = Document()
    doc.add_paragraph("排除标准")
    doc.save(str(path))

    with pytest.raises(StructureExtractionError, match="哈希不一致"):
        extract_docx_structure(
            path,
            snapshot_id="changed-snap",
            source_artifact=artifact,
            output_dir=tmp_path,
        )
