"""链级反过拟合确定性补充：全候选链、提示词、合成夹具与重放包词汇中立。

独立反过拟合审查（worker_03，`phase5-protocol-control-live-model-smoke-20260831`）：

- 静态扫描整个共享候选链（两个传输、解构器、三个服务、无模型重放 harness、
  两个运行器脚本、运行时配置与示例环境），证明不存在 D001/SAR 项目标识、
  药物、疾病、评分量表、项目特异性时间点或阈值的硬编码。
- 行为验证：发现/深析/两类同会话修复提示词的生成结果、冒烟运行器内置合成
  方案 DOCX、重放 harness 落盘工件均不含上述词汇。
- 将异构真实协议只读回放的前置门禁编码为可执行规格：模型身份正向核验
  （不一致即失效关闭）、合成默认输入与真实项目/数据目录守卫、词汇中立夹具、
  相同提示词/合同/门禁的可复现重放、正式目录不物化且 `claims_complete=false`。

本文件不调用任何真实模型，不读取真实方案，不写入任何受控目录。
"""

from __future__ import annotations

import hashlib
import importlib.util
import re
from pathlib import Path

import pytest
from docx import Document

from app.agents.protocol_control_deconstructor import (
    ProtocolControlAgentInput,
    ProtocolControlDiscoveryAgentInput,
    build_protocol_control_agent_prompt,
    build_protocol_control_discovery_prompt,
    build_protocol_control_discovery_repair_prompt,
    build_protocol_control_repair_prompt,
)
from app.domain.contracts.enums import PhaseScope, ReviewStage, StudyPhase
from app.domain.contracts.protocol_controls import (
    ProtocolSectionCoverageManifest,
    ProtocolStructureUnit,
)
from app.domain.contracts.rules import WorkflowStage
from app.protocols.protocol_control_planning import (
    plan_protocol_control_batches,
    plan_protocol_control_discovery,
)
from app.protocols.protocol_replay_harness import (
    REPLAY_HARNESS_CONFIG_SCHEMA,
    ReplayHarnessConfig,
    build_protocol_replay_pack,
    replay_pack_fingerprint,
    verify_replay_pack,
)


ROOT = Path(__file__).resolve().parents[3]
_SMOKE_SCRIPT = ROOT / "scripts" / "run_protocol_control_smoke.py"

# ---------------------------------------------------------------------------
# 词表与模式
# ---------------------------------------------------------------------------

# 与真实语料（D001 银屑病方案、MG-K10-SAR 特应性皮炎方案）绑定的词汇。
# 单词边界的拉丁标记避免误伤 "obligation" 等普通英文单词。
_WORD_BOUNDED_MARKERS = (
    "D001",
    "MG-K10-SAR",
    "CMS-D001",
    "SAR-001",
    "PASI",
    "EASI",
    "IGA",
    "SCORAD",
    "DLQI",
    "PGA",
    "BSA",
    "HADS",
    "NRS",
    "JAK",
    "IL-4",
    "IL-13",
    "IL-17",
    "IL-23",
)
_SUBSTRING_MARKERS = (
    "银屑病",
    "特应性皮炎",
    "斑块",
    "皮炎",
    "度普利尤",
    "达必妥",
    "司库奇尤",
    "乌司奴",
    "第12周",
    "12周",
    "24周",
    "52周",
    "8周",
    "16周",
    "W12",
    "Week 12",
    "week 12",
    "第24周",
    "第52周",
    "75%",
    "50%",
)
# 项目特异性时间点/阈值字面量模式（如 "第8周"、"week 12"、">=75%"）。
# 注意：不含裸 "数字周/天" 模式——领域模型中的通用示例（如 ``4 周``）不是
# 项目特异时间点，由上方显式端点标记（12/24/52/8/16 周）覆盖。
_TIMEPOINT_OR_THRESHOLD_PATTERNS = (
    re.compile(r"第\s*[0-9一二三四五六七八九十百]+\s*(?:周|天)"),
    re.compile(r"\bW\s*[0-9]{1,2}\b", re.IGNORECASE),
    re.compile(r"\bweek\s*[0-9]{1,2}\b", re.IGNORECASE),
    re.compile(r"[0-9]+(?:\.[0-9]+)?\s*%"),
)

# 共享候选链（不含遗留项目脚本与测试自身）。
_CHAIN_FILES = (
    ROOT / "app/agents/protocol_control_agent_transport.py",
    ROOT / "app/agents/protocol_control_discovery_transport.py",
    ROOT / "app/agents/protocol_control_deconstructor.py",
    ROOT / "app/services/protocol_control_execution.py",
    ROOT / "app/services/protocol_control_executor.py",
    ROOT / "app/services/protocol_control_job_service.py",
    ROOT / "app/protocols/protocol_replay_harness.py",
    ROOT / "app/protocols/adaptive_batch_budget.py",
    ROOT / "scripts/run_protocol_control_smoke.py",
    ROOT / "scripts/run_protocol_replay_harness.py",
    ROOT / "app/config.py",
    ROOT / ".env.example",
)


def _marker_hits(text: str) -> tuple[str, ...]:
    lowered = text.casefold()
    hits: list[str] = []
    for marker in _SUBSTRING_MARKERS:
        if marker.casefold() in lowered:
            hits.append(marker)
    for marker in _WORD_BOUNDED_MARKERS:
        if re.search(rf"\b{re.escape(marker)}\b", text, re.IGNORECASE):
            hits.append(marker)
    for pattern in _TIMEPOINT_OR_THRESHOLD_PATTERNS:
        if pattern.search(text):
            hits.append(f"pattern:{pattern.pattern}")
    return tuple(dict.fromkeys(hits))


def _assert_neutral(text: str, *, label: str) -> None:
    hits = _marker_hits(text)
    assert not hits, f"{label} 含有语料/项目特异词汇：{hits}"


# ---------------------------------------------------------------------------
# 中性合成夹具（与真实语料无关的通用语句）
# ---------------------------------------------------------------------------


def _unit(
    unit_id: str,
    order: int,
    excerpt: str,
) -> ProtocolStructureUnit:
    return ProtocolStructureUnit(
        structure_unit_id=unit_id,
        source_ref=f"body.p{order}",
        member_source_refs=[f"body.p{order}"],
        source_span_ids=[f"span:{unit_id}"],
        unit_kind="paragraph",
        heading_path=["通用章节"],
        source_order=order,
        study_phase=StudyPhase.PHASE_II,
        phase_scopes=[PhaseScope.SHARED],
        excerpt=excerpt,
    )


def _manifest(
    *,
    excerpts: tuple[str, ...] = ("必须记录对象α的状态。", "补充语境，不形成独立控制。"),
) -> ProtocolSectionCoverageManifest:
    return ProtocolSectionCoverageManifest(
        manifest_id="manifest:neutral-chain-wide",
        protocol_version_id="protocol:neutral-chain-wide",
        protocol_document_sha256="d" * 64,
        study_phase=StudyPhase.PHASE_II,
        snapshot_id="snapshot:neutral-chain-wide",
        units=[_unit(f"unit-{index:02d}", index, excerpt) for index, excerpt in enumerate(excerpts, start=1)],
        dispositions=[],
        claims_full_coverage=False,
    )


def _workflow() -> list[WorkflowStage]:
    return [
        WorkflowStage(
            workflow_stage_id="stage:neutral:screening",
            stage=ReviewStage.SCREENING,
            display_name="通用筛选节点",
            visit_instance="visit-neutral-1",
        )
    ]


def _load_smoke_module():
    spec = importlib.util.spec_from_file_location(
        "run_protocol_control_smoke_under_test", _SMOKE_SCRIPT
    )
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


# ---------------------------------------------------------------------------
# 1) 静态扫描：全链无语料/项目特异词汇
# ---------------------------------------------------------------------------


def test_chain_wide_static_scan_has_no_corpus_vocabulary() -> None:
    for path in _CHAIN_FILES:
        assert path.is_file(), path
        source = path.read_text(encoding="utf-8")
        hits = _marker_hits(source)
        assert not hits, f"{path} 含有语料/项目特异词汇：{hits}"


@pytest.mark.parametrize("path", _CHAIN_FILES, ids=lambda p: p.name)
def test_chain_wide_static_scan_timepoint_threshold_patterns(path: Path) -> None:
    """时间点与阈值字面量不得进入共享链（含提示词模板与运行器脚本）。"""
    source = path.read_text(encoding="utf-8")
    for pattern in _TIMEPOINT_OR_THRESHOLD_PATTERNS:
        match = pattern.search(source)
        assert match is None, f"{path} 命中项目特异性模式 {pattern.pattern}: {match.group(0)!r}"


# ---------------------------------------------------------------------------
# 2) 行为验证：生成提示词、合成夹具、重放包工件均词汇中立
# ---------------------------------------------------------------------------


def test_generated_discovery_and_deep_prompts_are_vocabulary_neutral() -> None:
    manifest = _manifest()
    workflow = _workflow()

    discovery_plan = plan_protocol_control_discovery(
        manifest, max_units_per_batch=2, context_radius=1
    )
    discovery_batch = discovery_plan.batches[0]
    discovery_prompt = build_protocol_control_discovery_prompt(discovery_batch)
    discovery_input = ProtocolControlDiscoveryAgentInput.from_batch(discovery_batch)
    discovery_repair = build_protocol_control_discovery_repair_prompt(
        discovery_input, problem="中性修复问题"
    )
    _assert_neutral(discovery_prompt, label="发现提示词")
    _assert_neutral(discovery_repair, label="发现修复提示词")
    assert "输出结构：" in discovery_prompt
    assert "输出结构：" not in discovery_repair
    assert '"$defs"' not in discovery_repair
    assert "只返回包含 wire_version 和 decisions 的单个 JSON 对象" in discovery_repair

    deep_plan = plan_protocol_control_batches(
        manifest,
        max_owned_units_per_batch=1,
        context_radius=1,
        workflow_stages=workflow,
    )
    deep_batch = deep_plan.batches[0]
    deep_prompt = build_protocol_control_agent_prompt(deep_batch)
    deep_input = ProtocolControlAgentInput.from_batch(deep_batch)
    deep_repair = build_protocol_control_repair_prompt(deep_batch, problem="中性修复问题")
    _assert_neutral(deep_prompt, label="深析提示词")
    _assert_neutral(deep_repair, label="深析修复提示词")

    # 冻结输入确实进入提示词且保持中性标识。
    assert "protocol:neutral-chain-wide" in deep_prompt
    assert "protocol:neutral-chain-wide" in discovery_prompt


def test_prompts_separate_subject_controls_from_study_administration() -> None:
    manifest = _manifest()
    workflow = _workflow()
    discovery_batch = plan_protocol_control_discovery(
        manifest, max_units_per_batch=2, context_radius=1
    ).batches[0]
    deep_batch = plan_protocol_control_batches(
        manifest,
        max_owned_units_per_batch=1,
        context_radius=1,
        workflow_stages=workflow,
    ).batches[0]

    discovery_prompt = build_protocol_control_discovery_prompt(discovery_batch)
    deep_prompt = build_protocol_control_agent_prompt(deep_batch)

    assert "若未明确约束某一受试者的资格判断、节点放行或个例证据" in discovery_prompt
    assert "‘研究启动前’或‘任何研究操作前’只说明项目或中心义务的时序" in discovery_prompt
    assert "当前原文不足以区分义务对象时标为 uncertain" in discovery_prompt
    assert "目录标题和页码仅是定位指针" in discovery_prompt
    assert "修订记录若仅描述曾经修改或优化某项标准" in discovery_prompt
    assert "摘要或背景中真正陈述了当前筛选、导入" in discovery_prompt
    assert "必须先区分义务对象是个例受试者" in deep_prompt
    assert "不得仅因时点或名称相似就关联" in deep_prompt
    assert "候选只能引用并表达受试者层面的真实增量" in deep_prompt


def test_bundled_smoke_fixture_docx_is_vocabulary_neutral(tmp_path) -> None:
    smoke = _load_smoke_module()
    docx_path = tmp_path / "synthetic.docx"
    smoke.build_synthetic_protocol_docx(docx_path)

    document = Document(str(docx_path))
    texts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            texts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    texts.append(cell.text)
    assert texts
    _assert_neutral("\n".join(texts), label="冒烟合成方案 DOCX")


def test_replay_harness_pack_artifacts_are_vocabulary_neutral(tmp_path) -> None:
    """无模型重放包从词汇中立输入构建后，所有落盘工件保持词汇中立且可复现。"""
    smoke = _load_smoke_module()
    docx_path = tmp_path / "neutral.docx"
    smoke.build_synthetic_protocol_docx(docx_path)
    sha256 = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    pins = {"protocol_version_id": "NEUTRAL-CHAIN:001:v1.0:phase-ii"}

    def build_pack(out_dir: str) -> Path:
        config = ReplayHarnessConfig(
            schema_version=REPLAY_HARNESS_CONFIG_SCHEMA,
            protocol_path=str(docx_path),
            expected_protocol_sha256=sha256,
            study_phase=StudyPhase.PHASE_II,
            protocol_version_id=pins["protocol_version_id"],
            source_artifact_id="neutral-chain-test",
            owned_source_refs=["body.p2", "body.p4"],
            attached_source_refs=["body.t0.r0"],
            out_dir=out_dir,
            pinned_created_at="2026-08-29T00:00:00+00:00",
        )
        result = build_protocol_replay_pack(config)
        assert verify_replay_pack(result.out_dir) == []
        return result.out_dir

    first = build_pack(str(tmp_path / "pack-a"))
    second = build_pack(str(tmp_path / "pack-b"))

    assert replay_pack_fingerprint(first) == replay_pack_fingerprint(second)
    for out_dir in (first, second):
        for artifact in out_dir.rglob("*"):
            # 只扫描文本工件；原始方案二进制 blob 是源输入副本，不参与词表扫描。
            if (
                not artifact.is_file()
                or artifact.name == "replay-summary.json"
                or artifact.suffix not in {".json", ".txt"}
            ):
                continue
            _assert_neutral(
                artifact.read_text(encoding="utf-8"),
                label=f"重放包工件 {artifact.relative_to(out_dir)}",
            )


# ---------------------------------------------------------------------------
# 3) 异构真实协议只读回放前置门禁（可执行规格）
# ---------------------------------------------------------------------------


def test_heterogeneous_real_protocol_replay_precondition_gate(
    tmp_path, monkeypatch
) -> None:
    """真实协议只读回放前必须满足的全部前置条件，逐项映射到已强制行为。

    G1 模型身份：配置模型必须与 /v1/models 实际加载模型正向匹配；缺失、
        歧义或不一致在任何语义请求之前失效关闭（退出码 2）。
    G2 输入边界：默认只消费词汇中立合成协议；`projects/` 与真实数据目录
        下的方案文件被守卫拒绝（退出码 4）。
    G3 词汇中立：内置合成夹具与重放包工件通过全词表扫描。
    G4 同合同可复现：重放包在相同冻结输入下字节级指纹一致且可校验。
    G5 不物化：任何路径都不建立正式目录，运行记录恒为
        `formal_catalog_materialized=false` 且 `claims_complete=false`。
    """

    smoke = _load_smoke_module()

    # G1：失效关闭退出码与运行前身份核验入口存在。
    assert smoke.EXIT_IDENTITY_FAIL_CLOSED == 2
    assert callable(smoke.verify_live_model_identity)
    assert issubclass(smoke.ProtocolControlModelIdentityError, RuntimeError)

    def raise_mismatch(**kwargs):
        raise smoke.ProtocolControlModelIdentityError(
            "配置模型与服务实际加载模型不一致",
            configured_model="quality-target",
            served_model_ids=["speed-variant-loaded"],
            reason="mismatch",
        )

    monkeypatch.setattr(smoke, "verify_live_model_identity", raise_mismatch)
    record, exit_code = smoke.run_smoke(out_dir=tmp_path / "out")
    assert exit_code == smoke.EXIT_IDENTITY_FAIL_CLOSED
    assert record["model_identity"]["verified_before_run"] is False
    assert record["model_identity"]["failure_class"] == "model_identity_mismatch"
    # 失效关闭必须发生在任何任务建立之前：不创建数据目录与合成文件。
    assert not (tmp_path / "out" / "data_v2").exists()

    # G2：真实项目/真实数据目录守卫存在且拒绝外部方案文件。
    assert smoke.EXIT_USAGE_OR_GUARD == 4
    assert callable(smoke._assert_fixture_allowed)
    fake_repo = tmp_path / "repo"
    projects_dir = fake_repo / "projects"
    projects_dir.mkdir(parents=True)
    real_like = projects_dir / "real.docx"
    real_like.write_bytes(b"docx")
    monkeypatch.setattr(smoke, "REPO_ROOT", fake_repo)
    monkeypatch.setattr(
        smoke,
        "resolve_data_paths",
        lambda **kwargs: type("P", (), {"root": tmp_path / "real_data"})(),
    )
    with pytest.raises(smoke.SmokeGuardError):
        smoke._assert_fixture_allowed(real_like)
    with pytest.raises(smoke.SmokeGuardError):
        smoke._assert_fixture_allowed(tmp_path / "real_data" / "uploaded.docx")
    allowed = tmp_path / "scratch" / "synthetic.docx"
    allowed.parent.mkdir(parents=True)
    allowed.write_bytes(b"docx")
    smoke._assert_fixture_allowed(allowed)

    # G3：内置合成夹具与重放包工件词汇中立（复用前两个行为测试的扫描）。
    docx_path = tmp_path / "synthetic.docx"
    smoke.build_synthetic_protocol_docx(docx_path)
    document = Document(str(docx_path))
    texts = [
        p.text for p in document.paragraphs if p.text.strip()
    ] + [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip()
    ]
    _assert_neutral("\n".join(texts), label="门禁 G3 内置合成夹具")

    # G4：同合同可复现——相同冻结输入的两个重放包指纹一致。
    sha256 = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    pack_a = build_protocol_replay_pack(
        ReplayHarnessConfig(
            schema_version=REPLAY_HARNESS_CONFIG_SCHEMA,
            protocol_path=str(docx_path),
            expected_protocol_sha256=sha256,
            study_phase=StudyPhase.PHASE_II,
            protocol_version_id="NEUTRAL-GATE:001:v1.0:phase-ii",
            source_artifact_id="neutral-gate-test",
            owned_source_refs=["body.p2", "body.p4"],
            attached_source_refs=["body.t0.r0"],
            out_dir=str(tmp_path / "pack-a"),
            pinned_created_at="2026-08-29T00:00:00+00:00",
        )
    )
    pack_b = build_protocol_replay_pack(
        ReplayHarnessConfig(
            schema_version=REPLAY_HARNESS_CONFIG_SCHEMA,
            protocol_path=str(docx_path),
            expected_protocol_sha256=sha256,
            study_phase=StudyPhase.PHASE_II,
            protocol_version_id="NEUTRAL-GATE:001:v1.0:phase-ii",
            source_artifact_id="neutral-gate-test",
            owned_source_refs=["body.p2", "body.p4"],
            attached_source_refs=["body.t0.r0"],
            out_dir=str(tmp_path / "pack-b"),
            pinned_created_at="2026-08-29T00:00:00+00:00",
        )
    )
    assert replay_pack_fingerprint(pack_a.out_dir) == replay_pack_fingerprint(
        pack_b.out_dir
    )
    assert verify_replay_pack(pack_a.out_dir) == []
    assert verify_replay_pack(pack_b.out_dir) == []

    # G5：不物化——运行记录边界恒为未物化且 claims_complete=false。
    assert record["claims_complete"] is False
    assert record["boundaries"]["formal_catalog_materialized"] is False
    assert record["boundaries"]["real_clinical_content_used"] is False
