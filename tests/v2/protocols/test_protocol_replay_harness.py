"""通用重放 harness 的确定性合同测试：模型无关构建、指纹、可复现与校验。"""
from __future__ import annotations

import hashlib
import json
import platform
import subprocess
import sys
from pathlib import Path

import pydantic
import pytest
from pydantic import ValidationError
from docx import Document

from app.domain.contracts.enums import ReviewStage, StudyPhase
from app.domain.contracts.protocol_controls import (
    KnownOfficialRuleTarget,
    ProtocolControlDispositionBatch,
)
from app.protocols.protocol_replay_harness import (
    REPLAY_HARNESS_CONFIG_SCHEMA,
    KnownWorkflowStageTarget,
    ReplayHarnessConfig,
    ReplayHarnessError,
    build_protocol_replay_pack,
    load_replay_harness_config,
    replay_pack_fingerprint,
    verify_replay_pack,
    _unit_by_source_ref,
)

PINNED_CREATED_AT = "2026-08-29T00:00:00+00:00"


def _build_protocol_docx(path: Path) -> None:
    """合成方案：期别标记段落 + 表题表，结构稳定可预测。"""

    doc = Document()
    doc.add_heading("第一章 入选标准", level=1)
    doc.add_paragraph("Ⅱ期参与者必须为18至75岁的受试者。")
    doc.add_paragraph("Ⅱ期筛选血肌酐须小于1.5倍正常上限。")
    doc.add_heading("第二章 排除标准", level=1)
    doc.add_paragraph("既往有重大心血管事件者不得入组。")
    doc.add_paragraph("对研究药物过敏者不得入组。")
    doc.add_heading("第三章 实验室检查", level=1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "Ⅱ期病毒学检查"
    table.cell(1, 1).text = "HBsAg阴性"
    doc.save(str(path))


def _manifest_units(out_dir: Path) -> list[dict]:
    return json.loads((out_dir / "coverage-manifest.json").read_text(encoding="utf-8"))[
        "units"
    ]


def _config(tmp_path: Path, *, docx: Path, **overrides) -> ReplayHarnessConfig:
    values = {
        "schema_version": REPLAY_HARNESS_CONFIG_SCHEMA,
        "protocol_path": str(docx),
        "expected_protocol_sha256": hashlib.sha256(docx.read_bytes()).hexdigest(),
        "study_phase": StudyPhase.PHASE_II,
        "protocol_version_id": "GEN-TEST-001:v1.0:phase-ii",
        "source_artifact_id": "psa-harness-test",
        "owned_source_refs": ["body.p2", "body.p4"],
        "attached_source_refs": ["body.t0.r0"],
        "out_dir": str(tmp_path / "pack"),
        "pinned_created_at": PINNED_CREATED_AT,
    }
    values.update(overrides)
    return ReplayHarnessConfig(**values)


def _rewrite_recorded_artifact(out_dir: Path, rel_path: str, payload: dict) -> None:
    path = out_dir / rel_path
    data = (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    path.write_bytes(data)
    summary_path = out_dir / "replay-summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    summary["artifacts"][rel_path] = {
        "sha256": hashlib.sha256(data).hexdigest(),
        "size_bytes": len(data),
    }
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def test_build_replay_pack_from_raw_docx_resolves_specified_refs(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    config = _config(tmp_path, docx=docx)
    result = build_protocol_replay_pack(config)

    manifest = _manifest_units(result.out_dir)
    expected_owned = {
        unit["source_ref"]: unit
        for unit in manifest
        if "血肌酐" in unit["excerpt"] or "心血管" in unit["excerpt"]
    }
    assert sorted(expected_owned) == sorted(config.owned_source_refs)

    batch = ProtocolControlDispositionBatch.model_validate(
        json.loads((result.out_dir / "replay-batch.json").read_text(encoding="utf-8"))
    )
    # 批次恰好拥有配置范围，且保持原文顺序。
    document_order = [
        unit["source_ref"]
        for unit in sorted(
            (expected_owned[ref] for ref in expected_owned),
            key=lambda item: item["source_order"],
        )
    ]
    assert [unit.source_ref for unit in batch.owned_units] == document_order
    assert [unit.source_ref for unit in batch.context_units] == ["body.t0.r0"]
    assert batch.batch_total == 1
    source_copy = result.out_dir / "source-input" / "blobs" / "protocol_sources"
    source_copy = source_copy / f"{result.protocol_document_sha256}.docx"
    assert source_copy.read_bytes() == docx.read_bytes()
    # 摘要指纹覆盖每个落盘文件（摘要自身除外）。
    summary = json.loads(
        (result.out_dir / "replay-summary.json").read_text(encoding="utf-8")
    )
    on_disk = {
        path.relative_to(result.out_dir).as_posix()
        for path in result.out_dir.rglob("*")
        if path.is_file()
    }
    assert on_disk - {"replay-summary.json"} == set(summary["artifacts"])
    assert summary["mode"] == "model_free"
    assert summary["transport_instantiated"] is False
    assert summary["toolchain"] == {
        "python_version": platform.python_version(),
        "pydantic_version": pydantic.__version__,
        "parser_name": "docx-ooxml",
        "parser_version": summary["toolchain"]["parser_version"],
    }
    replay_input = json.loads(
        (result.out_dir / "replay-input.json").read_text(encoding="utf-8")
    )
    assert "resolved" not in replay_input
    assert str(tmp_path) not in (result.out_dir / "replay-input.json").read_text(
        encoding="utf-8"
    )
    assert verify_replay_pack(result.out_dir) == []


def test_replay_pack_is_byte_for_byte_reproducible(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    first = build_protocol_replay_pack(_config(tmp_path, docx=docx))
    second = build_protocol_replay_pack(
        _config(tmp_path, docx=docx, out_dir=str(tmp_path / "pack-b"))
    )
    assert first.protocol_document_sha256 == second.protocol_document_sha256
    assert first.manifest_id == second.manifest_id
    assert first.batch_id == second.batch_id
    assert first.prompt_sha256 == second.prompt_sha256
    assert replay_pack_fingerprint(first.out_dir) == replay_pack_fingerprint(
        second.out_dir
    )


def test_prompt_is_rebuilt_deterministically_from_saved_batch(tmp_path):
    from app.agents.protocol_control_deconstructor import (
        build_protocol_control_agent_prompt,
    )

    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    result = build_protocol_replay_pack(_config(tmp_path, docx=docx))
    batch = ProtocolControlDispositionBatch.model_validate(
        json.loads((result.out_dir / "replay-batch.json").read_text(encoding="utf-8"))
    )
    rebuilt = build_protocol_control_agent_prompt(batch).encode("utf-8")
    assert rebuilt == (result.out_dir / "prompt.txt").read_bytes()
    summary = json.loads(
        (result.out_dir / "replay-summary.json").read_text(encoding="utf-8")
    )
    assert hashlib.sha256(rebuilt).hexdigest() == summary["prompt_sha256"]


def test_verify_detects_tamper_unexpected_and_missing_files(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    result = build_protocol_replay_pack(_config(tmp_path, docx=docx))
    original_fingerprint = replay_pack_fingerprint(result.out_dir)

    summary_path = result.out_dir / "replay-summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    summary["batch_id"] = "tampered-summary"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False), encoding="utf-8")
    assert any(
        item.startswith("pack_fingerprint:")
        for item in verify_replay_pack(
            result.out_dir, expected_fingerprint=original_fingerprint
        )
    )
    summary_path.write_text(
        json.dumps(summary | {"batch_id": result.batch_id}, ensure_ascii=False, indent=2)
        + "\n",
        encoding="utf-8",
    )

    prompt = result.out_dir / "prompt.txt"
    prompt.write_bytes(prompt.read_bytes() + b"tampered")
    mismatches = verify_replay_pack(result.out_dir)
    assert "sha256:prompt.txt" in mismatches

    prompt.write_bytes((result.out_dir / "replay-batch.json").read_bytes())
    (result.out_dir / "unexpected-note.txt").write_text("stray")
    mismatches = verify_replay_pack(result.out_dir)
    assert "unexpected:unexpected-note.txt" in mismatches
    (result.out_dir / "unexpected-note.txt").unlink()

    (result.out_dir / "replay-batch.json").unlink()
    mismatches = verify_replay_pack(result.out_dir)
    assert "missing:replay-batch.json" in mismatches


def test_build_imports_no_model_transport(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    before = set(sys.modules)
    build_protocol_replay_pack(_config(tmp_path, docx=docx))
    imported = set(sys.modules) - before
    forbidden = (
        "app.agents.protocol_control_agent_transport",
        "app.llm",
        "openai",
        "httpx",
        "requests",
    )
    hits = {
        name
        for name in imported
        if name in forbidden or any(name.startswith(prefix + ".") for prefix in forbidden)
    }
    assert not hits, f"构建重放包期间导入了模型路径：{sorted(hits)}"


def test_cold_import_does_not_load_model_transport() -> None:
    repo_root = Path(__file__).resolve().parents[3]
    code = """
import sys
import app.protocols.protocol_replay_harness
forbidden = ('app.agents.protocol_control_agent_transport', 'app.llm', 'openai', 'httpx', 'requests')
hits = sorted(name for name in sys.modules if name in forbidden or any(name.startswith(prefix + '.') for prefix in forbidden))
assert not hits, hits
"""
    completed = subprocess.run(
        [sys.executable, "-c", code],
        cwd=repo_root,
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr or completed.stdout


def test_pdf_input_rejected_after_entry_decommission(tmp_path):
    """方案 PDF 上传与结构分派入口已下线：PDF 输入必须在构建前显式失败。"""
    pdf = tmp_path / "protocol.pdf"
    pdf.write_bytes(b"%PDF-1.7 decommissioned entry check")
    config = _config(tmp_path, docx=pdf)

    with pytest.raises(ReplayHarnessError, match="仅支持 DOCX 结构提取"):
        build_protocol_replay_pack(config)
    assert not (tmp_path / "pack").exists()


def test_unknown_source_ref_fails_and_leaves_no_pack(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    config = _config(tmp_path, docx=docx, owned_source_refs=["body.p99"])
    with pytest.raises(
        ReplayHarnessError, match="无法在从原始 DOCX 重建的覆盖清单中解析"
    ):
        build_protocol_replay_pack(config)
    assert not (tmp_path / "pack").exists()


def test_duplicate_manifest_source_ref_is_rejected(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    result = build_protocol_replay_pack(_config(tmp_path, docx=docx))
    manifest = json.loads(
        (result.out_dir / "coverage-manifest.json").read_text(encoding="utf-8")
    )
    from app.domain.contracts.protocol_controls import ProtocolSectionCoverageManifest

    contract = ProtocolSectionCoverageManifest.model_validate(manifest)
    duplicate = contract.units[0].model_copy(
        update={"structure_unit_id": "duplicate-unit", "source_order": 999999}
    )
    malformed = contract.model_copy(update={"units": [*contract.units, duplicate]})
    with pytest.raises(ReplayHarnessError, match="重复 source_ref"):
        _unit_by_source_ref(malformed)


def test_source_hash_mismatch_fails_before_output(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    config = _config(tmp_path, docx=docx, expected_protocol_sha256="0" * 64)
    with pytest.raises(ReplayHarnessError, match="SHA-256"):
        build_protocol_replay_pack(config)
    assert not (tmp_path / "pack").exists()


def test_non_empty_out_dir_is_rejected(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    first = build_protocol_replay_pack(_config(tmp_path, docx=docx))
    config = _config(tmp_path, docx=docx, out_dir=str(first.out_dir))
    with pytest.raises(ReplayHarnessError, match="拒绝覆盖既有重放包"):
        build_protocol_replay_pack(config)


def test_ref_range_validation_rejects_duplicates_and_overlap(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    with pytest.raises(ValidationError):
        _config(tmp_path, docx=docx, owned_source_refs=["body.p2", "body.p2"])
    with pytest.raises(ValidationError):
        _config(
            tmp_path,
            docx=docx,
            owned_source_refs=["body.p2"],
            attached_source_refs=["body.p2"],
        )
    with pytest.raises(ValidationError):
        _config(tmp_path, docx=docx, owned_source_refs=["  "])
    with pytest.raises(ValidationError):
        _config(tmp_path, docx=docx, pinned_created_at="2026-08-29T00:00:00")
    with pytest.raises(ValidationError):
        payload = _config(tmp_path, docx=docx).model_dump(mode="json")
        payload.pop("pinned_created_at")
        ReplayHarnessConfig.model_validate(payload)


def test_config_json_loader_rejects_unknown_schema(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    payload = _config(tmp_path, docx=docx).model_dump(mode="json")
    payload["schema_version"] = "phase5/unknown-schema/v9"
    path = tmp_path / "config.json"
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ValidationError):
        load_replay_harness_config(path)


def test_owned_units_keep_document_order_regardless_of_config_order(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    # 配置顺序与文档顺序相反。
    config = _config(tmp_path, docx=docx, owned_source_refs=["body.p4", "body.p2"])
    result = build_protocol_replay_pack(config)
    batch = ProtocolControlDispositionBatch.model_validate(
        json.loads((result.out_dir / "replay-batch.json").read_text(encoding="utf-8"))
    )
    assert [unit.source_ref for unit in batch.owned_units] == ["body.p2", "body.p4"]
    assert verify_replay_pack(result.out_dir) == []


def test_frozen_targets_are_carried_on_replay_batch(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    official = KnownOfficialRuleTarget(
        catalog_item_id="rule:IN-01",
        official_code="IN-01",
        label="入选标准 1",
        position=0,
        source_span_ids=["body.p2"],
        source_excerpts=["Ⅱ期筛选血肌酐须小于1.5倍正常上限。"],
    )
    stage = KnownWorkflowStageTarget(
        workflow_stage_id="stage:screening:1",
        review_stage=ReviewStage.SCREENING,
        display_name="筛选期",
        visit_instance="screening-1",
    )
    config = _config(
        tmp_path,
        docx=docx,
        official_targets=[official],
        workflow_stages=[stage],
    )
    result = build_protocol_replay_pack(config)
    batch = ProtocolControlDispositionBatch.model_validate(
        json.loads((result.out_dir / "replay-batch.json").read_text(encoding="utf-8"))
    )
    assert [item.official_code for item in batch.known_official_targets] == ["IN-01"]
    assert [item.workflow_stage_id for item in batch.known_workflow_stage_targets] == [
        "stage:screening:1"
    ]
    agent_input = json.loads(
        (result.out_dir / "agent-input.json").read_text(encoding="utf-8")
    )
    assert agent_input["known_official_targets"][0]["official_code"] == "IN-01"
    assert verify_replay_pack(result.out_dir) == []


def test_immutability_summary_lists_source_blobs_with_fingerprints(tmp_path):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    result = build_protocol_replay_pack(_config(tmp_path, docx=docx))
    summary = json.loads(
        (result.out_dir / "replay-summary.json").read_text(encoding="utf-8")
    )
    source_blobs = [
        rel
        for rel in summary["artifacts"]
        if rel.startswith("source-input/blobs/protocol_sources/")
    ]
    assert len(source_blobs) == 1
    assert source_blobs[0].endswith(".docx")
    entry = summary["artifacts"][source_blobs[0]]
    blob = (result.out_dir / source_blobs[0]).read_bytes()
    assert entry["sha256"] == hashlib.sha256(blob).hexdigest()
    assert entry["size_bytes"] == len(blob)


@pytest.mark.parametrize(
    ("rel_path", "field", "value", "expected"),
    [
        (
            "replay-input.json",
            "protocol_document_sha256",
            "0" * 64,
            "replay_input.protocol_document_sha256:mismatch",
        ),
        ("coverage-manifest.json", "manifest_id", "tampered", "manifest.manifest_id:mismatch"),
        ("replay-batch.json", "batch_id", "tampered", "batch.batch_id:mismatch"),
        ("agent-input.json", "batch_id", "tampered", "agent_input:batch_projection_mismatch"),
    ],
)
def test_verify_cross_checks_saved_contract_identities(
    tmp_path, rel_path, field, value, expected
):
    docx = tmp_path / "protocol.docx"
    _build_protocol_docx(docx)
    result = build_protocol_replay_pack(_config(tmp_path, docx=docx))
    payload = json.loads((result.out_dir / rel_path).read_text(encoding="utf-8"))
    payload[field] = value
    _rewrite_recorded_artifact(result.out_dir, rel_path, payload)
    assert expected in verify_replay_pack(result.out_dir)


def test_d001_p803_p805_read_only_checkpoint_rebuilds(tmp_path):
    repo_root = Path(__file__).resolve().parents[3]
    checkpoint_dir = (
        repo_root
        / ".trellis/tasks/08-22-phase5-clinical-facts-profile/research/"
        "d001-ii-phase-closure/checkpoints"
    )
    config = load_replay_harness_config(
        checkpoint_dir / "p803-p805-model-free-replay-config.v1.json"
    ).model_copy(update={"out_dir": str(tmp_path / "pack")})
    checkpoint = json.loads(
        (
            checkpoint_dir / "p803-p805-model-free-replay-checkpoint.v4.json"
        ).read_text(encoding="utf-8")
    )

    result = build_protocol_replay_pack(config, base_dir=repo_root)

    assert result.protocol_document_sha256 == checkpoint["protocol_document_sha256"]
    assert result.snapshot_id == checkpoint["snapshot_id"]
    assert result.manifest_id == checkpoint["manifest_id"]
    assert result.batch_id == checkpoint["batch_id"]
    assert result.prompt_sha256 == checkpoint["prompt_sha256"]
    assert result.unit_counts == checkpoint["unit_counts"]
    summary = json.loads(
        (result.out_dir / "replay-summary.json").read_text(encoding="utf-8")
    )
    assert summary["prompt_version"] == checkpoint["prompt_version"]
    assert summary["prompt_template_sha256"] == checkpoint["prompt_template_sha256"]
    assert replay_pack_fingerprint(result.out_dir) == checkpoint[
        "expected_pack_fingerprint"
    ]
    assert verify_replay_pack(
        result.out_dir,
        expected_fingerprint=checkpoint["expected_pack_fingerprint"],
    ) == []
    assert not list(result.out_dir.glob(".staging-*"))


def test_previous_d001_replay_checkpoint_remains_immutable_history() -> None:
    checkpoint_dir = (
        Path(__file__).resolve().parents[3]
        / ".trellis/tasks/08-22-phase5-clinical-facts-profile/research/"
        "d001-ii-phase-closure/checkpoints"
    )
    previous = json.loads(
        (
            checkpoint_dir / "p803-p805-model-free-replay-checkpoint.v1.json"
        ).read_text(encoding="utf-8")
    )
    superseded = json.loads(
        (
            checkpoint_dir / "p803-p805-model-free-replay-checkpoint.v2.json"
        ).read_text(encoding="utf-8")
    )
    superseded_v3 = json.loads(
        (
            checkpoint_dir / "p803-p805-model-free-replay-checkpoint.v3.json"
        ).read_text(encoding="utf-8")
    )
    current = json.loads(
        (
            checkpoint_dir / "p803-p805-model-free-replay-checkpoint.v4.json"
        ).read_text(encoding="utf-8")
    )

    assert previous["prompt_sha256"] == (
        "35d5d5fb9f62ecc752869960b75fa94d48cf578e8da7a8a3e5b5730c303c13a5"
    )
    # v2 冻结后已登记为不可变历史：其提示词哈希与整包指纹均不得被改写。
    assert superseded["prompt_sha256"] == (
        "115812e731ffd40cba7b7d21e022303b68c7777b1468b916f8e22fca57c8453d"
    )
    assert superseded["expected_pack_fingerprint"] == (
        "e842e33cbd8615ac747e35c8d6ee5c0ffd8c3840cbc413a68b908e77d73ecb77"
    )
    assert previous["expected_pack_fingerprint"] != superseded[
        "expected_pack_fingerprint"
    ]
    assert superseded["expected_pack_fingerprint"] != superseded_v3[
        "expected_pack_fingerprint"
    ]
    assert superseded_v3["expected_pack_fingerprint"] != current[
        "expected_pack_fingerprint"
    ]
    assert previous["protocol_document_sha256"] == superseded[
        "protocol_document_sha256"
    ]
    assert superseded["protocol_document_sha256"] == superseded_v3[
        "protocol_document_sha256"
    ]
    assert superseded_v3["protocol_document_sha256"] == current[
        "protocol_document_sha256"
    ]
    assert superseded["unit_counts"] == superseded_v3["unit_counts"] == current[
        "unit_counts"
    ]
    assert superseded["batch_id"] == superseded_v3["batch_id"] == current[
        "batch_id"
    ]
