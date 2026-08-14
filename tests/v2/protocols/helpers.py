"""合成 DOCX 构建助手：用于结构提取的确定性测试夹具。"""
from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn


def build_nested_table_docx(path: Path) -> None:
    """外层 2x2 表格，单元格 (0,0) 内嵌 1x2 表格。"""
    doc = Document()
    outer = doc.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "outer00"
    outer.cell(0, 1).text = "outer01"
    inner = doc.add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "inner00"
    inner.cell(0, 1).text = "inner01"
    cell_tc = outer.cell(0, 0)._tc
    inner._tbl.getparent().remove(inner._tbl)
    cell_tc.append(inner._tbl)
    doc.save(str(path))


def build_gridspan_docx(path: Path) -> None:
    """单行三列表格，水平合并为单个跨 3 列的单元格。"""
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(0, 0).text = "merged"
    doc.save(str(path))


def build_header_docx(path: Path) -> None:
    """default/first/even 页眉页脚 + 第二个继承 section。"""
    doc = Document()
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = True
    section.header.paragraphs[0].text = "DEFAULT_H"
    section.first_page_header.paragraphs[0].text = "FIRST_H"
    section.even_page_header.paragraphs[0].text = "EVEN_H"
    section.footer.paragraphs[0].text = "DEFAULT_F"
    section.first_page_footer.paragraphs[0].text = "FIRST_F"
    doc.add_paragraph("body")
    doc.add_section()  # 第二 section 继承前一 section 的部件
    doc.save(str(path))


def build_tracked_change_header_docx(path: Path) -> None:
    """页眉段落中注入一处未接受修订（``w:ins``）。"""
    doc = Document()
    header = doc.sections[0].header
    paragraph = header.paragraphs[0]
    paragraph.text = "HDR"
    ins = parse_xml(
        f'<w:ins {nsdecls("w")} w:id="1" w:author="tester" '
        f'w:date="2026-08-14T00:00:00Z"><w:r><w:t>inserted</w:t></w:r></w:ins>'
    )
    paragraph._p.append(ins)
    doc.add_paragraph("body")
    doc.save(str(path))


def build_style_numbering_docx(path: Path) -> None:
    """段落用 ``List Number`` 样式（样式定义含 ``w:numPr``，段落无直接编号）。"""
    doc = Document()
    doc.add_paragraph("styled item", style="List Number")
    doc.save(str(path))


def build_based_on_numbering_docx(path: Path) -> None:
    """子样式自身无 numPr，仅通过正确层级的 ``w:basedOn`` 继承编号。"""
    doc = Document()
    parent = doc.styles.add_style("NumberedParent", WD_STYLE_TYPE.PARAGRAPH)
    parent.base_style = doc.styles["List Number"]
    child = doc.styles.add_style("NumberedChild", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    doc.add_paragraph("based-on item", style=child)
    doc.save(str(path))


def build_direct_numbering_docx(path: Path, num_id: int, text: str = "numbered item") -> None:
    """段落直接注入 ``w:numPr``，可指定 ``numId``（含 0 或悬空 id）。"""
    doc = Document()
    p = doc.add_paragraph(text)
    ppr = p._p.get_or_add_pPr()
    numpr = parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
    )
    ppr.append(numpr)
    doc.save(str(path))


def build_shared_role_header_docx(path: Path) -> None:
    """同一页眉部件被 default 与 first 两种角色引用（共享 rId）。"""
    doc = Document()
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.header.paragraphs[0].text = "SHARED_ROLE_H"
    doc.add_paragraph("body")
    sect_pr = section._sectPr
    default_ref = None
    for ref in sect_pr.findall(qn("w:headerReference")):
        if ref.get(qn("w:type")) == "default":
            default_ref = ref
            break
    assert default_ref is not None, "default headerReference 未建立"
    rid = default_ref.get(qn("r:id"))
    first_ref = parse_xml(
        f'<w:headerReference {nsdecls("w", "r")} w:type="first" r:id="{rid}"/>'
    )
    default_ref.addnext(first_ref)
    doc.save(str(path))


def build_empty_body_docx(path: Path) -> None:
    """无正文段落/表格的最小合法 DOCX。"""
    Document().save(str(path))


def build_sdt_body_docx(path: Path) -> None:
    """正文唯一临床段落位于 ``w:sdtContent``，验证内容控件不会被静默漏提。"""
    doc = Document()
    paragraph = doc.add_paragraph("内容控件中的基线前必做检查")
    body = doc.element.body
    body.remove(paragraph._p)
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr/><w:sdtContent/></w:sdt>'
    )
    sdt.find(qn("w:sdtContent")).append(paragraph._p)
    body.insert(0, sdt)
    doc.save(str(path))


def build_altchunk_body_docx(path: Path) -> None:
    """注入未支持的外部内容块，提取必须显式要求核对。"""
    doc = Document()
    body = doc.element.body
    body.insert(0, parse_xml(f'<w:altChunk {nsdecls("w", "r")} r:id="rId999"/>'))
    doc.save(str(path))


def build_corrupt_docx(path: Path) -> None:
    """ZIP 含 ``word/document.xml`` 但内容非 XML（损坏 DOCX）。"""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", "this is not valid xml <<<")
