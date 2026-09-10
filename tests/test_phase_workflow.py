import unittest
import tempfile
import io
import json
import shutil
import asyncio
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient

from app.main import app
from app.deconstructor import (
    DECONSTRUCT_PROMPT,
    extract_docx_criteria,
    extract_docx_protocol_workflow,
    extract_protocol_metadata,
    sanitize_selected_stage_scope_language,
    strengthen_compound_condition_summaries,
    strip_outer_markdown_fence,
)
from app.models import ReviewResult, SubjectInfo, SubjectStatus
from app.markdown_export import ISSUE_VERDICTS, VERDICT_LABELS, generate_markdown_report
from app.batch_sources import collect_review_files, collect_review_files_from_folders, discover_subject_folders
from app.pipeline.bundler import build_evidence_bundle
from app.pipeline.ocr import ocr_documents_parallel, text_requires_high_precision_review
from app.pipeline.reviewer import build_review_messages, extract_rule_ids, parse_review_response, pass_verify_label, run_review
from app.pipeline.reviewer import (
    apply_current_phase_required_evidence_adjustments,
    apply_missing_phase_anchor_date_adjustments,
    apply_phase_timing_adjustments,
)
from app.subject_dates import backfill_subject_icf_date, extract_icf_date_from_cache, subject_anchor_dates
from app.phases import load_review_workflow
import app.router.projects as projects_router
from app.shared import clear_subject_outputs, load_subject_info, save_subject_info


ROOT = Path(__file__).resolve().parents[1]
MGK10_PROTOCOL = Path(
    "/Users/smkzw/Documents/康哲项目资料/MG-K10/SAR/4. Protocol/"
    "MG-K10-SAR-001_临床研究方案_ V2.1_20250919_clean版 .docx"
)
D001_PROTOCOL = Path(
    "/Users/smkzw/Documents/康哲项目资料/AI/入排/test-D001项目/"
    "CMS-D001 银屑病2、3期临床方案 v1.0-2025.12.21.docx"
)
D001_SOURCE_ROOT = Path(
    "/Users/smkzw/Documents/康哲项目资料/AI/入排/test-D001项目/全量-入组"
)


def admin_headers(client: TestClient) -> dict:
    resp = client.post("/api/auth/login", json={"username": "admin", "password": "20121116"})
    assert resp.status_code == 200, resp.text
    data = resp.json()
    return {
        "X-Enrollment-User": data["username"],
        "X-Enrollment-Token": data["token"],
    }


def login_headers(client: TestClient, username: str, password: str = "") -> dict:
    resp = client.post("/api/auth/login", json={"username": username, "password": password})
    assert resp.status_code == 200, resp.text
    data = resp.json()
    return {
        "X-Enrollment-User": data["username"],
        "X-Enrollment-Token": data["token"],
    }


@contextmanager
def temporary_api_project(client: TestClient, headers: dict, code: str):
    """Create an isolated legacy-API project instead of relying on local clinical data."""
    client.delete(f"/api/projects/{code}", headers=headers)
    response = client.post(
        "/api/projects",
        json={"project_code": code, "protocol_id": "UT-001", "name": "上传行为测试"},
        headers=headers,
    )
    assert response.status_code == 201, response.text
    try:
        yield code
    finally:
        client.delete(f"/api/projects/{code}", headers=headers)


class ProtocolWorkflowTests(unittest.TestCase):
    def test_metadata_prioritizes_header_and_front_page_version_date_over_filename(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "HDR-TEST-001_临床研究方案_V1.0_20200101.docx"
            doc = Document()
            header = doc.sections[0].header
            header.paragraphs[0].text = "申办者\t方案编号：HDR-TEST-001"
            header.add_paragraph("试验药物\t版本号/版本日期：V9.9//2026年02月03日")
            table = doc.add_table(rows=3, cols=2)
            table.cell(0, 0).text = "方案标题："
            table.cell(0, 1).text = "一项测试用临床研究方案"
            table.cell(1, 0).text = "方案编号："
            table.cell(1, 1).text = "HDR-TEST-001"
            table.cell(2, 0).text = "方案版本/日期："
            table.cell(2, 1).text = "V9.9//2026年02月03日"
            doc.add_paragraph("临床研究方案")
            doc.save(path)

            metadata = extract_protocol_metadata(str(path))

        self.assertEqual(metadata["project_code"], "HDR-TEST")
        self.assertEqual(metadata["protocol_id"], "HDR-TEST-001")
        self.assertEqual(metadata["protocol_version"], "V9.9")
        self.assertEqual(metadata["protocol_date"], "2026-02-03")
        self.assertEqual(metadata["name"], "一项测试用临床研究方案")

    def test_metadata_ignores_template_and_database_versions_when_protocol_version_pair_exists(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "CMS-D001 银屑病2、3期临床方案 v1.0-2025.12.21.docx"
            doc = Document()
            header = doc.sections[0].header
            header.paragraphs[0].text = "海南德镁医药科技有限责任公司 文件编号：CMSS-SOP-MD-5101-T04"
            header.add_paragraph("版 本 号：00")
            header.add_paragraph("方案编号：D001-02-002 版本号：1.0 版本日期：2025年12月10日")
            doc.add_paragraph("方案标题：评价CMS-D001片治疗中度至重度斑块状银屑病成人患者的有效性和安全性的多中心研究")
            doc.add_paragraph("数据库版本 V28.0")
            doc.add_paragraph("临床研究方案")
            doc.save(path)

            metadata = extract_protocol_metadata(str(path))

        self.assertEqual(metadata["project_code"], "D001-02")
        self.assertEqual(metadata["protocol_id"], "D001-02-002")
        self.assertEqual(metadata["protocol_version"], "V1.0")
        self.assertEqual(metadata["protocol_date"], "2025-12-10")

    @unittest.skipUnless(D001_PROTOCOL.exists(), "CMS-D001 protocol fixture not available")
    def test_extracts_project_metadata_from_d001_protocol_without_database_version_contamination(self):
        metadata = extract_protocol_metadata(str(D001_PROTOCOL))

        self.assertEqual(metadata["project_code"], "D001-02")
        self.assertEqual(metadata["protocol_id"], "D001-02-002")
        self.assertEqual(metadata["protocol_version"], "V1.0")
        self.assertEqual(metadata["protocol_date"], "2025-12-10")
        self.assertIn("CMS-D001", metadata["name"])

    @unittest.skipUnless(D001_PROTOCOL.exists(), "CMS-D001 protocol fixture not available")
    def test_d001_criteria_preserves_common_parent_counts_and_subitems(self):
        extracted = extract_docx_criteria(str(D001_PROTOCOL), "Ⅱ期")

        self.assertEqual(extracted["study_stage"], "Ⅱ期")
        self.assertEqual(extracted["inclusion_count"], 6)
        self.assertEqual(extracted["exclusion_count"], 30)
        self.assertIn("1）自愿签署知情同意书", extracted["inclusion_text"])
        self.assertIn("  - 银屑病面积与严重程度指数（PASI）评分≥12分", extracted["inclusion_text"])
        self.assertIn("30）研究者认为不合适参加本研究的其他原因。", extracted["exclusion_text"])
        self.assertNotIn("31）", extracted["exclusion_text"])

    @unittest.skipUnless(D001_PROTOCOL.exists(), "CMS-D001 protocol fixture not available")
    def test_d001_workflow_extracts_phase_specific_screening_and_baseline_review_points(self):
        workflow = extract_docx_protocol_workflow(str(D001_PROTOCOL))

        self.assertEqual(workflow["study_stages"], ["Ⅱ期", "Ⅲ期"])
        self.assertTrue(workflow["requires_study_stage_selection"])

        phase2 = [phase for phase in workflow["review_phases"] if phase.get("study_stage") == "Ⅱ期"]
        phases = {phase["phase_id"]: phase for phase in phase2}
        self.assertIn("screening_run_in", phases)
        self.assertIn("baseline_randomization", phases)

        screening = phases["screening_run_in"]
        baseline = phases["baseline_randomization"]
        self.assertEqual(screening["visit"], "筛选")
        self.assertEqual(screening["day_window"], "D-28~D-1")
        self.assertNotIn("Ⅲ期", screening["stage"])
        self.assertIn("签署知情同意书", screening["required_items"])
        self.assertIn("PASI评分", screening["required_items"])
        self.assertIn("入排标准审核", screening["required_items"])
        self.assertIn("基线2", baseline["visit"])
        self.assertIn("D≤-7", baseline["day_window"])
        self.assertNotIn("Ⅲ期", baseline["stage"])
        self.assertIn("随机", baseline["required_items"])

    @unittest.skipUnless(D001_PROTOCOL.exists(), "CMS-D001 protocol fixture not available")
    def test_d001_project_scoped_workflow_filters_other_stage_review_points(self):
        workflow = extract_docx_protocol_workflow(str(D001_PROTOCOL))
        scoped = projects_router._project_scoped_workflow(workflow, "Ⅱ期")

        self.assertEqual(scoped["study_stages"], ["Ⅱ期"])
        self.assertFalse(scoped["requires_study_stage_selection"])
        self.assertTrue(scoped["review_phases"])
        self.assertTrue(all(phase.get("study_stage") == "Ⅱ期" for phase in scoped["review_phases"]))
        self.assertFalse(any("Ⅲ期" in (phase.get("stage") or "") for phase in scoped["review_phases"]))

        with tempfile.TemporaryDirectory() as tmp:
            pd = Path(tmp)
            (pd / "config.json").write_text(
                json.dumps({"study_stage": "Ⅱ期"}, ensure_ascii=False),
                encoding="utf-8",
            )
            (pd / "review_phases.json").write_text(
                json.dumps(workflow, ensure_ascii=False),
                encoding="utf-8",
            )
            loaded = load_review_workflow(pd)

        self.assertEqual(loaded["study_stages"], ["Ⅱ期"])
        self.assertTrue(all(phase.get("study_stage") == "Ⅱ期" for phase in loaded["review_phases"]))
        self.assertFalse(any("Ⅲ期" in (phase.get("stage") or "") for phase in loaded["review_phases"]))

    @unittest.skipUnless(MGK10_PROTOCOL.exists(), "MG-K10-SAR protocol fixture not available")
    def test_extracts_project_metadata_from_mgk10_protocol(self):
        metadata = extract_protocol_metadata(str(MGK10_PROTOCOL))

        self.assertEqual(metadata["project_code"], "MG-K10-SAR")
        self.assertEqual(metadata["protocol_id"], "MG-K10-SAR-001")
        self.assertEqual(metadata["protocol_version"], "V2.1")
        self.assertEqual(metadata["protocol_date"], "2025-09-19")
        self.assertIn("MG-K10", metadata["name"])
        self.assertIn("季节性过敏性鼻炎", metadata["name"])

    @unittest.skipUnless(MGK10_PROTOCOL.exists(), "MG-K10-SAR protocol fixture not available")
    def test_extracts_study_stages_and_review_phases_from_mgk10_protocol(self):
        workflow = extract_docx_protocol_workflow(str(MGK10_PROTOCOL))

        self.assertEqual(workflow["study_stages"], ["Ⅱ期", "Ⅲ期"])
        self.assertTrue(workflow["requires_study_stage_selection"])

        phases = {phase["phase_id"]: phase for phase in workflow["review_phases"]}
        self.assertIn("screening_run_in", phases)
        self.assertIn("baseline_randomization", phases)

        screening = phases["screening_run_in"]
        baseline = phases["baseline_randomization"]

        self.assertEqual(screening["visit"], "V1")
        self.assertEqual(screening["day_window"], "D-7~D-1")
        self.assertIn("签署知情同意书", screening["required_items"])
        self.assertIn("过敏原筛查", screening["required_items"])

        self.assertEqual(baseline["visit"], "V2（基线2）")
        self.assertEqual(baseline["day_window"], "D1")
        self.assertIn("随机", baseline["required_items"])
        self.assertIn("血常规", baseline["required_items"])

        self.assertTrue(all(not phase.get("study_stage") for phase in workflow["review_phases"]))
        scoped = projects_router._project_scoped_workflow(workflow, "Ⅲ期")
        self.assertEqual(scoped["study_stages"], ["Ⅲ期"])
        self.assertFalse(scoped["requires_study_stage_selection"])
        self.assertTrue(all(phase.get("study_stage") == "Ⅲ期" for phase in scoped["review_phases"]))
        self.assertFalse(any("Ⅱ期" in (phase.get("stage") or "") for phase in scoped["review_phases"]))

    @unittest.skipUnless(MGK10_PROTOCOL.exists(), "MG-K10-SAR protocol fixture not available")
    def test_mgk10_phase3_criteria_preserves_parent_counts_and_order(self):
        extracted = extract_docx_criteria(str(MGK10_PROTOCOL), "Ⅲ期")

        self.assertEqual(extracted["inclusion_count"], 7)
        self.assertEqual(extracted["exclusion_count"], 16)
        self.assertIn("5）基线时血EOS≥300/μL；", extracted["inclusion_text"])
        self.assertIn("16）经研究者判断，因其他原因不适合参加本研究者。", extracted["exclusion_text"])
        self.assertNotIn("17）", extracted["exclusion_text"])

    def test_deconstruct_prompt_requires_selected_stage_and_exact_parent_numbering(self):
        self.assertIn("父级IN/EX数量、顺序、编号必须与方案原文", DECONSTRUCT_PROMPT)
        self.assertIn("用户必须先选择本次解构Ⅱ期还是Ⅲ期", DECONSTRUCT_PROMPT)
        self.assertIn("把该期别视为一个独立项目规则集", DECONSTRUCT_PROMPT)
        self.assertIn("不得输出一致/差异说明", DECONSTRUCT_PROMPT)
        self.assertIn("IN-04a", DECONSTRUCT_PROMPT)
        self.assertIn("不得把子项、检查项目、流程表项目升级成", DECONSTRUCT_PROMPT)
        self.assertIn("逻辑关系保真", DECONSTRUCT_PROMPT)
        self.assertIn("不得弱化成 OR", DECONSTRUCT_PROMPT)
        self.assertIn("尿糖1+、潜血阳性、GGT升高", DECONSTRUCT_PROMPT)
        self.assertIn("存在重大或不稳定系统性疾病", DECONSTRUCT_PROMPT)
        self.assertIn("研究者明确判断不具备临床研究条件", DECONSTRUCT_PROMPT)
        self.assertIn("TPPA/TP-Ab/梅毒特异性抗体阳性", DECONSTRUCT_PROMPT)
        self.assertIn("研究者明确判断既往感染已治愈", DECONSTRUCT_PROMPT)

    def test_deconstructed_rules_header_uses_extracted_metadata_not_llm_guess(self):
        raw = """---
项目代号: MG-K10-SAR-Ⅲ
方案编码: MG-K10-SAR-301（假设）
---

# MG-K10-SAR-Ⅲ 入排审核规则

#### IN-01 年龄
- **判断点**：年龄
"""
        sanitized = projects_router._sanitize_deconstructed_rules_header(
            raw,
            {
                "project_code": "MG-K10-SAR",
                "protocol_id": "MG-K10-SAR-001",
                "protocol_version": "V2.1",
                "protocol_date": "2025-09-19",
            },
            "Ⅲ期",
        )

        self.assertIn("方案编号: MG-K10-SAR-001", sanitized)
        self.assertIn("解构期别: Ⅲ期", sanitized)
        self.assertNotIn("假设", sanitized)
        self.assertNotIn("MG-K10-SAR-301", sanitized)

    def test_deconstruction_strengthens_ex20h_all_of_summary(self):
        raw = """#### EX-20 实验室检查异常

- **判断点**：筛选或基线时是否存在以下实验室检查异常：
- **子项**：
  - a. 外周血白细胞计数 <3×10⁹/L
  - h. 任何其它实验室检查结果异常且有临床意义，研究者评估参与研究将带来不可接受的风险。
- **通过**：筛选和基线时均无上述任何一项实验室检查异常。
- **不通过**：筛选或基线时存在任一项上述实验室检查异常。
"""
        cleaned = strengthen_compound_condition_summaries(raw)

        self.assertIn("其它实验室检查异常 + 有临床意义 + 经研究者评估参与研究将可能构成不可接受风险", cleaned)
        self.assertIn("h项缺少任一组件时不得直接判不通过", cleaned)
        self.assertNotIn("存在任一项上述实验室检查异常。", cleaned)

    def test_deconstruction_strengthens_researcher_judgment_compound_summary(self):
        raw = """#### EX-11 重大或不稳定的系统性疾病

- **判断点**：首次给药前6个月内是否存在重大或不稳定的系统性疾病。
- **通过**：首次给药前6个月内无上述疾病。
- **不通过**：存在上述系统性疾病。
- **证据不足**：无详细病史记录。
"""
        cleaned = strengthen_compound_condition_summaries(raw)

        self.assertIn("存在相关系统性疾病 + 研究者明确判断不具备临床研究条件", cleaned)
        self.assertIn("仅有疾病存在", cleaned)
        self.assertNotIn("- **不通过**：存在上述系统性疾病。", cleaned)

    def test_deconstruction_strengthens_syphilis_exception_summary(self):
        raw = """#### EX-22 特定感染筛查阳性

- **判断点**：筛选访视时是否存在下列感染：
- **子项**：
  - d. 梅毒：梅毒特异性抗体试验阳性（梅毒非特异性抗体阴性且研究者判断为既往感染已治愈的除外）。
- **通过**：筛选时无上述任何一项感染。
- **不通过**：筛选时存在任一项上述感染。
- **证据不足**：无规定的感染筛查全套检查结果。
"""
        cleaned = strengthen_compound_condition_summaries(raw)

        self.assertIn("TPPA/TP-Ab/梅毒特异性抗体阳性", cleaned)
        self.assertIn("非特异性抗体阴性 + 研究者明确判断既往感染已治愈", cleaned)
        self.assertIn("缺少治愈判断时不得按例外通过", cleaned)

    def test_selected_stage_sanitizer_removes_misleading_only_stage_labels(self):
        raw = """## 一、入选标准（Inclusion）
#### IN-06 避孕要求
- **判断点**：避孕要求
- **适用期别**：仅Ⅲ期。
#### IN-07 知情同意
- **说明**：仅III期使用本编号。
- **说明**：本条与Ⅱ期对应避孕要求实质一致。
## 三、基线及以前方案流程核查
- **本次解构期别**：Ⅲ期
"""
        cleaned = sanitize_selected_stage_scope_language(
            raw,
            "Ⅲ期",
            {"study_stages": ["Ⅱ期", "Ⅲ期"]},
        )

        self.assertNotIn("适用期别", cleaned)
        self.assertNotIn("仅III期", cleaned)
        self.assertNotIn("按本次所选", cleaned)
        self.assertNotIn("对应", cleaned)
        self.assertNotIn("实质一致", cleaned)
        self.assertIn("本次解构期别**：Ⅲ期", cleaned)

    def test_deconstruction_strips_outer_markdown_code_fence(self):
        raw = """```markdown
---
项目代号: D001
---

#### IN-01 知情同意
```"""

        cleaned = strip_outer_markdown_fence(raw)

        self.assertTrue(cleaned.startswith("---"))
        self.assertIn("#### IN-01 知情同意", cleaned)
        self.assertNotIn("```", cleaned)

    @unittest.skipUnless(
        (ROOT / "projects" / "MG-K10-SAR" / "criteria_rules.md").exists(),
        "MG-K10-SAR saved-rules anchor not available",
    )
    def test_mgk10_saved_rules_do_not_mark_shared_items_as_only_phase3(self):
        rules_path = ROOT / "projects" / "MG-K10-SAR" / "criteria_rules.md"
        text = rules_path.read_text(encoding="utf-8")

        self.assertNotRegex(text, r"仅\s*(?:Ⅲ期|III期|III\s*期)")
        self.assertNotRegex(text, r"适用期别|对应条款|对应避孕|对应知情同意|实质一致|按本次所选|另一阶段|另一期别|一致/差异")
        self.assertIn("IN-06 避孕要求", text)
        self.assertIn("IN-07 知情同意、理解和方案依从性", text)


class EvidencePhaseTests(unittest.TestCase):
    @unittest.skipUnless(
        (ROOT / "projects" / "MG-K10-SAR" / "subjects" / "06003").exists(),
        "MG-K10-SAR/06003 evidence anchor not available",
    )
    def test_screening_phase_bundle_excludes_baseline_named_documents(self):
        subject_dir = ROOT / "projects" / "MG-K10-SAR" / "subjects" / "06003"
        if not (subject_dir / "cache").exists():
            self.skipTest("MG-K10-SAR/06003 OCR cache fixture not available")

        screening_bundle = build_evidence_bundle(
            subject_dir=subject_dir,
            project_code="MG-K10-SAR",
            review_phase="screening_run_in",
        )
        baseline_bundle = build_evidence_bundle(
            subject_dir=subject_dir,
            project_code="MG-K10-SAR",
            review_phase="baseline_randomization",
        )

        self.assertIn("审核阶段：**筛选/导入期**", screening_bundle)
        self.assertNotIn("06003基线血常规", screening_bundle)
        self.assertIn("06003基线血常规", baseline_bundle)


class ReviewPromptPhaseTests(unittest.TestCase):
    def test_extract_rule_ids_preserves_order_and_deduplicates(self):
        rules = "#### IN-01 年龄\n正文 IN-01 重复\n#### EX-02 旅行\n#### EX-02 重复"
        self.assertEqual(extract_rule_ids(rules), ["IN-01", "EX-02"])

    def test_review_prompt_includes_phase_and_study_stage_scope(self):
        messages = build_review_messages(
            project_code="MG-K10-SAR",
            anchor_dates={},
            criteria_rules="#### IN-01 年龄\n#### EX-02 旅行",
            evidence_bundle="# bundle",
            review_phase={
                "phase_id": "screening_run_in",
                "name": "筛选/导入期",
                "visit": "V1",
                "day_window": "D-7~D-1",
                "required_items": ["签署知情同意书", "审核入选/排除标准"],
            },
            study_stage="Ⅲ期",
        )
        user_prompt = messages[1]["content"]

        self.assertIn("当前研究阶段：Ⅲ期", user_prompt)
        self.assertIn("当前审核阶段：筛选/导入期", user_prompt)
        self.assertIn("D-7~D-1", user_prompt)
        self.assertIn("必须覆盖的规则ID清单", user_prompt)
        self.assertIn("`IN-01`", user_prompt)
        self.assertIn("`EX-02`", user_prompt)
        self.assertIn("尚未到达的基线/随机前要求", user_prompt)
        self.assertIn("本阶段只判定筛选期组件", user_prompt)
        self.assertIn("基线待后续阶段复核", user_prompt)
        self.assertNotIn("若只有筛选期结果而缺少应到达的基线/D1源文件", user_prompt)
        self.assertIn("语义触发判断", user_prompt)
        self.assertIn("可能性不是排除触发", user_prompt)
        self.assertIn("时间窗交接不清、需确认、待补充", user_prompt)
        self.assertIn("检验项目名精确匹配", user_prompt)
        self.assertIn("GGT、ALP、直接胆红素", user_prompt)
        self.assertIn("证据主题必须匹配当前条款主题", user_prompt)
        self.assertIn("尿糖、潜血、GGT等孤立检验异常不能证明活动性感染", user_prompt)
        self.assertIn("肺功能/FEV1排除条款必须引用可解释数值证据", user_prompt)
        self.assertIn("OCR模糊、数值不完整", user_prompt)
        self.assertIn("全部组件作为AND同时满足", user_prompt)
        self.assertIn("尿糖1+或“未排除临床意义”均不足以判", user_prompt)
        self.assertIn("未排除临床意义、需研究者评估", user_prompt)
        self.assertIn("触发判断：已触发", messages[0]["content"])
        self.assertIn("可能性不是触发", messages[0]["content"])
        self.assertIn("复杂组合/实验室条款", messages[0]["content"])
        self.assertIn("不能用相邻项目、同属肝功能项目", messages[0]["content"])
        self.assertIn("条款主题必须匹配证据主题", messages[0]["content"])
        self.assertIn("合取条件必须全部满足", messages[0]["content"])
        self.assertIn("临床意义和不可接受风险必须有明确判断来源", messages[0]["content"])
        self.assertIn("梅毒筛查例外必须完整", messages[0]["content"])
        self.assertIn("研究者判断类复合条件不能缺项", messages[0]["content"])
        self.assertIn("DeepSeek输出前自检", messages[0]["content"])
        self.assertIn("官方父级规则不可遗漏", messages[0]["content"])
        self.assertIn("整体IE结论不能替代单条证据", messages[0]["content"])
        self.assertIn("例外缺项不能硬套", messages[0]["content"])
        self.assertIn("只有疾病/异常/用药存在", user_prompt)
        self.assertIn("TPPA/TP-Ab/梅毒特异性抗体阳性", user_prompt)
        self.assertIn("研究者明确判断既往感染已治愈", user_prompt)
        self.assertIn("输出前必须逐行执行系统提示中的DeepSeek输出前自检", user_prompt)
        self.assertIn("病史来源需溯源验证", messages[0]["content"])
        self.assertIn("整体也按 pass", messages[0]["content"])
        self.assertIn("整体总结论仍写pass", user_prompt)
        self.assertIn("既往诊断、最早症状、病程时长", messages[0]["content"])

    def test_deepseek_review_kwargs_use_native_thinking_controls(self):
        import app.llm.client as llm_client

        with patch.object(llm_client, "REVIEW_BACKEND", "deepseek"), patch.object(
            llm_client, "REVIEW_REASONING_EFFORT", "high"
        ):
            kwargs = llm_client._review_completion_kwargs(
                model="deepseek-v4-pro",
                messages=[{"role": "user", "content": "json"}],
                temperature=0.3,
                max_tokens=12000,
            )

        self.assertEqual(kwargs["reasoning_effort"], "high")
        self.assertEqual(kwargs["extra_body"], {"thinking": {"type": "enabled"}})
        self.assertNotIn("temperature", kwargs)

    def test_deepseek_review_kwargs_can_use_provider_default_thinking(self):
        import app.llm.client as llm_client

        with patch.object(llm_client, "REVIEW_BACKEND", "deepseek"), patch.object(
            llm_client, "REVIEW_REASONING_EFFORT", "default"
        ):
            kwargs = llm_client._review_completion_kwargs(
                model="deepseek-v4-flash",
                messages=[{"role": "user", "content": "json"}],
                temperature=0.3,
                max_tokens=12000,
            )

        self.assertNotIn("reasoning_effort", kwargs)
        self.assertNotIn("extra_body", kwargs)
        self.assertNotIn("temperature", kwargs)

    def test_local_review_kwargs_do_not_receive_deepseek_only_params(self):
        import app.llm.client as llm_client

        with patch.object(llm_client, "REVIEW_BACKEND", "omlx"):
            kwargs = llm_client._review_completion_kwargs(
                model="Qwen3.6-27B-oQ8-mtp",
                messages=[{"role": "user", "content": "review"}],
                temperature=0.3,
                max_tokens=12000,
            )

        self.assertEqual(kwargs["temperature"], 0.3)
        self.assertNotIn("reasoning_effort", kwargs)
        self.assertNotIn("extra_body", kwargs)

    def test_review_prompt_for_baseline_requires_reached_baseline_evidence(self):
        messages = build_review_messages(
            project_code="MG-K10-SAR",
            anchor_dates={},
            criteria_rules="#### IN-04 病情严重程度",
            evidence_bundle="# bundle",
            review_phase={
                "phase_id": "baseline_randomization",
                "name": "基线/随机前",
                "visit": "V2",
                "day_window": "D1",
                "required_items": ["审核入选/排除标准", "随机"],
            },
            study_stage="Ⅲ期",
        )
        user_prompt = messages[1]["content"]

        self.assertIn("若只有筛选期结果而缺少应到达的基线/D1源文件", user_prompt)
        self.assertIn("不得用筛选期结果替代最终判定", user_prompt)
        self.assertNotIn("本阶段只判定筛选期组件", user_prompt)
        self.assertIn("未提供本次基线/随机前审核锚点日期", user_prompt)
        self.assertIn("随机前/基线前/给药前时间窗", user_prompt)
        self.assertNotIn("请从证据中自行识别关键日期", user_prompt)

    def test_review_prompt_for_baseline_accepts_explicit_washout_denials_as_evidence(self):
        messages = build_review_messages(
            project_code="MG-K10-SAR",
            anchor_dates={"review_phase_anchor_date": "2025-09-10"},
            criteria_rules="#### EX-12 合并禁止用药/治疗规范和洗脱期",
            evidence_bundle="# bundle",
            review_phase={
                "phase_id": "baseline_randomization",
                "name": "基线/随机前",
                "visit": "V2",
                "day_window": "D1",
                "required_items": ["随机", "合并用药核查"],
            },
            study_stage="Ⅲ期",
        )
        user_prompt = messages[1]["content"]

        self.assertIn("本次基线/随机前审核锚点日期", user_prompt)
        self.assertIn("2025-09-10", user_prompt)
        self.assertIn("明确否认相关禁用药物/治疗类别及对应时间窗", user_prompt)
        self.assertIn("不得仅因缺少单独的合并用药记录表", user_prompt)
        self.assertIn("溯源提醒", user_prompt)

    def test_screening_phase_adjustment_does_not_downgrade_future_baseline_gap(self):
        results = [
            ReviewResult(
                rule_id="IN-04",
                rule_name="病情严重程度",
                rule_type="inclusion",
                verdict="insufficient",
                reasoning="筛选期PASI/PGA/BSA均满足标准，但基线评估尚未到达，基线待复核。",
            )
        ]

        changed = apply_phase_timing_adjustments(results, {"phase_id": "screening_run_in"})

        self.assertTrue(changed)
        self.assertEqual(results[0].verdict, "pass_verify")
        self.assertIn("基线待复核", results[0].reasoning)
        self.assertIn("后续阶段复核提醒", results[0].reasoning)
        self.assertEqual(pass_verify_label(results[0].reasoning), "通过（后续阶段复核）")

    def test_screening_phase_adjustment_marks_random_window_as_future_review(self):
        results = [
            ReviewResult(
                rule_id="EX-06f",
                rule_name="随机前10周/5个半衰期内使用生物制剂",
                rule_type="exclusion",
                verdict="investigator",
                reasoning="筛选期病历无实际使用风险；随机日期尚未明确，需随机前常规核实。",
            )
        ]

        changed = apply_phase_timing_adjustments(results, {"phase_id": "screening_run_in"})

        self.assertTrue(changed)
        self.assertEqual(results[0].verdict, "pass_verify")
        self.assertIn("不得用筛选日期替代", results[0].reasoning)
        self.assertEqual(pass_verify_label(results[0].reasoning), "通过（后续阶段复核）")

    def test_screening_phase_adjustment_marks_pass_with_future_component(self):
        results = [
            ReviewResult(
                rule_id="IN-05",
                rule_name="基线血EOS",
                rule_type="inclusion",
                verdict="pass",
                reasoning="筛选期EOS 550/μL，当前符合阈值；基线时需核对数值。",
            )
        ]

        changed = apply_phase_timing_adjustments(results, {"phase_id": "screening_run_in"})

        self.assertTrue(changed)
        self.assertEqual(results[0].verdict, "pass_verify")
        self.assertIn("后续阶段复核提醒", results[0].reasoning)
        self.assertEqual(pass_verify_label(results[0].reasoning), "通过（后续阶段复核）")

    def test_screening_phase_adjustment_marks_na_future_component_as_review_reminder(self):
        results = [
            ReviewResult(
                rule_id="EX-11",
                rule_name="导入期依从性",
                rule_type="exclusion",
                verdict="na",
                reasoning="当前为导入期开始阶段，尚未到评价依从性时间点，待基线随机前确认。",
            )
        ]

        changed = apply_phase_timing_adjustments(results, {"phase_id": "screening_run_in"})

        self.assertTrue(changed)
        self.assertEqual(results[0].verdict, "pass_verify")
        self.assertEqual(pass_verify_label(results[0].reasoning), "通过（后续阶段复核）")

    def test_screening_phase_adjustment_keeps_washout_risk_as_investigator_attention(self):
        results = [
            ReviewResult(
                rule_id="EX-06f",
                rule_name="随机前10周/5个半衰期内使用生物制剂",
                rule_type="exclusion",
                verdict="investigator",
                reasoning="筛选期病历记录既往使用生物制剂；半衰期未知，需研究者随机前核实洗脱期是否足够。",
            )
        ]

        changed = apply_phase_timing_adjustments(results, {"phase_id": "screening_run_in"})

        self.assertTrue(changed)
        self.assertEqual(results[0].verdict, "investigator")
        self.assertIn("随机前重点关注", results[0].reasoning)
        self.assertIn("不得仅凭筛选日期替代", results[0].reasoning)

    def test_screening_phase_adjustment_converts_pass_verify_washout_risk_to_investigator(self):
        results = [
            ReviewResult(
                rule_id="EX-06",
                rule_name="治疗史与用药限制",
                rule_type="exclusion",
                verdict="pass_verify",
                reasoning="曾使用司普奇拜单抗，洗脱期计算（生物制剂半衰期未知）存在潜在不确定性，需随机前确认。",
            )
        ]

        changed = apply_phase_timing_adjustments(results, {"phase_id": "screening_run_in"})

        self.assertTrue(changed)
        self.assertEqual(results[0].verdict, "investigator")
        self.assertIn("随机前重点关注", results[0].reasoning)

    def test_review_parser_separates_future_phase_verify_from_traceability_verify(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-03 | 病程时长 | 入选 | ✅通过（需验证：病史来源需溯源验证） | 【筛选病历】转述病史超过6个月，需补充更早既往源文件。 |
| IN-04 | 基线评分 | 入选 | ✅通过（需验证：基线待后续阶段复核） | 筛选期PASI/PGA/BSA满足，基线/D1尚未到达，后续阶段复核。 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)
        labels = {r.rule_id: pass_verify_label(r.reasoning) for r in parsed["rule_results"]}

        self.assertEqual(parsed["overall_verdict"], "pass")
        self.assertEqual(labels["IN-03"], "通过（溯源提醒）")
        self.assertEqual(labels["IN-04"], "通过（后续阶段复核）")
        self.assertIn("溯源提醒", parsed["summary"])
        self.assertIn("后续阶段复核提醒", parsed["summary"])

    def test_baseline_phase_adjustment_keeps_reached_baseline_gap(self):
        results = [
            ReviewResult(
                rule_id="IN-04",
                rule_name="病情严重程度",
                rule_type="inclusion",
                verdict="insufficient",
                reasoning="筛选期PASI/PGA/BSA均满足标准，但基线评估尚未到达，基线待复核。",
            )
        ]

        changed = apply_phase_timing_adjustments(results, {"phase_id": "baseline_randomization"})

        self.assertFalse(changed)
        self.assertEqual(results[0].verdict, "insufficient")

    def test_baseline_phase_adjustment_blocks_screening_only_current_phase_pass(self):
        results = [
            ReviewResult(
                rule_id="IN-05",
                rule_name="基线血嗜酸性粒细胞（EOS）计数",
                rule_type="inclusion",
                verdict="pass",
                reasoning=(
                    "【筛选-基线检验报告单】p3显示EO# 0.32×10^9/L，换算为320/μL。"
                    "该检查样本采血于筛选期（2025-08-08），目前无基线时独立血常规报告，"
                    "但该值≥300/μL，按现有资料判为通过。"
                ),
            )
        ]

        changed = apply_current_phase_required_evidence_adjustments(
            results,
            {"phase_id": "baseline_randomization", "name": "基线/随机前", "visit": "D1", "required_items": ["随机"]},
            {},
        )

        self.assertTrue(changed)
        self.assertEqual(results[0].verdict, "insufficient")
        self.assertIn("不能用筛选期结果替代普通通过", results[0].reasoning)

    def test_baseline_phase_adjustment_blocks_approximate_random_date_pass(self):
        results = [
            ReviewResult(
                rule_id="EX-06f",
                rule_name="随机前10周或5个半衰期内接受过生物制剂治疗",
                rule_type="exclusion",
                verdict="pass",
                reasoning="末次司普奇拜单抗2025-06-04，距随机日(约2025-08-15)约10周，未触发排除。",
            )
        ]

        changed = apply_current_phase_required_evidence_adjustments(
            results,
            {"phase_id": "baseline_randomization", "name": "基线/随机前", "visit": "D1", "required_items": ["随机"]},
            {},
        )

        self.assertTrue(changed)
        self.assertEqual(results[0].verdict, "insufficient")
        self.assertIn("未提供基线/随机前锚点日期", results[0].reasoning)

    def test_baseline_phase_adjustment_ignores_source_label_baseline_wording(self):
        results = [
            ReviewResult(
                rule_id="EX-07",
                rule_name="疾病史",
                rule_type="exclusion",
                verdict="pass",
                reasoning="【筛选-基线病历】p3-p4 否认急性感染、严重系统疾病、免疫抑制状态。触发判断：未触发。",
            )
        ]

        changed = apply_current_phase_required_evidence_adjustments(
            results,
            {"phase_id": "baseline_randomization", "name": "基线/随机前", "visit": "D1", "required_items": ["随机"]},
            {},
        )

        self.assertFalse(changed)
        self.assertEqual(results[0].verdict, "pass")

    def test_run_review_retries_when_llm_omits_expected_rule_ids(self):
        criteria = """#### IN-01 年龄范围
#### IN-02 诊断
#### EX-01 过敏史
"""
        first = """## 审核结论
**✅ 可入组**

## 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-01 | 年龄范围 | 入选 | ✅ 通过 | 年龄符合。 |
"""
        second = """## 审核结论
**✅ 可入组**

## 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-01 | 年龄范围 | 入选 | ✅ 通过 | 年龄符合。 |
| IN-02 | 诊断 | 入选 | ✅ 通过 | 诊断符合。 |
| EX-01 | 过敏史 | 排除 | ✅ 通过 | 未见过敏史。 |
"""

        async def run_case():
            with tempfile.TemporaryDirectory() as tmp:
                with patch("app.pipeline.reviewer.review_chat", new=AsyncMock(side_effect=[first, second])) as mocked:
                    report = await run_review(
                        subject_id="S001",
                        project_code="UT",
                        anchor_dates={},
                        criteria_rules=criteria,
                        evidence_bundle="# evidence",
                        output_dir=Path(tmp),
                        review_phase={"phase_id": "screening_run_in"},
                    )

                self.assertEqual(mocked.await_count, 2)
                self.assertEqual([r.rule_id for r in report.rule_results], ["IN-01", "IN-02", "EX-01"])
                self.assertTrue((Path(tmp) / "review_attempt_1_raw.md").exists())
                self.assertTrue((Path(tmp) / "review_attempt_2_raw.md").exists())

        asyncio.run(run_case())

    def test_missing_baseline_anchor_does_not_use_screening_date_for_random_window_fail(self):
        results = [
            ReviewResult(
                rule_id="EX-06f",
                rule_name="随机前10周或5个半衰期内接受过生物制剂治疗",
                rule_type="exclusion",
                verdict="fail",
                reasoning="触发判断：已触发。2025年6月4日使用生物制剂，从2025年6月4日至筛选日期2025年8月8日，间隔约9.3周，不足10周。",
            )
        ]

        changed = apply_missing_phase_anchor_date_adjustments(
            results,
            {"phase_id": "baseline_randomization", "name": "基线/随机", "visit": "随机", "required_items": ["随机"]},
            {},
        )

        self.assertTrue(changed)
        self.assertEqual(results[0].verdict, "insufficient")
        self.assertIn("不得用筛选日期", results[0].reasoning)

    def test_missing_baseline_anchor_keeps_true_non_timing_fail(self):
        results = [
            ReviewResult(
                rule_id="EX-15",
                rule_name="筛选前3个月内大量饮酒",
                rule_type="exclusion",
                verdict="fail",
                reasoning="触发判断：已触发。筛选病历记录确认3个月内有大量饮酒。",
            )
        ]

        changed = apply_missing_phase_anchor_date_adjustments(
            results,
            {"phase_id": "baseline_randomization", "name": "基线/随机", "visit": "随机", "required_items": ["随机"]},
            {},
        )

        self.assertFalse(changed)
        self.assertEqual(results[0].verdict, "fail")

    def test_review_parser_downgrades_pass_when_rules_need_confirmation(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-01 | 年龄 | 入选 | ✅ 通过 | 原文支持 |
| IN-02 | 基线评分 | 入选 | ⚠️ 需研究者 | 缺少基线评分明细 |

### 总结论
判定结果：pass
仍需研究者确认。
"""
        parsed = parse_review_response(raw)
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_keeps_overall_pass_when_only_history_source_needs_verification(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-03 | 过敏性鼻炎病史 | 入选 | ✅ 通过（需验证：病史来源需溯源验证） | 【筛选-基线病历】p1 提及："既往过敏性鼻炎病史4年"；未见更早既往病历或诊断证明。 |

### 总结论
判定结果：pass
病史时长需补充既往源文件。
"""
        parsed = parse_review_response(raw)
        self.assertEqual(parsed["rule_results"][0].verdict, "pass_verify")
        self.assertEqual(parsed["overall_verdict"], "pass")

    def test_review_parser_maps_plain_not_pass_before_pass(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-01 | 年龄 | 入选 | 不通过 | 年龄不符合 |

### 总结论
判定结果：needs_evidence
"""
        parsed = parse_review_response(raw)
        self.assertEqual(parsed["rule_results"][0].verdict, "fail")
        self.assertEqual(parsed["overall_verdict"], "fail")

    def test_review_parser_keeps_not_applicable_as_na(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-01 | 随机前项目 | 排除 | 不适用 | 未到本阶段 |

### 总结论
判定结果：pass
全部当前阶段规则通过。
"""
        parsed = parse_review_response(raw)
        self.assertEqual(parsed["rule_results"][0].verdict, "na")
        self.assertEqual(parsed["overall_verdict"], "pass")

    def test_review_parser_maps_ascii_dash_to_na(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-01 | 随机前项目 | 排除 | - | 未到本阶段 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)
        self.assertEqual(parsed["rule_results"][0].verdict, "na")
        self.assertEqual(parsed["overall_verdict"], "pass")

    def test_review_parser_merges_duplicate_parent_rule_rows_conservatively(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-04 | 症状评分 | 入选 | ✅ 通过 | 筛选评分达标 |
| IN-04 | 症状评分 | 入选 | ⚠️ 需研究者 | 基线评分缺少精确均值 |
| EX-06 | 禁用药物 | 排除 | ✅ 通过 | 无禁用药 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)
        self.assertEqual([r.rule_id for r in parsed["rule_results"]], ["IN-04", "EX-06"])
        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("筛选评分达标", parsed["rule_results"][0].reasoning)
        self.assertIn("基线评分缺少精确均值", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_accepts_bold_overall_verdict(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-01 | 排除项 | 排除 | ⚠️ 证据不足 | 缺少关键源文件 |

### 总结论
判定结果：**needs_evidence**
需补充资料。
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["overall_verdict"], "insufficient")

    def test_review_parser_prioritizes_missing_evidence_over_investigator_overall(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-02 | 实验室检查 | 入选 | ⚠️ 证据不足 | 缺少当前阶段应完成的检验报告。 |
| EX-11 | 合并疾病 | 排除 | ⚠️ 需研究者 | 已有病史资料，但缺少研究者对临床研究条件的明确判断。 |

### 总结论
判定结果：needs_evidence
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["overall_verdict"], "insufficient")

    def test_review_parser_corrects_fail_cell_when_reasoning_explicitly_passes(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-20 | 实验室异常 | 排除 | ❌不通过 | HGB 106 g/L，高于排除阈值100 g/L，其余指标未见达到排除标准，故本条通过。 |

### 总结论
判定结果：**fail（不可入组）**
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "pass")
        self.assertNotIn("系统一致性校正", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "pass")

    def test_review_parser_corrects_fail_cell_to_investigator_when_reasoning_requires_confirmation(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-08 | 合并疾病 | 排除 | ❌ 不通过 | 检查结果需研究者评估临床意义，暂不直接判定不通过。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("当前依据不足以作为definitive不通过", parsed["rule_results"][0].reasoning)
        self.assertNotIn("判定为不通过", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_keeps_true_fail_when_reasoning_contains_trigger(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-20 | 实验室异常 | 排除 | ❌ 不通过 | 未触发EX-20e；ALT 210 U/L超过3倍ULN，触发EX-20g。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "fail")
        self.assertEqual(parsed["overall_verdict"], "fail")

    def test_review_parser_downgrades_uncertain_washout_possibility(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-18 | 禁用治疗 | 排除 | ❌不通过 | 局部用药记录截至2026.04.14，首次给药前2周时间窗为2026.04.27后；记录未明确2026.04.27后是否再次使用禁用局部药，存在违规可能，因时间窗交接不清判定为不通过。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("当前依据不足以作为definitive不通过", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_keeps_true_washout_trigger(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-18 | 禁用治疗 | 排除 | ❌不通过 | 【合并用药】显示2026.05.02使用哈西奈德溶液，首次给药前2周时间窗为2026.04.27后，触发判断：已触发EX-18a禁用局部糖皮质激素时间窗。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "fail")
        self.assertEqual(parsed["overall_verdict"], "fail")

    def test_review_parser_does_not_append_researcher_compound_text_to_washout_rule(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-06f | 随机前10周/5个半衰期内使用生物制剂 | 排除 | ❌不通过 | 筛选病历提示既往使用生物制剂，半衰期需确认；随机日期尚未明确，待随机日期核实后确认。当前不能直接判不通过。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)
        result = parsed["rule_results"][0]

        self.assertEqual(result.verdict, "investigator")
        self.assertNotIn("疾病存在、异常存在或用药存在本身", result.reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_does_not_keep_missing_conmed_log_insufficient_when_exact_denials_exist(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-12 | 合并禁止用药/治疗规范和洗脱期 | 排除 | ⚠️证据不足 | 证据包未提供随机前4天的合并用药记录表。但【筛选-基线病历】明确记录：否认近8周或5个半衰期内使用过全身性免疫抑制剂治疗炎症性疾病或自身免疫性疾病；否认近6个月内接受免疫治疗或变应原特异性免疫治疗；否认3个月内接种过或计划在研究期间接种活疫苗/减毒活疫苗。上述病历已逐项否认方案列明禁用用药/治疗并覆盖对应时间窗，未见实际使用证据。 |

### 总结论
判定结果：needs_evidence
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "pass_verify")
        self.assertIn("单独合并用药记录表缺失", parsed["rule_results"][0].reasoning)
        self.assertIn("不应作为证据不足", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "pass")

    def test_review_parser_keeps_positive_biologic_washout_uncertainty_despite_other_denials(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-06 | 治疗史与用药限制 | 排除 | ⚠️证据不足 | 病历逐一否认近4周糖皮质激素、近8周免疫抑制剂等；但用药史显示2025.6.4使用司普奇拜单抗（抗IL-4Rα），未提供洗脱期具体计算，需研究者判断是否在随机前10周或5个半衰期内。缺少单独合并用药记录表。 |

### 总结论
判定结果：needs_evidence
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "insufficient")
        self.assertNotIn("单独合并用药记录表缺失不应作为证据不足", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "insufficient")

    def test_review_parser_downgrades_ggt_substitution_for_alt_ast_tbil_ex20g(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-20 | 实验室检查异常 | 排除 | ❌不通过 | 生化示：GGT 210 U/L（>ULN 60），已超过1.5×ULN，触发EX-20g标准：ALT或AST或总胆红素 ≥1.5×ULN。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("不能据此触发EX-20g标准", parsed["rule_results"][0].reasoning)
        self.assertIn("GGT不能替代ALT/AST/总胆红素", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_keeps_true_alt_ex20g_trigger(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-20 | 实验室检查异常 | 排除 | ❌不通过 | 生化示：ALT 135 U/L（ULN 40），超过3倍ULN，触发EX-20g标准：ALT或AST或总胆红素 ≥1.5×ULN。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "fail")
        self.assertEqual(parsed["overall_verdict"], "fail")

    def test_review_parser_downgrades_urinalysis_substitution_for_ex07(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-07 | 活动性感染或急性疾病状态 | 排除 | ❌不通过 | 【筛选-基线检验报告单】p2 尿常规示：“葡萄糖 1+ ↑”、“潜血 1+ ↑”。【筛选-基线病历】p2 提及：“参与者否认一周存在活动性感染或急性疾病状态”。鉴于尿检异常且有临床意义，应按不通过处理，除非后续补充明确排除证据。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("尿糖/潜血等尿检异常不能直接证明活动性感染", parsed["rule_results"][0].reasoning)
        self.assertNotIn("应按不通过处理", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_keeps_true_ex07_infection_trigger(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-07 | 活动性感染或急性疾病状态 | 排除 | ❌不通过 | 【筛选-基线病历】p2 提及：“首次给药前3天诊断急性支气管炎，仍有发热并接受抗感染治疗”。触发判断：已触发活动性感染/急性疾病状态。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "fail")
        self.assertEqual(parsed["overall_verdict"], "fail")

    def test_review_parser_requires_numeric_fev1_support_for_lung_function_pass(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-08 | 肺功能 | 排除 | ✅通过 | 肺功能报告虽OCR模糊但未提示FEV1%预计值≤50%，未触发排除。 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "insufficient")
        self.assertIn("FEV1占预计值百分比需要可解释数值证据", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "insufficient")

    def test_review_parser_keeps_fev1_pass_when_numeric_percent_predicted_is_present(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-08 | 肺功能 | 排除 | ✅通过 | 肺功能报告显示支气管舒张剂使用前FEV1占预计值百分比为82.4%，>50%，未触发排除。 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "pass")
        self.assertEqual(parsed["overall_verdict"], "pass")

    def test_review_parser_marks_screening_only_history_duration_as_traceability_verify(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-02 | 过敏性鼻炎诊断 | 入选 | ✅通过 | 【筛选-基线病历】p1记载“自2020年开始出现鼻塞、鼻痒、流涕，于2020年确诊为季节性过敏性鼻炎”，病史≥2年；sIgE阳性。 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "pass_verify")
        self.assertIn("病史来源需溯源验证", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "pass")

    def test_review_parser_marks_screening_only_history_since_year_as_traceability_verify(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-02 | 过敏性鼻炎诊断 | 入选 | ✅通过 | a)【筛选-基线病历】p1 诊断“季节性过敏性鼻炎”；b) 病史自2000年起≥2年；c)【筛选-基线检验报告单】提示季节相关过敏原阳性。 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "pass_verify")
        self.assertIn("病史来源需溯源验证", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "pass")

    def test_review_parser_keeps_history_duration_pass_with_prior_source(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-02 | 过敏性鼻炎诊断 | 入选 | ✅通过 | 【既往病历】p1记载“主诉：间断打喷嚏、流涕10年余，诊断：变应性鼻炎”；【筛选-基线检验报告单】sIgE阳性。 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "pass")
        self.assertEqual(parsed["overall_verdict"], "pass")

    def test_review_parser_rewrites_ex07_urinalysis_reasoning_without_infection_bridge(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-07 | 活动性感染或急性疾病状态 | 排除 | ❌不通过 | 【筛选-基线检验报告单】p2 尿常规示：“葡萄糖 1+ ↑”、“潜血 1+ ↑”，此为活动性实验室异常，需研究者判断是否为急性疾病状态。同时，【筛选-基线病历】p2提及：“参与者否认一周存在活动性感染或急性疾病状态”。鉴于尿检异常，需进一步评估，除非后续补充明确排除证据。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("尿糖/潜血等尿检异常不能直接证明活动性感染", parsed["rule_results"][0].reasoning)
        self.assertNotIn("活动性实验室异常", parsed["rule_results"][0].reasoning)
        self.assertNotIn("是否为急性疾病状态", parsed["rule_results"][0].reasoning)

    def test_review_parser_downgrades_syphilis_exception_without_cured_judgment(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-22 | 特定感染筛查阳性 | 排除 | ✅通过 | 【筛选-基线检验报告单】p25梅毒特异性抗体(TPPA)阳性，【筛选-基线检验报告单】p12 TRUST阴性；【筛选-基线病历】p5记录研究者评估为“TPPA阳性...目前无不适，未予治疗”且非活动性感染，符合“非特异性抗体阴性且研究者判断为既往感染已治愈”的例外。余乙肝、丙肝、HIV均阴性。 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("梅毒特异性抗体阳性", parsed["rule_results"][0].reasoning)
        self.assertIn("既往感染已治愈", parsed["rule_results"][0].reasoning)
        self.assertNotIn("符合“非特异性抗体阴性且研究者判断为既往感染已治愈”的例外", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")
        self.assertIn("EX-22", parsed["summary"])

    def test_review_parser_downgrades_syphilis_incomplete_exception_fail_to_investigator(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-22 | 特定感染筛查阳性 | 排除 | ❌不通过 | 【筛选-基线检验报告单】p25梅毒特异性抗体(TPPA)阳性，【筛选-基线检验报告单】p12 TRUST阴性；【筛选-基线病历】p5记录“TPPA阳性，目前无不适，未予治疗”，但未见研究者明确判断既往感染已治愈，因此未完整满足例外，触发判断：已触发EX-22。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("不能因缺少“既往感染已治愈”判断而直接硬判不通过", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_keeps_syphilis_exception_with_explicit_cured_judgment(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-22 | 特定感染筛查阳性 | 排除 | ✅通过 | 【筛选-基线检验报告单】p25梅毒特异性抗体(TPPA)阳性，【筛选-基线检验报告单】p12 TRUST阴性；【筛选-基线病历】p5记录研究者明确判断“梅毒非特异性抗体阴性，既往梅毒感染已治愈，当前无活动性感染”。符合判断：满足方案例外，不触发EX-22d。 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "pass")
        self.assertEqual(parsed["overall_verdict"], "pass")

    def test_review_parser_downgrades_ex11_when_systemic_disease_lacks_unsuitable_judgment(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-11 | 重大或不稳定的系统性疾病 | 排除 | ❌不通过 | 【筛选-基线病历】p5记载既往症状包括“心衰”“冠状动脉供血不足”“左室增大”。这些是方案列出的重大心血管系统疾病，且在首次给药前6个月内病史持续，研究者虽嘱“必要时专科就诊”但未明确判断具备临床研究条件。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("疾病存在", parsed["rule_results"][0].reasoning)
        self.assertIn("研究者明确判断不具备临床研究条件", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")
        self.assertIn("EX-11", parsed["summary"])

    def test_review_parser_keeps_ex11_fail_with_explicit_unsuitable_judgment(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-11 | 重大或不稳定的系统性疾病 | 排除 | ❌不通过 | 【筛选-基线病历】p5记载“症状性充血性心力衰竭，NYHA III级”。研究者明确判断该心血管疾病不稳定且受试者不具备临床研究条件，触发判断：已触发EX-11。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "fail")
        self.assertEqual(parsed["overall_verdict"], "fail")

    def test_review_parser_downgrades_compound_fail_when_required_researcher_judgment_is_unknown(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-16 | 既往IL-12/IL-17/IL-23靶向药物疗效不佳 | 排除 | ❌不通过 | 【既往病历】显示既往使用过司库奇尤单抗，但疗效记录不详，未见研究者明确评估疗效不佳，因此判定不通过。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_downgrades_researcher_compound_rules_without_adverse_component(self):
        cases = [
            (
                "EX-02",
                "其他影响评估的皮肤病",
                "【病历】记载合并湿疹，触发判断：已触发EX-02其他皮肤病排除标准。",
                "影响研究评估",
            ),
            (
                "EX-08",
                "慢性或复发性感染性疾病",
                "【病历】记载慢性前列腺炎，触发判断：已触发EX-08慢性或复发性感染性疾病。",
                "增加安全性风险",
            ),
            (
                "EX-15",
                "精神疾病",
                "【病历】记载既往抑郁症病史，触发判断：已触发EX-15精神疾病史。",
                "影响用药依从性",
            ),
            (
                "EX-21",
                "生命体征/体格检查/ECG/CT异常",
                "【心电图】提示ST-T改变，触发判断：已触发EX-21检查异常。",
                "不可接受风险",
            ),
        ]
        for rule_id, rule_name, reasoning, expected_hint in cases:
            with self.subTest(rule_id=rule_id):
                raw = f"""### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| {rule_id} | {rule_name} | 排除 | ❌不通过 | {reasoning} |

### 总结论
判定结果：fail
"""
                parsed = parse_review_response(raw)

                self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
                self.assertIn(expected_hint, parsed["rule_results"][0].reasoning)
                self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_does_not_apply_d001_id_guard_to_mgk10_ex15_alcohol(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-15 | 有吸毒史、药物滥用史，或筛选前3个月内大量饮酒 | 排除 | ❌不通过 | 触发判断：已触发。【筛选-基线病历】p4记载“饮酒史：确认3个月内有大量饮酒”。这直接构成方案中“筛选前3个月内大量饮酒”的排除情形。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "fail")
        self.assertNotIn("疾病存在、异常存在或用药存在本身不等于完整排除触发", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "fail")

    def test_review_parser_does_not_apply_ex11_systemic_guard_to_other_protocol_ex11(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-11 | 导入期不愿意记录或随机时依从性<80% | 排除 | ❌不通过 | 触发判断：已触发。导入期日志卡记录依从性为72%，低于方案要求80%。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "fail")
        self.assertNotIn("不具备临床研究条件", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "fail")

    def test_review_parser_recomputes_parent_from_same_prefix_children_only(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-07 | 患有以下疾病或疾病史 | 排除 | 🟡需研究者 | 此条为父级标准，结论由子项决定。子项EX-07e、EX-07l、EX-09g存在需研究者进一步评估的情况。 |
| EX-07a | 其他鼻合并症且研究者评估可能影响疗效 | 排除 | ✅通过 | 触发判断：未触发。 |
| EX-07e | 对宠物毛发过敏的PAR患者 | 排除 | 🟡需研究者 | 需研究者明确是否符合例外。 |
| EX-07l | 伴有严重疾病且研究者认为可能影响疗效和安全性 | 排除 | 🟡需研究者 | 需研究者明确判断。 |
| EX-09g | 其他实验室检查异常有临床意义且研究者判断不适合入组 | 排除 | 🟡需研究者 | 需研究者明确是否不适合入组。 |

### 总结论
判定结果：investigator
"""
        parsed = parse_review_response(raw)
        by_id = {row.rule_id: row for row in parsed["rule_results"]}

        self.assertEqual(by_id["EX-07"].verdict, "investigator")
        self.assertIn("EX-07e、EX-07l", by_id["EX-07"].reasoning)
        self.assertNotIn("EX-09g", by_id["EX-07"].reasoning)

    def test_review_parser_inserts_missing_parent_rows_from_children(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-04a | 筛选评分 | 入选 | ✅通过 | 筛选评分满足。 |
| IN-04b | 基线评分 | 入选 | ⚠️证据不足 | 缺少基线日志卡。 |
| EX-09a | 活动性肝炎 | 排除 | ✅通过 | 触发判断：未触发。 |
| EX-09g | 其它实验室异常 | 排除 | 🟡需研究者 | 需研究者明确是否不适合入组。 |

### 总结论
判定结果：insufficient
"""
        parsed = parse_review_response(raw)
        by_id = {row.rule_id: row for row in parsed["rule_results"]}

        self.assertIn("IN-04", by_id)
        self.assertEqual(by_id["IN-04"].verdict, "insufficient")
        self.assertIn("IN-04b", by_id["IN-04"].reasoning)
        self.assertIn("EX-09", by_id)
        self.assertEqual(by_id["EX-09"].verdict, "investigator")
        self.assertIn("EX-09g", by_id["EX-09"].reasoning)

    def test_review_parser_keeps_child_rule_local_when_other_child_needs_review(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-09e | AST或ALT >2×ULN，或TBil ≥1.5×ULN | 排除 | 🟡需研究者 | ALT约1.16×ULN，AST正常，TBIL约1.17×ULN，均未达到EX-09e排除阈值。但数值仍存在异常，需按EX-09g进行评估。本条判定为未触发。 |
| EX-09g | 其他实验室检查异常有临床意义且研究者判断不适合入组 | 排除 | 🟡需研究者 | 肝功能异常具有临床意义，但尚缺研究者明确判断不适合入组。 |

### 总结论
判定结果：investigator
"""
        parsed = parse_review_response(raw)
        by_id = {row.rule_id: row for row in parsed["rule_results"]}

        self.assertEqual(by_id["EX-09e"].verdict, "pass")
        self.assertIn("其他子项另行审核", by_id["EX-09e"].reasoning)
        self.assertEqual(by_id["EX-09g"].verdict, "investigator")

    def test_review_parser_corrects_fail_when_lab_threshold_reasoning_says_not_triggered(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-09e | AST或ALT >2×ULN，或TBil ≥1.5×ULN | 排除 | ❌不通过 | 【筛选-基线检验报告单】p20显示：ALT 76.1 U/L (参考19-80)，AST 50.0 U/L (参考15-40)；ALT >2×ULN (76.1>?? 但未标记ULN值，实际参考上限ALT为80、AST为40，ALT 76.1未超2倍ULN，AST 50.0已超1.25倍ULN但未达2倍)。然而【筛选-基线病历】p7研究者明确评估为：“天门冬氨酸氨基转移酶 50.0 U/L (CS肝功能不全)，丙氨酸氨基转移酶 76.1 U/L (CS肝功能不全)”。虽未直接写明≥2×ULN，但肝功能异常已有临床意义CS。需等待方案精确阈值判定。实验室数据ALT 76.1未超过2倍上限80；AST 50.0也未超过2倍上限40；TBil 10.2正常。触发判断：未达实验室阈值，不触发EX-09e。 |
| EX-09g | 其他实验室检查异常有临床意义且研究者判断不适合入组 | 排除 | 🟡需研究者 | 肝功能异常被标注CS，但尚缺研究者明确判断不适合入组。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)
        by_id = {row.rule_id: row for row in parsed["rule_results"]}

        self.assertEqual(by_id["EX-09e"].verdict, "pass")
        self.assertEqual(by_id["EX-09g"].verdict, "investigator")
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_ignores_protocol_threshold_symbols_in_final_correction(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-09 | 肝功能异常父级 | 排除 | ❌不通过 | 父级汇总项，结论由子项决定。 |
| EX-09e | AST或ALT >2×ULN，或TBil ≥1.5×ULN | 排除 | ❌ 不通过 | 【筛选-基线检验报告单】p8 提及：“ALT 57.8 U/L ↑”（参考范围9-50 U/L，1.16×ULN）、“AST 34.9 U/L”（参考范围15-40 U/L，0.87×ULN）、“TBIL 30.5 μmol/L ↑”（参考范围≤26 μmol/L，1.17×ULN）。触发判断：已触发。ALT虽高于上限但未超2倍，AST正常，TBil高于上限但也未超1.5倍。但综合来看，并未达到排除阈值。重判：ALT 57.8 < 100 U/L (2×ULN)，TBil 30.5 < 39 μmol/L (1.5×ULN)。符合判断：数值未超排除标准限值，故不触发。--- 纠正：数值未达方案设定的“>2×ULN”或“≥1.5×ULN”标准，本项通过。 |
| EX-09g | 其它实验室异常且需研究者判断 | 排除 | ⚠️需研究者 | GGT、胆红素分项异常已有CS标记，但尚缺研究者明确判断参与研究构成不可接受风险。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)
        by_id = {row.rule_id: row for row in parsed["rule_results"]}

        self.assertEqual(by_id["EX-09e"].verdict, "pass")
        self.assertEqual(by_id["EX-09g"].verdict, "investigator")
        self.assertEqual(by_id["EX-09"].verdict, "investigator")
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_downgrades_current_phase_missing_evidence_fail_to_insufficient(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-09e | AST/ALT >2×ULN 或 TBil ≥1.5×ULN | 排除 | ❌不通过 | ALT 75.1 U/L，约1.5×ULN且未达到方案禁止的实验室阈值；当前基线报告单未显示有效数据，D1/基线期血生化结果缺失，无法确证ALT/AST是否进一步升高至排除标准或仍处于合格范围。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)
        by_id = {row.rule_id: row for row in parsed["rule_results"]}

        self.assertEqual(by_id["EX-09e"].verdict, "insufficient")
        self.assertIn("当前阶段关键证据缺失", by_id["EX-09e"].reasoning)
        self.assertNotIn("疾病存在、异常存在或用药存在本身不等于完整排除触发", by_id["EX-09e"].reasoning)
        self.assertEqual(parsed["overall_verdict"], "insufficient")

    def test_review_parser_downgrades_unreadable_current_phase_inclusion_fail_to_insufficient(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-05 | 基线血EOS≥300/μL | 入选 | ❌不通过 | 【筛选-基线检验报告单】p1 为基线期血常规，但OCR严重损坏无法提取EOS绝对值；筛选期血常规EO#为0.25×10⁹/L，邮件提示“血EOS250，D1需进行检测”。基线期血常规报告虽已提供但不可读，且无其余D1报告证明EOS≥300。符合判断：不满足IN-05。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)
        by_id = {row.rule_id: row for row in parsed["rule_results"]}

        self.assertEqual(by_id["IN-05"].verdict, "insufficient")
        self.assertIn("当前阶段关键证据缺失", by_id["IN-05"].reasoning)
        self.assertEqual(parsed["overall_verdict"], "insufficient")

    def test_review_parser_generalizes_syphilis_exception_guard_beyond_ex22(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-09c | 筛选时TP-Ab阳性者，RPR或TRUST阴性者除外 | 排除 | ✅通过 | TPPA阳性，TRUST阴性；病历仅记录目前无不适，未予治疗，符合非特异性抗体阴性的例外。 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)
        by_id = {row.rule_id: row for row in parsed["rule_results"]}

        self.assertEqual(by_id["EX-09c"].verdict, "investigator")
        self.assertIn("既往感染已治愈", by_id["EX-09c"].reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_uses_ex11_specific_researcher_component_not_generic_risk(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-11 | 重大或不稳定的系统性疾病 | 排除 | ⚠️需研究者 | 受试者有高血压、肝囊肿、胆囊结石、高脂血症、高尿酸血症等病史。研究者尚未明确评估这些情况是否对受试者参与研究构成不可接受的风险。 |

### 总结论
判定结果：needs_evidence
"""
        parsed = parse_review_response(raw)

        reasoning = parsed["rule_results"][0].reasoning
        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("研究者明确判断不具备临床研究条件", reasoning)
        self.assertNotIn("不可接受的风险", reasoning)

    def test_review_parser_downgrades_ex11_fail_when_only_generic_unacceptable_risk_is_stated(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-11 | 重大或不稳定的系统性疾病 | 排除 | ❌不通过 | 受试者有高血压、肝囊肿和胆囊结石等病史。研究者评估这些情况对参与研究构成不可接受风险，触发判断：已触发EX-11。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        reasoning = parsed["rule_results"][0].reasoning
        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("研究者明确判断不具备临床研究条件", reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_downgrades_global_ie_substitution_for_lab_researcher_component(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-09g | 其它有临床意义实验室异常且研究者判断不适合入组 | 排除 | ✅通过 | AST/ALT升高被医生标注为肝功能不全（有临床意义），但筛选病历中研究者明确“初步符合入排标准，且不符合排除标准”并发放导入期药物，表明未判断为不适合入组；触发条件不成立。 |

### 总结论
判定结果：pass
"""
        parsed = parse_review_response(raw)
        by_id = {row.rule_id: row for row in parsed["rule_results"]}

        self.assertEqual(by_id["EX-09g"].verdict, "investigator")
        self.assertIn("整体IE/可入组结论不能替代单条研究者判断", by_id["EX-09g"].reasoning)
        self.assertEqual(by_id["EX-09"].verdict, "investigator")
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_keeps_researcher_compound_fail_with_adverse_component(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-02 | 其他影响评估的皮肤病 | 排除 | ❌不通过 | 【病历】记载合并湿疹，研究者明确判断该皮肤病可能影响研究评估，触发判断：已触发EX-02。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "fail")
        self.assertEqual(parsed["overall_verdict"], "fail")

    def test_review_parser_downgrades_unconfirmed_clinical_significance_ex20h(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-20 | 实验室检查异常 | 排除 | ❌不通过 | 尿常规示：“葡萄糖 1+ ↑”。此为EX-20h中“任何其它实验室检查结果异常且有临床意义”，且未排除临床意义，因此不通过。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("必须同时满足实验室异常、有临床意义和研究者不可接受风险评估", parsed["rule_results"][0].reasoning)
        self.assertNotIn("因此不通过", parsed["rule_results"][0].reasoning)
        self.assertNotIn("其存在已触发排除", parsed["rule_results"][0].reasoning)
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_uses_ex20h_specific_message_when_investigator_assessment_is_pending(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-20 | 实验室检查异常 | 排除 | ❌不通过 | 尿常规示：“葡萄糖 1+ ↑”。此为EX-20h中“任何其它实验室检查结果异常且有临床意义”，且未排除临床意义，因此需研究者评估。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertIn("EX-20h等复合条款必须同时满足实验室异常、有临床意义和研究者不可接受风险评估", parsed["rule_results"][0].reasoning)
        self.assertNotIn("疾病存在、异常存在或用药存在本身不等于完整排除触发", parsed["rule_results"][0].reasoning)

    def test_review_parser_downgrades_ex20h_when_cs_is_explicit_but_risk_is_missing(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-20 | 实验室检查异常 | 排除 | ❌不通过 | 尿常规示：“葡萄糖 1+ ↑”。研究者评估该实验室异常有临床意义，触发判断：已触发EX-20h。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "investigator")
        self.assertEqual(parsed["overall_verdict"], "investigator")

    def test_review_parser_keeps_explicit_investigator_cs_ex20h_fail(self):
        raw = """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-20 | 实验室检查异常 | 排除 | ❌不通过 | 尿常规示：“蛋白 3+ ↑”。研究者评估该实验室异常有临床意义且参与研究将带来不可接受风险，触发判断：已触发EX-20h。 |

### 总结论
判定结果：fail
"""
        parsed = parse_review_response(raw)

        self.assertEqual(parsed["rule_results"][0].verdict, "fail")
        self.assertEqual(parsed["overall_verdict"], "fail")

    def test_markdown_export_treats_pass_verify_as_traceability_note(self):
        self.assertEqual(VERDICT_LABELS["pass_verify"], "通过（需验证）")
        self.assertNotIn("pass_verify", ISSUE_VERDICTS)
        self.assertEqual(pass_verify_label("病史来源需溯源验证"), "通过（溯源提醒）")
        self.assertEqual(pass_verify_label("基线/D1/随机前尚未到达，后续阶段复核"), "通过（后续阶段复核）")

    def test_markdown_export_note_uses_neutral_reminder_wording(self):
        with tempfile.TemporaryDirectory() as tmp:
            text = generate_markdown_report(
                project_path=Path(tmp),
                project={"project_code": "UT", "name": "UT"},
                scope="project",
                phase="baseline_randomization",
            )

        self.assertIn("同属提醒项", text)
        self.assertNotIn("低干预", text)
        self.assertNotIn("低强度", text)

    def test_default_parallel_ocr_allows_up_to_eight_concurrent_vlm_calls(self):
        async def run_case():
            with tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                files = []
                for idx in range(9):
                    path = root / f"source_{idx}.png"
                    path.write_bytes(b"not-a-real-image-but-passed-to-vlm")
                    files.append({"path": str(path), "stem": path.stem})

                active = 0
                max_active = 0
                lock = asyncio.Lock()

                async def fake_vision_ocr(*args, **kwargs):
                    nonlocal active, max_active
                    async with lock:
                        active += 1
                        max_active = max(max_active, active)
                    await asyncio.sleep(0.02)
                    async with lock:
                        active -= 1
                    return "OCR文本"

                with patch("app.pipeline.ocr.call_vision_ocr", side_effect=fake_vision_ocr):
                    result = await ocr_documents_parallel(files, root / "cache")

                self.assertEqual(len(result["results"]), 9)
                self.assertEqual(max_active, 8)

        asyncio.run(run_case())

    def test_parallel_ocr_uses_one_global_limit_across_batches(self):
        async def run_case():
            with tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)

                def make_files(prefix: str):
                    files = []
                    for idx in range(9):
                        path = root / f"{prefix}_{idx}.png"
                        path.write_bytes(b"not-a-real-image-but-passed-to-vlm")
                        files.append({"path": str(path), "stem": path.stem})
                    return files

                active = 0
                max_active = 0
                lock = asyncio.Lock()

                async def fake_vision_ocr(*args, **kwargs):
                    nonlocal active, max_active
                    async with lock:
                        active += 1
                        max_active = max(max_active, active)
                    await asyncio.sleep(0.02)
                    async with lock:
                        active -= 1
                    return "OCR文本"

                with patch("app.pipeline.ocr.call_vision_ocr", side_effect=fake_vision_ocr):
                    first, second = await asyncio.gather(
                        ocr_documents_parallel(make_files("a"), root / "cache_a"),
                        ocr_documents_parallel(make_files("b"), root / "cache_b"),
                    )

                self.assertEqual(len(first["results"]), 9)
                self.assertEqual(len(second["results"]), 9)
                self.assertEqual(max_active, 8)

        asyncio.run(run_case())

    def test_native_text_with_critical_polarity_requires_high_precision_review(self):
        self.assertTrue(text_requires_high_precision_review("饮酒史：确认3个月内有大量饮酒"))
        self.assertTrue(text_requires_high_precision_review("饮酒史：否认3个月内有大量饮酒"))
        self.assertTrue(text_requires_high_precision_review("否认近8周或5个半衰期内使用过全身性免疫抑制剂治疗炎症性疾病"))
        self.assertTrue(text_requires_high_precision_review("梅毒特异性抗体(TPPA)阳性，TRUST阴性"))
        self.assertFalse(text_requires_high_precision_review("本页为空白说明文字，仅用于流程说明。"))


class ProjectRouterTests(unittest.TestCase):
    def test_health_endpoint_exposes_enrollment_service_identity(self):
        with patch("app.llm.client.check_omlx", new=AsyncMock(return_value=True)), patch(
            "app.llm.client.check_deepseek", new=AsyncMock(return_value=True)
        ):
            response = TestClient(app).get("/api/health")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["service"], "enrollment-review-app")
        self.assertEqual(payload["version"], "2.0.0")

    def test_project_router_has_module_level_json_import_for_create_project(self):
        self.assertEqual(projects_router.json.dumps({"ok": True}), '{"ok": true}')

    def test_generic_protocol_title_is_not_saved_as_project_name(self):
        self.assertEqual(
            projects_router._project_display_name({"name": "临床研究方案"}, "MG-K10-SAR"),
            "MG-K10-SAR",
        )

    def test_static_logo_is_served(self):
        with TestClient(app) as client:
            resp = client.get("/static/header_logo.svg")
        self.assertEqual(resp.status_code, 200, resp.text[:100])
        self.assertIn("<svg", resp.text)
        self.assertIn("ER", resp.text)

    def test_report_endpoint_prefers_postprocessed_report_over_raw_response(self):
        from app.shared import projects

        code = "UT-REPORT-PREF"
        sid = "S001"
        pd = ROOT / "projects" / code
        if pd.exists():
            shutil.rmtree(pd)
        projects.pop(code, None)
        try:
            with TestClient(app) as client:
                headers = admin_headers(client)
                created = client.post(
                    "/api/projects",
                    json={"project_code": code, "protocol_id": "UT-RPT-001", "name": "Report Preference"},
                    headers=headers,
                )
                self.assertEqual(created.status_code, 201, created.text)
                subject = client.post(
                    f"/api/projects/{code}/subjects",
                    json={"subject_id": sid},
                    headers=headers,
                )
                self.assertEqual(subject.status_code, 201, subject.text)

                report_dir = pd / "subjects" / sid / "llm" / "screening_run_in"
                report_dir.mkdir(parents=True, exist_ok=True)
                (report_dir / "review_raw.md").write_text(
                    """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-22 | 感染筛查 | 排除 | ✅通过 | 乙肝、丙肝、HIV、梅毒筛查均阴性。 |

### 总结论
判定结果：pass
""",
                    encoding="utf-8",
                )
                (report_dir / "review_report.md").write_text(
                    """## 审核结论

**⚠️ 待补证**  
未见可直接确认的不可入组项；EX-22仍需补充资料或研究者判断，当前判定为待补证。

---

## 逐条审核结果

| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-22 | 感染筛查 | 排除 | ⚠️ 需研究者 | TPPA阳性，TRUST阴性；尚缺研究者明确判断既往感染已治愈。 |

---
""",
                    encoding="utf-8",
                )

                resp = client.get(
                    f"/api/projects/{code}/subjects/{sid}/report?phase=screening_run_in",
                    headers=headers,
                )
                self.assertEqual(resp.status_code, 200, resp.text)
                data = resp.json()
                self.assertEqual(data["verdict"], "investigator")
                self.assertIn("EX-22", data["summary"])
                self.assertIn("需研究者", data["summary"])
                self.assertEqual(data["rule_results"][0]["result"], "investigator")
                self.assertIn("raw_response", data)
        finally:
            if pd.exists():
                shutil.rmtree(pd)
            projects.pop(code, None)

    def test_subject_phase_anchor_dates_are_saved_and_used_for_review(self):
        from app.shared import projects

        code = "UT-PHASE-ANCHOR"
        sid = "S001"
        pd = ROOT / "projects" / code
        if pd.exists():
            shutil.rmtree(pd)
        projects.pop(code, None)
        try:
            with TestClient(app) as client:
                headers = admin_headers(client)
                created = client.post(
                    "/api/projects",
                    json={"project_code": code, "protocol_id": "UT-ANCHOR-001", "name": "Anchor Project"},
                    headers=headers,
                )
                self.assertEqual(created.status_code, 201, created.text)
                subject = client.post(
                    f"/api/projects/{code}/subjects",
                    json={"subject_id": sid},
                    headers=headers,
                )
                self.assertEqual(subject.status_code, 201, subject.text)
                patched = client.patch(
                    f"/api/projects/{code}/subjects/{sid}",
                    json={"phase_anchor_dates": {"baseline_randomization": "2025年9月10日"}},
                    headers=headers,
                )
                self.assertEqual(patched.status_code, 200, patched.text)
                self.assertEqual(patched.json()["phase_anchor_dates"]["baseline_randomization"], "2025-09-10")

                anchors = subject_anchor_dates(pd / "subjects" / sid, review_phase="baseline_randomization")
                self.assertEqual(anchors["review_phase_anchor_date"], "2025-09-10")
        finally:
            if pd.exists():
                shutil.rmtree(pd)
            projects.pop(code, None)

    def test_subject_list_exposes_phase_review_summaries(self):
        from app.shared import projects

        code = "UT-PHASE-LIST"
        sid = "S001"
        pd = ROOT / "projects" / code
        if pd.exists():
            shutil.rmtree(pd)
        projects.pop(code, None)
        try:
            with TestClient(app) as client:
                headers = admin_headers(client)
                created = client.post(
                    "/api/projects",
                    json={"project_code": code, "protocol_id": "UT-LIST-001", "name": "Phase List"},
                    headers=headers,
                )
                self.assertEqual(created.status_code, 201, created.text)
                (pd / "review_phases.json").write_text(
                    json.dumps({
                        "study_stages": ["Ⅲ期"],
                        "requires_study_stage_selection": False,
                        "review_phases": [
                            {"phase_id": "screening_run_in", "name": "筛选期", "visit": "V1", "day_window": "D-7~D-1"},
                            {"phase_id": "baseline_randomization", "name": "基线/随机前", "visit": "V2", "day_window": "D1"},
                        ],
                    }, ensure_ascii=False),
                    encoding="utf-8",
                )
                subject = client.post(
                    f"/api/projects/{code}/subjects",
                    json={"subject_id": sid},
                    headers=headers,
                )
                self.assertEqual(subject.status_code, 201, subject.text)
                for phase, verdict in [("screening_run_in", "pass"), ("baseline_randomization", "insufficient")]:
                    report_dir = pd / "subjects" / sid / "llm" / phase
                    report_dir.mkdir(parents=True, exist_ok=True)
                    report_dir.joinpath("review_report.md").write_text(
                        f"""## 审核结论

**{verdict}**

---

## 逐条审核结果

| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| IN-01 | 知情同意 | 入选 | {'✅通过' if verdict == 'pass' else '⚠️证据不足'} | 测试依据。 |

### 总结论
判定结果：{verdict}
""",
                        encoding="utf-8",
                    )

                resp = client.get(f"/api/projects/{code}/subjects", headers=headers)
                self.assertEqual(resp.status_code, 200, resp.text)
                row = resp.json()[0]
                self.assertIn("phase_reviews", row)
                self.assertEqual(row["phase_reviews"]["screening_run_in"]["verdict"], "pass")
                self.assertEqual(row["phase_reviews"]["baseline_randomization"]["verdict"], "insufficient")
                self.assertTrue(row["phase_reviews"]["baseline_randomization"]["has_report"])
        finally:
            if pd.exists():
                shutil.rmtree(pd)
            projects.pop(code, None)

    def test_report_endpoint_marks_parent_child_rule_hierarchy(self):
        from app.shared import projects

        code = "UT-RULE-HIER"
        sid = "S001"
        pd = ROOT / "projects" / code
        if pd.exists():
            shutil.rmtree(pd)
        projects.pop(code, None)
        try:
            with TestClient(app) as client:
                headers = admin_headers(client)
                created = client.post(
                    "/api/projects",
                    json={"project_code": code, "protocol_id": "UT-HIER-001", "name": "Rule Hierarchy"},
                    headers=headers,
                )
                self.assertEqual(created.status_code, 201, created.text)
                subject = client.post(
                    f"/api/projects/{code}/subjects",
                    json={"subject_id": sid},
                    headers=headers,
                )
                self.assertEqual(subject.status_code, 201, subject.text)
                report_dir = pd / "subjects" / sid / "llm" / "screening_run_in"
                report_dir.mkdir(parents=True, exist_ok=True)
                report_dir.joinpath("review_report.md").write_text(
                    """## 审核结论

**✅ 通过**

---

## 逐条审核结果

| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-06 | 禁用治疗 | 排除 | ✅通过 | 父级汇总项。 |
| EX-06f | 生物制剂洗脱 | 排除 | ✅通过 | 未见随机前时间窗内使用。 |

### 总结论
判定结果：pass
""",
                    encoding="utf-8",
                )

                resp = client.get(
                    f"/api/projects/{code}/subjects/{sid}/report?phase=screening_run_in",
                    headers=headers,
                )
                self.assertEqual(resp.status_code, 200, resp.text)
                rows = {row["rule_id"]: row for row in resp.json()["rule_results"]}
                self.assertTrue(rows["EX-06"]["is_parent_rule"])
                self.assertFalse(rows["EX-06"].get("is_child_rule"))
                self.assertTrue(rows["EX-06f"]["is_child_rule"])
                self.assertEqual(rows["EX-06f"]["parent_rule_id"], "EX-06")
                self.assertEqual(rows["EX-06f"]["hierarchy_level"], 1)
        finally:
            if pd.exists():
                shutil.rmtree(pd)
            projects.pop(code, None)

    def test_backfill_reports_rewrites_legacy_needs_evidence_to_specific_verdict(self):
        from app.report_backfill import backfill_project_reports

        with tempfile.TemporaryDirectory() as tmp:
            project_path = Path(tmp) / "UT-BACKFILL"
            subject_dir = project_path / "subjects" / "S001"
            report_dir = subject_dir / "llm" / "screening_run_in"
            report_dir.mkdir(parents=True)
            save_subject_info(
                subject_dir,
                SubjectInfo(
                    subject_id="S001",
                    project_code="UT-BACKFILL",
                    status=SubjectStatus.REVIEWED.value,
                    overall_verdict="needs_evidence",
                ),
            )
            (report_dir / "review_report.md").write_text(
                """## 审核结论

**⚠️ 待补证**  
未见可直接确认的不可入组项；EX-22仍需补充资料或研究者判断，当前判定为待补证。

---

## 逐条审核结果

| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-22 | 感染筛查 | 排除 | ⚠️ 需研究者 | TPPA阳性，TRUST阴性；尚缺研究者明确判断既往感染已治愈。 |

---
""",
                encoding="utf-8",
            )

            result = backfill_project_reports(
                project_path,
                project_code="UT-BACKFILL",
                phase="screening_run_in",
            )

            info = load_subject_info(subject_dir)
            rewritten = (report_dir / "review_report.md").read_text(encoding="utf-8")
            self.assertEqual(result["updated"], 1)
            self.assertEqual(result["subjects"][0]["new_overall"], "investigator")
            self.assertEqual(info.overall_verdict, "investigator")
            self.assertIn("**🟡 需研究者判定**", rewritten)
            self.assertNotIn("待补证", rewritten)

    def test_backfill_reports_prefers_raw_response_over_old_postprocessed_report(self):
        from app.report_backfill import backfill_project_reports

        with tempfile.TemporaryDirectory() as tmp:
            project_path = Path(tmp) / "UT-BACKFILL-RAW"
            subject_dir = project_path / "subjects" / "S001"
            report_dir = subject_dir / "llm" / "baseline_randomization"
            report_dir.mkdir(parents=True)
            save_subject_info(
                subject_dir,
                SubjectInfo(
                    subject_id="S001",
                    project_code="UT-BACKFILL-RAW",
                    status=SubjectStatus.REVIEWED.value,
                    overall_verdict="investigator",
                ),
            )
            (report_dir / "review_raw.md").write_text(
                """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-15 | 有吸毒史、药物滥用史，或筛选前3个月内大量饮酒 | 排除 | ❌不通过 | 触发判断：已触发。病历记载筛选前3个月内确认有大量饮酒。 |

### 总结论
判定结果：fail
""",
                encoding="utf-8",
            )
            (report_dir / "review_report.md").write_text(
                """## 审核结论

**🟡 需研究者判定**  
旧后处理误把EX-15大量饮酒当成研究者复合条件。

---

## 逐条审核结果

| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-15 | 有吸毒史、药物滥用史，或筛选前3个月内大量饮酒 | 排除 | 🟡 需研究者判定 | 疾病存在、异常存在或用药存在本身不等于完整排除触发。 |
""",
                encoding="utf-8",
            )

            result = backfill_project_reports(
                project_path,
                project_code="UT-BACKFILL-RAW",
                phase="baseline_randomization",
            )

            info = load_subject_info(subject_dir)
            rewritten = (report_dir / "review_report.md").read_text(encoding="utf-8")
            self.assertEqual(result["subjects"][0]["new_overall"], "fail")
            self.assertEqual(info.overall_verdict, "fail")
            self.assertIn("**❌ 不可入组**", rewritten)
            self.assertNotIn("旧后处理误把", rewritten)

    def test_multistage_project_code_uses_selected_stage_as_identity(self):
        workflow = {"study_stages": ["Ⅱ期", "Ⅲ期"], "requires_study_stage_selection": True}
        self.assertEqual(
            projects_router._project_code_base_for_stage(
                {"project_code": "MG-K10-SAR", "study_stage": "Ⅲ期"},
                workflow,
            ),
            "MG-K10-SAR-III",
        )
        self.assertEqual(
            projects_router._project_code_base_for_stage(
                {"project_code": "MG-K10-SAR", "study_stage": "Ⅱ期"},
                workflow,
            ),
            "MG-K10-SAR-II",
        )
        self.assertEqual(
            projects_router._project_code_base_for_stage(
                {"project_code": "MG-K10-SAR", "study_stage": "Ⅲ期"},
                {"protocol_study_stages": ["Ⅱ期", "Ⅲ期"], "study_stages": ["Ⅲ期"]},
            ),
            "MG-K10-SAR-III",
        )

    def test_load_review_workflow_scopes_stage_from_project_config(self):
        with tempfile.TemporaryDirectory() as tmp:
            pd = Path(tmp)
            (pd / "config.json").write_text(
                projects_router.json.dumps({"project_code": "UT", "study_stage": "Ⅲ期"}, ensure_ascii=False),
                encoding="utf-8",
            )
            (pd / "review_phases.json").write_text(
                projects_router.json.dumps(
                    {
                        "study_stages": ["Ⅱ期", "Ⅲ期"],
                        "requires_study_stage_selection": True,
                        "review_phases": [{"phase_id": "screening_run_in", "name": "筛选期"}],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            workflow = load_review_workflow(pd)

        self.assertEqual(workflow["study_stage"], "Ⅲ期")
        self.assertEqual(workflow["study_stages"], ["Ⅲ期"])
        self.assertFalse(workflow["requires_study_stage_selection"])

    def test_local_auth_allows_first_registration_and_passwordless_login(self):
        from app.router import auth as auth_router

        users_path = auth_router._users_path()
        backup = users_path.read_text(encoding="utf-8") if users_path.exists() else None
        if users_path.exists():
            users_path.unlink()
        try:
            with TestClient(app) as client:
                status = client.get("/api/auth/status")
                self.assertEqual(status.status_code, 200, status.text)
                self.assertFalse(status.json()["has_users"])
                self.assertTrue(status.json()["admin_available"])

                unauth_projects = client.get("/api/projects")
                self.assertEqual(unauth_projects.status_code, 401)

                admin_register = client.post("/api/auth/register", json={"username": "admin", "password": ""})
                self.assertEqual(admin_register.status_code, 409)

                admin = client.post("/api/auth/login", json={"username": "admin", "password": "20121116"})
                self.assertEqual(admin.status_code, 200, admin.text)
                self.assertEqual(admin.json()["role"], "admin")
                self.assertFalse(admin.json()["requires_help"])
                second_admin = client.post("/api/auth/login", json={"username": "admin", "password": "20121116"})
                self.assertEqual(second_admin.status_code, 200, second_admin.text)
                self.assertNotEqual(admin.json()["token"], second_admin.json()["token"])

                wrong_admin = client.post("/api/auth/login", json={"username": "admin", "password": "wrong"})
                self.assertEqual(wrong_admin.status_code, 401)

                reg = client.post("/api/auth/register", json={"username": "__unit_auth__", "password": ""})
                self.assertEqual(reg.status_code, 201, reg.text)
                self.assertTrue(reg.json()["first_user"])

                login = client.post("/api/auth/login", json={"username": "__unit_auth__", "password": ""})
                self.assertEqual(login.status_code, 200, login.text)
                self.assertEqual(login.json()["username"], "__unit_auth__")
                self.assertTrue(login.json()["requires_help"])

                headers = {
                    "X-Enrollment-User": login.json()["username"],
                    "X-Enrollment-Token": login.json()["token"],
                }
                seen = client.post("/api/auth/help-seen", json={"username": "__unit_auth__"}, headers=headers)
                self.assertEqual(seen.status_code, 200, seen.text)

                second_login = client.post("/api/auth/login", json={"username": "__unit_auth__", "password": ""})
                self.assertEqual(second_login.status_code, 200, second_login.text)
                self.assertFalse(second_login.json()["requires_help"])
        finally:
            if backup is None:
                users_path.unlink(missing_ok=True)
            else:
                users_path.write_text(backup, encoding="utf-8")

    def test_projects_are_globally_readable_but_writable_only_by_owner_or_admin(self):
        from app.router import auth as auth_router
        from app.shared import projects

        codes = ["UT-AUTH-ALICE", "UT-AUTH-BOB"]
        users_path = auth_router._users_path()
        backup = users_path.read_text(encoding="utf-8") if users_path.exists() else None
        if users_path.exists():
            users_path.unlink()
        for code in codes:
            client_pd = ROOT / "projects" / code
            if client_pd.exists():
                shutil.rmtree(client_pd)
            projects.pop(code, None)

        try:
            with TestClient(app) as client:
                self.assertEqual(client.post("/api/auth/register", json={"username": "alice", "password": ""}).status_code, 201)
                self.assertEqual(client.post("/api/auth/register", json={"username": "bob", "password": ""}).status_code, 201)
                alice = login_headers(client, "alice")
                bob = login_headers(client, "bob")
                admin = admin_headers(client)

                a_create = client.post(
                    "/api/projects",
                    json={"project_code": codes[0], "protocol_id": "A-001", "name": "Alice Project"},
                    headers=alice,
                )
                self.assertEqual(a_create.status_code, 201, a_create.text)
                self.assertEqual(a_create.json()["owner_username"], "alice")

                b_create = client.post(
                    "/api/projects",
                    json={"project_code": codes[1], "protocol_id": "B-001", "name": "Bob Project"},
                    headers=bob,
                )
                self.assertEqual(b_create.status_code, 201, b_create.text)
                self.assertEqual(b_create.json()["owner_username"], "bob")

                alice_codes = {p["project_code"] for p in client.get("/api/projects", headers=alice).json()}
                bob_codes = {p["project_code"] for p in client.get("/api/projects", headers=bob).json()}
                admin_codes = {p["project_code"] for p in client.get("/api/projects", headers=admin).json()}

                self.assertIn(codes[0], alice_codes)
                self.assertIn(codes[1], alice_codes)
                self.assertIn(codes[1], bob_codes)
                self.assertIn(codes[0], bob_codes)
                self.assertTrue(set(codes).issubset(admin_codes))

                bob_read_alice = client.get(f"/api/projects/{codes[0]}", headers=bob)
                self.assertEqual(bob_read_alice.status_code, 200, bob_read_alice.text)
                self.assertFalse(bob_read_alice.json()["can_modify"])
                self.assertEqual(client.delete(f"/api/projects/{codes[0]}", headers=bob).status_code, 403)
                self.assertEqual(
                    client.patch(f"/api/projects/{codes[0]}", json={"name": "Bob edit"}, headers=bob).status_code,
                    403,
                )
                self.assertEqual(
                    client.post(f"/api/projects/{codes[0]}/rules", json={"rules": "#### IN-01"}, headers=bob).status_code,
                    403,
                )
                self.assertEqual(client.delete(f"/api/projects/{codes[0]}", headers=alice).status_code, 200)
                self.assertEqual(client.delete(f"/api/projects/{codes[1]}", headers=admin).status_code, 200)
        finally:
            for code in codes:
                pd = ROOT / "projects" / code
                if pd.exists():
                    shutil.rmtree(pd)
                projects.pop(code, None)
            if backup is None:
                users_path.unlink(missing_ok=True)
            else:
                users_path.write_text(backup, encoding="utf-8")

    def test_subject_ownership_and_global_centers_are_shared_readonly(self):
        from app.router import auth as auth_router
        from app.shared import projects

        project_code = "UT-SHARED-READ"
        sid = "SUBJ-B01"
        users_path = auth_router._users_path()
        centers_path = ROOT / "projects" / "_system" / "centers.json"
        users_backup = users_path.read_text(encoding="utf-8") if users_path.exists() else None
        centers_backup = centers_path.read_text(encoding="utf-8") if centers_path.exists() else None
        if users_path.exists():
            users_path.unlink()
        if centers_path.exists():
            centers_path.unlink()
        pd = ROOT / "projects" / project_code
        if pd.exists():
            shutil.rmtree(pd)
        projects.pop(project_code, None)

        try:
            with TestClient(app) as client:
                self.assertEqual(client.post("/api/auth/register", json={"username": "owner_a", "password": ""}).status_code, 201)
                self.assertEqual(client.post("/api/auth/register", json={"username": "uploader_b", "password": ""}).status_code, 201)
                owner = login_headers(client, "owner_a")
                uploader = login_headers(client, "uploader_b")
                admin = admin_headers(client)

                create_project = client.post(
                    "/api/projects",
                    json={"project_code": project_code, "protocol_id": "UT-001", "name": "共享只读项目"},
                    headers=owner,
                )
                self.assertEqual(create_project.status_code, 201, create_project.text)

                create_subject = client.post(
                    f"/api/projects/{project_code}/subjects",
                    json={"subject_id": sid, "center_code": "8", "center_name": "第八中心医院"},
                    headers=uploader,
                )
                self.assertEqual(create_subject.status_code, 201, create_subject.text)
                self.assertEqual(create_subject.json()["center_code"], "08")
                self.assertEqual(create_subject.json()["owner_username"], "uploader_b")

                owner_subjects = client.get(f"/api/projects/{project_code}/subjects", headers=owner)
                self.assertEqual(owner_subjects.status_code, 200, owner_subjects.text)
                owner_row = next(x for x in owner_subjects.json() if x["subject_id"] == sid)
                self.assertFalse(owner_row["can_modify"])

                blocked_upload = client.post(
                    f"/api/projects/{project_code}/subjects/{sid}/upload",
                    files=[
                        ("files", ("owner.txt", io.BytesIO(b"owner"), "text/plain")),
                        ("categories", (None, "screening_record")),
                    ],
                    headers=owner,
                )
                self.assertEqual(blocked_upload.status_code, 403, blocked_upload.text)

                allowed_upload = client.post(
                    f"/api/projects/{project_code}/subjects/{sid}/upload",
                    files=[
                        ("files", ("uploader.txt", io.BytesIO(b"uploader"), "text/plain")),
                        ("categories", (None, "screening_record")),
                    ],
                    headers=uploader,
                )
                self.assertEqual(allowed_upload.status_code, 200, allowed_upload.text)

                owner_centers = client.get(f"/api/projects/{project_code}/centers", headers=owner)
                self.assertEqual(owner_centers.status_code, 200, owner_centers.text)
                center = next(c for c in owner_centers.json()["centers"] if c["center_code"] == "08")
                self.assertEqual(center["center_name"], "第八中心医院")
                self.assertFalse(center["can_modify"])

                blocked_center_edit = client.put(
                    f"/api/projects/{project_code}/centers",
                    json={"centers": [{"center_code": "08", "center_name": "改名医院"}]},
                    headers=owner,
                )
                self.assertEqual(blocked_center_edit.status_code, 403, blocked_center_edit.text)

                self.assertEqual(client.delete(f"/api/projects/{project_code}/subjects/{sid}", headers=owner).status_code, 403)
                self.assertEqual(client.delete(f"/api/projects/{project_code}/subjects/{sid}", headers=uploader).status_code, 200)
                self.assertEqual(client.delete(f"/api/projects/{project_code}", headers=admin).status_code, 200)
        finally:
            if pd.exists():
                shutil.rmtree(pd)
            projects.pop(project_code, None)
            if users_backup is None:
                users_path.unlink(missing_ok=True)
            else:
                users_path.write_text(users_backup, encoding="utf-8")
            if centers_backup is None:
                centers_path.unlink(missing_ok=True)
            else:
                centers_path.parent.mkdir(parents=True, exist_ok=True)
                centers_path.write_text(centers_backup, encoding="utf-8")

    def test_protocol_draft_deconstruct_endpoint_keeps_metadata_and_rules(self):
        protocol_text = (
            "方案编号：UT-DRAFT-001\n版本：V1.0\n2026年1月2日\n"
            "UT-DRAFT 临床研究方案\n入选标准：1）年龄18至75岁。排除标准：1）严重肝肾异常。"
            * 8
        )
        with TestClient(app) as client, patch(
            "app.llm.client.deconstruct_chat",
            new=AsyncMock(return_value="#### IN-01 年龄\n- **判断点**：年龄\n"),
        ):
            headers = admin_headers(client)
            resp = client.post(
                "/api/projects/deconstruct-draft",
                data={"feedback": "请细化年龄标准", "current_rules": "#### IN-01 旧草稿"},
                files={"file": ("UT-DRAFT-001_临床研究方案_V1.0_20260102.txt", io.BytesIO(protocol_text.encode("utf-8")), "text/plain")},
                headers=headers,
            )

        self.assertEqual(resp.status_code, 200, resp.text)
        data = resp.json()
        self.assertEqual(data["metadata"]["project_code"], "UT-DRAFT")
        self.assertEqual(data["metadata"]["protocol_id"], "UT-DRAFT-001")
        self.assertEqual(data["metadata"]["protocol_version"], "V1.0")
        self.assertEqual(data["metadata"]["protocol_date"], "2026-01-02")
        self.assertIn("IN-01", data["rules"])

    @unittest.skipUnless(MGK10_PROTOCOL.exists(), "MG-K10-SAR protocol fixture not available")
    def test_multistage_protocol_draft_requires_study_stage_before_deconstruction(self):
        with TestClient(app) as client:
            headers = admin_headers(client)
            with MGK10_PROTOCOL.open("rb") as fh:
                resp = client.post(
                    "/api/projects/deconstruct-draft",
                    files={
                        "file": (
                            MGK10_PROTOCOL.name,
                            io.BytesIO(fh.read()),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    },
                    headers=headers,
                )

        self.assertEqual(resp.status_code, 409, resp.text)
        detail = resp.json()["detail"]
        self.assertEqual(detail["code"], "study_stage_required")
        self.assertEqual(detail["study_stages"], ["Ⅱ期", "Ⅲ期"])

    def test_save_protocol_draft_creates_project_rules_and_workflow(self):
        project_code = "UT-DRAFT-SAVE"
        metadata = {
            "project_code": project_code,
            "protocol_id": "UT-DRAFT-SAVE-001",
            "name": "UT 草稿保存测试方案",
            "protocol_version": "V1.0",
            "protocol_date": "2026-01-02",
            "protocol_source_filename": "UT-DRAFT-SAVE-001.txt",
        }
        workflow = {
            "review_phases": [
                {"phase_id": "screening_run_in", "name": "筛选期", "required_items": ["签署知情同意书"]}
            ]
        }
        rules = "#### IN-01 年龄\n- **判断点**：年龄\n"

        with TestClient(app) as client:
            headers = admin_headers(client)
            client.delete(f"/api/projects/{project_code}", headers=headers)
            try:
                resp = client.post(
                    "/api/projects/save-protocol-draft",
                    data={
                        "rules": rules,
                        "metadata_json": projects_router.json.dumps(metadata, ensure_ascii=False),
                        "workflow_json": projects_router.json.dumps(workflow, ensure_ascii=False),
                    },
                    files={"file": ("UT-DRAFT-SAVE-001.txt", io.BytesIO(b"protocol"), "text/plain")},
                    headers=headers,
                )
                self.assertEqual(resp.status_code, 201, resp.text)
                data = resp.json()
                self.assertEqual(data["project_code"], project_code)

                pd = ROOT / "projects" / project_code
                self.assertEqual((pd / "criteria_rules.md").read_text(encoding="utf-8"), rules)
                saved_workflow = projects_router.json.loads((pd / "review_phases.json").read_text(encoding="utf-8"))
                self.assertEqual(saved_workflow["review_phases"][0]["phase_id"], "screening_run_in")
                cfg = projects_router.json.loads((pd / "config.json").read_text(encoding="utf-8"))
                self.assertEqual(cfg["protocol_id"], "UT-DRAFT-SAVE-001")
                self.assertEqual(cfg["protocol_source_filename"], "UT-DRAFT-SAVE-001.txt")
            finally:
                client.delete(f"/api/projects/{project_code}", headers=headers)

    def test_save_protocol_draft_splits_multistage_protocol_into_stage_project(self):
        base_code = "UT-DRAFT-PHASE"
        expected_code = f"{base_code}-III"
        metadata = {
            "project_code": base_code,
            "protocol_id": "UT-DRAFT-PHASE-001",
            "name": "UT 分期草稿保存测试方案",
            "study_stage": "Ⅲ期",
        }
        workflow = {
            "study_stages": ["Ⅱ期", "Ⅲ期"],
            "requires_study_stage_selection": True,
            "review_phases": [
                {"phase_id": "screening_run_in", "name": "筛选期", "required_items": ["签署知情同意书"]}
            ],
        }
        rules = "#### IN-01 年龄\n- **判断点**：年龄\n"

        with TestClient(app) as client:
            headers = admin_headers(client)
            client.delete(f"/api/projects/{expected_code}", headers=headers)
            try:
                resp = client.post(
                    "/api/projects/save-protocol-draft",
                    data={
                        "rules": rules,
                        "metadata_json": projects_router.json.dumps(metadata, ensure_ascii=False),
                        "workflow_json": projects_router.json.dumps(workflow, ensure_ascii=False),
                    },
                    headers=headers,
                )
                self.assertEqual(resp.status_code, 201, resp.text)
                data = resp.json()
                self.assertEqual(data["project_code"], expected_code)

                pd = ROOT / "projects" / expected_code
                cfg = projects_router.json.loads((pd / "config.json").read_text(encoding="utf-8"))
                saved_workflow = projects_router.json.loads((pd / "review_phases.json").read_text(encoding="utf-8"))
                self.assertEqual(cfg["study_stage"], "Ⅲ期")
                self.assertEqual(saved_workflow["study_stage"], "Ⅲ期")
                self.assertEqual(saved_workflow["study_stages"], ["Ⅲ期"])
                self.assertFalse(saved_workflow["requires_study_stage_selection"])
            finally:
                client.delete(f"/api/projects/{expected_code}", headers=headers)


class SubjectUploadTests(unittest.TestCase):
    def test_all_folder_discovery_merges_subject_roots_and_ignores_nested_photo_dirs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            primary = root / "张三" / "SA01001入组"
            primary.mkdir(parents=True)
            (primary / "筛选病历.pdf").write_bytes(b"primary")
            photo_nested = primary / "SA01001筛选期照片"
            photo_nested.mkdir()
            (photo_nested / "SA01001照片.jpg").write_bytes(b"photo")

            duplicate_root = root / "李四" / "0519-SA01001"
            duplicate_root.mkdir(parents=True)
            (duplicate_root / "补充检查.pdf").write_bytes(b"supplement")

            screen_fail = root / "王五" / "SA01002筛败"
            screen_fail.mkdir(parents=True)
            (screen_fail / "筛败说明.pdf").write_bytes(b"failed")

            all_folders = discover_subject_folders(root, folder_keyword=None, study_stage="Ⅱ期")
            by_sid = {item.subject_id: item for item in all_folders}

            self.assertEqual([item.subject_id for item in all_folders], ["SA01001", "SA01002"])
            self.assertEqual(by_sid["SA01001"].folder, primary)
            self.assertEqual(set(by_sid["SA01001"].folders), {primary, duplicate_root})
            self.assertNotIn(photo_nested, set(by_sid["SA01001"].folders))

            retained, skipped = collect_review_files_from_folders(by_sid["SA01001"].folders, "SA01001")
            self.assertEqual([item.relative_path for item in retained], ["SA01001入组/筛选病历.pdf", "0519-SA01001/补充检查.pdf"])
            self.assertTrue(any(item["reason"] == "photo/image" for item in skipped))

            screen_failed_only = discover_subject_folders(root, folder_keyword="筛败", study_stage="Ⅱ期")
            self.assertEqual([item.subject_id for item in screen_failed_only], ["SA01002"])

    @unittest.skipUnless(D001_SOURCE_ROOT.exists(), "D001 source folder fixture not available")
    def test_d001_screen_failed_batch_uses_project_stage_not_filename_phase_markers(self):
        folders = discover_subject_folders(D001_SOURCE_ROOT, folder_keyword="筛败", study_stage="Ⅱ期")
        by_sid = {item.subject_id: item for item in folders}

        self.assertEqual(len(folders), 22)
        self.assertIn("SA07025", by_sid)
        self.assertIn("SA07030", by_sid)
        self.assertTrue(all(item.study_stage == "Ⅱ期" for item in folders))

        files_07025, skipped_07025 = collect_review_files(by_sid["SA07025"].folder, "SA07025")
        self.assertTrue(any("Ⅲ期" in item.relative_path for item in files_07025))
        self.assertFalse(any(item.relative_path.lower().endswith((".jpg", ".jpeg", ".png")) for item in files_07025))
        self.assertTrue(any("照片" in item["path"] and item["reason"] == "photo/image" for item in skipped_07025))

        files_07012, skipped_07012 = collect_review_files(by_sid["SA07012"].folder, "SA07012")
        self.assertFalse(any(".DS_Store" in item.relative_path for item in files_07012))
        self.assertTrue(any(item["path"] == ".DS_Store" for item in skipped_07012))

    def test_d001_batch_report_reads_postprocessed_report_before_raw_and_compares_manual_pass(self):
        from scripts import d001_phase2_batch_review as batch

        with tempfile.TemporaryDirectory() as tmp:
            project_dir = Path(tmp)
            report_dir = project_dir / "subjects" / "SA99999" / "llm" / "screening_run_in"
            report_dir.mkdir(parents=True)
            (report_dir / "review_raw.md").write_text(
                """### 逐条审核结果
| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-11 | 重大疾病 | 排除 | ❌不通过 | 存在系统性疾病。 |

### 总结论
判定结果：fail
""",
                encoding="utf-8",
            )
            (report_dir / "review_report.md").write_text(
                """## 审核结论

**✅ 通过**  
未见可直接确认的不可入组项。

---

## 逐条审核结果

| 规则ID | 规则名称 | 类型 | 判定结果 | 推理依据 |
|--------|----------|------|----------|----------|
| EX-11 | 重大疾病 | 排除 | ✅通过 | 未见“存在重大疾病”且“研究者明确判断不具备临床研究条件”同时满足。 |
""",
                encoding="utf-8",
            )

            with patch.object(batch, "PROJECT_DIR", project_dir):
                report = batch.read_subject_report("SA99999")

        self.assertEqual(report["overall"], "pass")
        self.assertEqual(report["rules"][0]["verdict"], "pass")

        comparison = batch.compare_manual(
            {"rules": [{"rule_id": "EX-20", "verdict": "investigator"}]},
            {"ieyn": "是", "manual_item": "", "manual_comment": ""},
        )
        self.assertEqual(comparison["status"], "存在差异：人工通过但系统仍有关注条目")

    def test_upload_merges_file_categories_for_incremental_batches(self):
        sid = "__upload_merge_test__"
        with TestClient(app) as client:
            headers = admin_headers(client)
            with temporary_api_project(client, headers, "UT-UPLOAD-MERGE") as project_code:
                create_resp = client.post(
                    f"/api/projects/{project_code}/subjects",
                    json={"subject_id": sid},
                    headers=headers,
                )
                self.assertEqual(create_resp.status_code, 201, create_resp.text)
                first = client.post(
                    f"/api/projects/{project_code}/subjects/{sid}/upload",
                    files=[
                        ("files", ("筛选期病历.pdf", io.BytesIO(b"screening"), "application/pdf")),
                        ("categories", (None, "screening_record")),
                    ],
                    headers=headers,
                )
                self.assertEqual(first.status_code, 200, first.text)

                second = client.post(
                    f"/api/projects/{project_code}/subjects/{sid}/upload",
                    files=[
                        ("files", ("基线血常规.pdf", io.BytesIO(b"baseline"), "application/pdf")),
                        ("categories", (None, "screening_lab")),
                    ],
                    headers=headers,
                )
                self.assertEqual(second.status_code, 200, second.text)

                cat_path = ROOT / "projects" / project_code / "subjects" / sid / "file_categories.json"
                cat_map = projects_router.json.loads(cat_path.read_text(encoding="utf-8"))
                self.assertEqual(cat_map["筛选期病历.pdf"], "screening_record")
                self.assertEqual(cat_map["基线血常规.pdf"], "screening_lab")

    def test_incremental_upload_preserves_same_named_files_and_marks_review_stale(self):
        sid = "__upload_same_name_test__"
        with TestClient(app) as client:
            headers = admin_headers(client)
            with temporary_api_project(client, headers, "UT-UPLOAD-SAME-NAME") as project_code:
                create_resp = client.post(
                    f"/api/projects/{project_code}/subjects",
                    json={"subject_id": sid},
                    headers=headers,
                )
                self.assertEqual(create_resp.status_code, 201, create_resp.text)
                subject_path = ROOT / "projects" / project_code / "subjects" / sid
                info = load_subject_info(subject_path)
                info.status = SubjectStatus.REVIEWED.value
                info.overall_verdict = "pass"
                save_subject_info(subject_path, info)

                first = client.post(
                    f"/api/projects/{project_code}/subjects/{sid}/upload",
                    files=[
                        ("files", ("同名报告.pdf", io.BytesIO(b"old"), "application/pdf")),
                        ("categories", (None, "screening_record")),
                    ],
                    headers=headers,
                )
                self.assertEqual(first.status_code, 200, first.text)
                second = client.post(
                    f"/api/projects/{project_code}/subjects/{sid}/upload",
                    files=[
                        ("files", ("同名报告.pdf", io.BytesIO(b"new"), "application/pdf")),
                        ("categories", (None, "screening_lab")),
                    ],
                    headers=headers,
                )
                self.assertEqual(second.status_code, 200, second.text)

                raw_names = sorted(p.name for p in (subject_path / "raw").iterdir() if p.is_file())
                self.assertEqual(raw_names, ["同名报告.pdf", "同名报告_2.pdf"])
                self.assertEqual((subject_path / "raw" / "同名报告.pdf").read_bytes(), b"old")
                self.assertEqual((subject_path / "raw" / "同名报告_2.pdf").read_bytes(), b"new")
                cat_map = projects_router.json.loads((subject_path / "file_categories.json").read_text(encoding="utf-8"))
                self.assertEqual(cat_map["同名报告.pdf"], "screening_record")
                self.assertEqual(cat_map["同名报告_2.pdf"], "screening_lab")
                updated = load_subject_info(subject_path)
                self.assertEqual(updated.status, SubjectStatus.PENDING.value)
                self.assertEqual(updated.overall_verdict, "")

    def test_subject_icf_date_can_be_manually_saved_and_not_overwritten_by_cache_backfill(self):
        sid = "__icf_date_test__"
        with TestClient(app) as client:
            headers = admin_headers(client)
            with temporary_api_project(client, headers, "UT-ICF-DATE") as project_code:
                create_resp = client.post(
                    f"/api/projects/{project_code}/subjects",
                    json={"subject_id": sid},
                    headers=headers,
                )
                self.assertEqual(create_resp.status_code, 201, create_resp.text)
                patch_resp = client.patch(
                    f"/api/projects/{project_code}/subjects/{sid}",
                    json={"icf_date": "2026-01-02"},
                    headers=headers,
                )
                self.assertEqual(patch_resp.status_code, 200, patch_resp.text)
                self.assertEqual(patch_resp.json()["icf_date"], "2026-01-02")
                self.assertTrue(patch_resp.json()["icf_date_manual"])

                subject_path = ROOT / "projects" / project_code / "subjects" / sid
                cache_dir = subject_path / "cache"
                cache_dir.mkdir(exist_ok=True)
                (cache_dir / "知情同意书_p1.md").write_text(
                    "受试者于2026年01月03日签署知情同意书。",
                    encoding="utf-8",
                )
                changed = backfill_subject_icf_date(subject_path)
                self.assertFalse(changed)
                self.assertEqual(load_subject_info(subject_path).icf_date, "2026-01-02")

    def test_extract_icf_date_from_cache_prefers_signed_informed_consent_context(self):
        with tempfile.TemporaryDirectory() as tmp:
            cache_dir = Path(tmp)
            (cache_dir / "other.md").write_text("筛选日期：2026-01-01", encoding="utf-8")
            (cache_dir / "screening_record.md").write_text(
                "筛选病历：受试者签署知情同意书日期为2026年01月05日，随后完成筛选。",
                encoding="utf-8",
            )
            self.assertEqual(extract_icf_date_from_cache(cache_dir), "2026-01-05")

    def test_extract_icf_date_from_cache_ignores_icf_version_date_in_same_sentence(self):
        with tempfile.TemporaryDirectory() as tmp:
            cache_dir = Path(tmp)
            record_dir = cache_dir / "31001-病历"
            record_dir.mkdir()
            (record_dir / "p2.md").write_text(
                "受试者于2025年8月8日09:43和研究者于2025年8月8日09:44签署了知情同意书"
                "（版本号：Master V2.0-site31v01，版本日期：2025年7月02日），筛选号：S31001。",
                encoding="utf-8",
            )
            categories = {"31001-病历.pdf": "screening_record"}

            self.assertEqual(
                extract_icf_date_from_cache(cache_dir, subject_id="31001", categories=categories),
                "2025-08-08",
            )


class SubjectResetTests(unittest.TestCase):
    def test_clear_subject_outputs_removes_phase_specific_bundles(self):
        with tempfile.TemporaryDirectory() as tmp:
            sd = Path(tmp)
            (sd / "cache").mkdir()
            (sd / "llm").mkdir()
            (sd / "evidence_bundle.md").write_text("legacy", encoding="utf-8")
            (sd / "evidence_bundle__screening_run_in.md").write_text("screening", encoding="utf-8")
            (sd / "evidence_bundle__baseline_randomization.md").write_text("baseline", encoding="utf-8")

            info = SubjectInfo(
                subject_id="TMP",
                project_code="MG-K10-SAR",
                status=SubjectStatus.REVIEWED.value,
                overall_verdict="符合",
            )
            save_subject_info(sd, info)

            clear_subject_outputs(sd, info)

            self.assertFalse((sd / "cache").exists())
            self.assertFalse((sd / "llm").exists())
            self.assertEqual(list(sd.glob("evidence_bundle*.md")), [])
            updated = projects_router.json.loads((sd / "info.json").read_text(encoding="utf-8"))
            self.assertEqual(updated["status"], SubjectStatus.PENDING.value)
            self.assertEqual(updated["overall_verdict"], "")


class SystemReviewRegressionTests(unittest.TestCase):
    def test_storage_id_rejects_path_traversal(self):
        from fastapi import HTTPException
        from app.shared import validate_storage_id

        self.assertEqual(validate_storage_id("SA07002", "受试者ID"), "SA07002")
        self.assertEqual(validate_storage_id("D001-02-II", "项目编号"), "D001-02-II")

        for bad in ["../../../etc/passwd", "..", "SA/07002", "SA\\07002", "SA 07002", ""]:
            with self.subTest(bad=bad):
                with self.assertRaises(HTTPException):
                    validate_storage_id(bad, "受试者ID")

    def test_upload_filename_validation_rejects_traversal_and_unsupported_extensions(self):
        from fastapi import HTTPException
        from app.router.subjects import _safe_upload_destination

        with tempfile.TemporaryDirectory() as tmp:
            raw_dir = Path(tmp)
            self.assertEqual(_safe_upload_destination(raw_dir, "筛选病历.pdf").name, "筛选病历.pdf")
            for bad in ["..", ".", "evil.sh", "archive.zip", "bad/name.pdf", "bad\\name.pdf"]:
                with self.subTest(bad=bad):
                    with self.assertRaises(HTTPException):
                        _safe_upload_destination(raw_dir, bad)

    def test_icf_date_extraction_ignores_enrollment_comm_and_other_subject_ids(self):
        with tempfile.TemporaryDirectory() as tmp:
            cache_dir = Path(tmp)
            comm = cache_dir / "SA07002合格性讨论表"
            record = cache_dir / "SA07002筛选期病历"
            comm.mkdir()
            record.mkdir()
            (comm / "p1.md").write_text(
                "受试者SA07202于2024年03月05日签署知情同意书。",
                encoding="utf-8",
            )
            (record / "p1.md").write_text(
                "筛选病历：受试者SA07002于2026年03月05日09时06分57秒签署了知情同意书。",
                encoding="utf-8",
            )
            categories = {
                "SA07002合格性讨论表.pdf": "enrollment_comm",
                "SA07002筛选期病历.pdf": "screening_record",
            }

            self.assertEqual(
                extract_icf_date_from_cache(cache_dir, subject_id="SA07002", categories=categories),
                "2026-03-05",
            )

    def test_ocr_hallucination_detection_deduplicates_repeated_lines(self):
        from app.pipeline.ocr import detect_ocr_hallucination

        text = "\n".join(["患者已开始进行免疫治疗并持续症状加重"] * 30 + ["真实检验结果：ALT 22 U/L"])
        flagged, cleaned, warning = detect_ocr_hallucination(text)

        self.assertTrue(flagged)
        self.assertIn("OCR质量警告", warning)
        self.assertLess(cleaned.count("患者已开始进行免疫治疗并持续症状加重"), 3)
        self.assertIn("真实检验结果：ALT 22 U/L", cleaned)

    def test_evidence_bundle_preserves_source_extension_and_warns_subject_mismatch(self):
        with tempfile.TemporaryDirectory() as tmp:
            sd = Path(tmp) / "SA07002"
            doc_dir = sd / "cache" / "SA07002合格性讨论表"
            doc_dir.mkdir(parents=True)
            (doc_dir / "p1.md").write_text("受试者SA07202于2024年03月05日签署知情同意书。", encoding="utf-8")
            (sd / "file_categories.json").write_text(
                json.dumps({"SA07002合格性讨论表.docx": "enrollment_comm"}, ensure_ascii=False),
                encoding="utf-8",
            )

            bundle = build_evidence_bundle(sd, "D001-02-II")

        self.assertIn("来源文件：SA07002合格性讨论表.docx", bundle)
        self.assertIn("资料一致性警告", bundle)
        self.assertIn("与当前文件夹 SA07002 不一致", bundle)

    def test_processing_lock_rejects_duplicate_subject_workflow(self):
        from fastapi import HTTPException
        from app.processing_locks import acquire_subject_processing_lock, release_subject_processing_lock

        async def scenario():
            key = await acquire_subject_processing_lock("D001-02-II", "SA07002")
            try:
                with self.assertRaises(HTTPException) as ctx:
                    await acquire_subject_processing_lock("D001-02-II", "SA07002")
                self.assertEqual(ctx.exception.status_code, 409)
            finally:
                release_subject_processing_lock(key)
            key2 = await acquire_subject_processing_lock("D001-02-II", "SA07002")
            release_subject_processing_lock(key2)

        asyncio.run(scenario())

    def test_project_list_stats_separate_insufficient_and_investigator(self):
        project_code = "__stats_review_test__"
        with TestClient(app) as client:
            headers = admin_headers(client)
            client.delete(f"/api/projects/{project_code}", headers=headers)
            created = client.post(
                "/api/projects",
                json={"project_code": project_code, "protocol_id": "STAT-001", "name": "stats"},
                headers=headers,
            )
            self.assertEqual(created.status_code, 201, created.text)
            try:
                project_path = ROOT / "projects" / project_code
                subjects_dir = project_path / "subjects"
                for sid, verdict in [("S1", "pass"), ("S2", "insufficient"), ("S3", "investigator"), ("S4", "")]:
                    sd = subjects_dir / sid
                    sd.mkdir(parents=True)
                    save_subject_info(sd, SubjectInfo(subject_id=sid, project_code=project_code, overall_verdict=verdict))
                listed = client.get("/api/projects", headers=headers)
                self.assertEqual(listed.status_code, 200, listed.text)
                item = next(p for p in listed.json() if p["project_code"] == project_code)
                self.assertEqual(item["stats"]["pass"], 1)
                self.assertEqual(item["stats"]["insufficient"], 1)
                self.assertEqual(item["stats"]["investigator"], 1)
                self.assertEqual(item["stats"]["pending"], 1)
                self.assertEqual(item["stats"]["pending_total"], 3)
            finally:
                client.delete(f"/api/projects/{project_code}", headers=headers)

    def test_project_delete_keeps_system_audit_record(self):
        project_code = "__audit_delete_test__"
        system_ledger = ROOT / "projects" / "_system" / "audit_ledger.jsonl"
        before = system_ledger.read_text(encoding="utf-8") if system_ledger.exists() else ""

        with TestClient(app) as client:
            headers = admin_headers(client)
            client.delete(f"/api/projects/{project_code}", headers=headers)
            created = client.post(
                "/api/projects",
                json={"project_code": project_code, "protocol_id": "AUDIT-001", "name": "audit"},
                headers=headers,
            )
            self.assertEqual(created.status_code, 201, created.text)
            deleted = client.delete(f"/api/projects/{project_code}", headers=headers)
            self.assertEqual(deleted.status_code, 200, deleted.text)

        after = system_ledger.read_text(encoding="utf-8") if system_ledger.exists() else ""
        new_lines = after[len(before):].strip().splitlines()
        events = [json.loads(line) for line in new_lines if line.strip()]
        self.assertTrue(
            any(
                item.get("project_code") == project_code and item.get("event") == "delete_project"
                for item in events
            )
        )


if __name__ == "__main__":
    unittest.main()
