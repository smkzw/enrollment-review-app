#!/usr/bin/env python3
"""词汇中立的真实生产 harness 冒烟运行器（方案控制点两阶段生产链）。

用法：

    python scripts/run_protocol_control_smoke.py
    python scripts/run_protocol_control_smoke.py --out-dir DIR
    python scripts/run_protocol_control_smoke.py --fixture PATH
    python scripts/run_protocol_control_smoke.py --out-dir DIR --resume-job-id JOB_ID

行为与边界（与执行任务 `phase5-protocol-control-live-model-smoke-20260831` 一致）：

- 先正向核实"配置模型 == 服务实际加载模型"（/v1/models），再执行任何运行；
  身份缺失、歧义或不一致时失效关闭并输出中文诊断，绝不静默替换模型。
- 词汇中立合成协议走真实生产来源链：工作台登记 -> 结构提取 -> 渲染对齐 ->
  方案信息与研究期别识别 -> 用户边界确认 -> 冻结结构快照；运行器不伪造结构
  快照检查点，也不把手写协议文本直接交给深析 Agent。
- 来源任务在冻结快照完成后立即以显式非重试边界停止，不驱动官方 IN/EX 语义
  解构（与本冒烟无关的另一条语义路线）。
- 随后用生产 ``ProtocolControlJobService`` 建立持久控制任务，并用生产
  ``create_protocol_control_executor`` 经真实 ``JobRunner`` 驱动发现与深析。
- 默认只消费本运行器内置的词汇中立合成协议；拒绝 ``projects/`` 与真实数据
  目录下的任何外部方案文件。
- 不物化正式 ``ProtocolReviewControl`` 目录；结果恒为水合候选控制点包，
  运行记录显式携带 ``claims_complete=false``。

- ``--resume-job-id`` 只打开既有 ``out-dir/data_v2`` SQLite，恢复过期租约，继续未完成控制步骤；
  不重新读取原始方案文件，也不创建第二个来源或控制任务。

退出码：0 完成；2 模型身份失效关闭；3 运行失败；4 使用/边界错误。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import traceback
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from docx import Document  # noqa: E402

from app.agents.protocol_control_agent_transport import (  # noqa: E402
    ProtocolControlModelIdentityError,
    protocol_control_transport_from_environment,
)
from app.agents.protocol_control_deconstructor import (  # noqa: E402
    ProtocolControlAgentRunResult,
    ProtocolControlDiscoveryAgentRunResult,
)
from app.agents.protocol_control_discovery_transport import (  # noqa: E402
    protocol_control_discovery_transport_from_environment,
)
from app.domain.contracts.enums import StudyPhase  # noqa: E402
from app.domain.contracts.protocol_metadata import (  # noqa: E402
    ProtocolIdentityDecision,
)
from app.services.protocol_control_execution import (  # noqa: E402
    CANDIDATE_CONTROL_PACKAGE_RESULT_KIND,
    FORMAL_CATALOG_STATUS_NOT_MATERIALIZED,
    PROTOCOL_CONTROL_EXECUTION_JOB_TYPE,
    ProtocolControlExecutionError,
    ProtocolControlExecutorConfig,
    ProtocolControlJobService,
    create_protocol_control_executor,
)
from app.services.protocol_deconstruction_executor import (  # noqa: E402
    ProtocolDeconstructionExecutorConfig,
    create_protocol_deconstruction_executor,
)
from app.services.protocol_workbench_service import (  # noqa: E402
    PROTOCOL_DECONSTRUCTION_JOB_TYPE,
    ProtocolWorkbenchError,
    ProtocolWorkbenchService,
    STEP_FREEZE,
    STEP_GENERATE,
)
from app.storage.codecs import utc_now, verify_payload_sha256  # noqa: E402
from app.storage.config import resolve_data_paths  # noqa: E402
from app.storage.db import build_engine, build_session_factory  # noqa: E402
from app.storage.migrate import MigrationManager  # noqa: E402
from app.workflow.errors import StepFailure  # noqa: E402
from app.workflow.jobstore import JobStore  # noqa: E402
from app.workflow.runner import JobRunner  # noqa: E402
from app.workflow.recovery import recover_expired_jobs  # noqa: E402


SMOKE_STOP_ERROR_CODE = "SMOKE_SOURCE_STOP_AFTER_FREEZE"
SOURCE_STOP_STEP = STEP_FREEZE
RUNNER_LABEL = "protocol-control-smoke"
_DISCOVERY_STEP_PREFIX = "discovery_"

EXIT_OK = 0
EXIT_IDENTITY_FAIL_CLOSED = 2
EXIT_RUN_FAILED = 3
EXIT_USAGE_OR_GUARD = 4


class SmokeGuardError(RuntimeError):
    """使用方式或边界守卫不满足（退出码 4）。"""


# ---------------------------------------------------------------------------
# 词汇中立合成协议（内置唯一默认输入）
# ---------------------------------------------------------------------------

_SYNTHETIC_DOCX_TITLE = "II期合成冒烟方案（词汇中立）"

_SYNTHETIC_FLOW_TABLE: tuple[tuple[str, ...], ...] = (
    ("操作项目", "筛选期", "基线期", "治疗期"),
    ("登记与评估", "X", "", ""),
    ("研究培训", "", "X", ""),
    ("状态记录", "", "", "X"),
)

_SYNTHETIC_DOCX_SECTIONS: tuple[tuple[str, tuple[str, ...], tuple[str, ...]], ...] = (
    (
        "II期方案信息",
        (
            "方案编号：SYN-SMOKE-001",
            "方案版本：v1.0",
            "方案版本日期：2026年1月1日",
            "本文件为本地冒烟验证专用的合成示例，不代表任何真实研究。",
            "本方案内容适用于II期阶段。",
        ),
        (),
    ),
    (
        "II期研究设计",
        (
            "本研究为II期探索性研究，仅覆盖II期阶段。",
            "研究设计只用于验证通用生产链路，不描述任何真实医学内容。",
        ),
        (),
    ),
    (
        "II期入选标准",
        (),
        (
            "受试者已完成规定的登记流程。",
            "受试者能够按照研究计划参加研究相关活动。",
        ),
    ),
    (
        "II期排除标准",
        (),
        (
            "受试者未完成规定的登记流程。",
            "根据研究者的判断，受试者不适合参加本研究。",
        ),
    ),
    (
        "II期研究流程",
        (
            "本部分流程适用于II期阶段。",
            "筛选期内完成统一的登记与评估流程。",
            "基线前完成必需的研究培训并记录完成情况。",
            "研究过程中按计划记录研究执行状态。",
        ),
        (),
    ),
    (
        "II期研究管理要求",
        (
            "本部分管理要求适用于II期阶段。",
            "该受试者进入基线节点前，与其相关的上一阶段全部待确认事项必须同时满足已完成处理且处理依据已记录；任一项未满足，该受试者不得进入基线节点。",
            "研究者在启动任何研究操作前必须完成方案培训并记录培训完成情况。",
            "所有研究文件应当保存在受控的文档系统中，并在研究结束前保持可查阅。",
            "研究结束后，研究资料按计划归档保存。",
        ),
        (),
    ),
    (
        "II期背景说明",
        (
            "本章背景说明适用于II期阶段。",
            "本章为背景描述，不包含可执行的研究管理要求。",
        ),
        (),
    ),
)


def build_synthetic_protocol_docx(path: Path) -> str:
    """在指定路径生成内置词汇中立合成方案 DOCX，返回文件 SHA-256。

    正文仅由 :data:`_SYNTHETIC_DOCX_SECTIONS` 的固定通用语句构成；任何外部
    输入都不会进入合成文档，从结构上保证词汇中立。入排标准使用 Word 官方
    编号列表段落，满足生产链对官方父规则列表的结构要求。
    """

    document = Document()
    document.add_heading(_SYNTHETIC_DOCX_TITLE, level=0)
    for heading, paragraphs, numbered_rules in _SYNTHETIC_DOCX_SECTIONS:
        document.add_heading(heading, level=1)
        for text in paragraphs:
            document.add_paragraph(text)
        for text in numbered_rules:
            document.add_paragraph(text, style="List Number")
        if heading == "II期研究设计":
            table = document.add_table(
                rows=len(_SYNTHETIC_FLOW_TABLE), cols=len(_SYNTHETIC_FLOW_TABLE[0])
            )
            for row_index, row in enumerate(_SYNTHETIC_FLOW_TABLE):
                for column_index, text in enumerate(row):
                    table.rows[row_index].cells[column_index].text = text
            table.style = "Table Grid"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _assert_fixture_allowed(path: Path) -> None:
    """拒绝真实项目/真实数据目录下的方案文件（默认只允许内置合成协议）。"""

    resolved = path.resolve()
    projects_root = (REPO_ROOT / "projects").resolve()
    if projects_root.exists() and (
        resolved == projects_root or projects_root in resolved.parents
    ):
        raise SmokeGuardError(
            f"冒烟运行器拒绝使用真实项目目录下的方案文件：{resolved}；"
            "请使用内置合成协议或位于受控临时目录的合成文件。"
        )
    try:
        default_root = resolve_data_paths().root.resolve()
    except Exception:  # noqa: BLE001 - 守卫为尽力而为；无法解析时跳过该层
        default_root = None
    if default_root is not None and (
        resolved == default_root or default_root in resolved.parents
    ):
        raise SmokeGuardError(
            f"冒烟运行器拒绝使用真实数据目录下的方案文件：{resolved}；"
            "请使用内置合成协议或位于受控临时目录的合成文件。"
        )


# ---------------------------------------------------------------------------
# 模型身份正向核验
# ---------------------------------------------------------------------------


def _transport_identity(transport: Any, *, stage: str, verified_model: str) -> dict[str, Any]:
    """记录非秘密的传输身份（与生产 executor 的持久化字段一致）。"""

    identity: dict[str, Any] = {
        "stage": stage,
        "transport_class": f"{type(transport).__module__}.{type(transport).__qualname__}",
        "configured_model": transport.model,
        "verified_model": verified_model,
    }
    for name in (
        "provider",
        "backend",
        "reasoning_effort",
        "max_tokens",
        "base_url",
        "temperature",
        "timeout",
        "max_retries",
        "response_format_sha256",
    ):
        try:
            value = getattr(transport, name)
        except Exception:  # noqa: BLE001 - 身份记录为尽力而为
            continue
        if callable(value) or value is None:
            continue
        if isinstance(value, (str, int, float, bool)):
            identity[name] = value
    return identity


def verify_live_model_identity(
    *,
    on_event: Callable[[str], None] | None = None,
) -> tuple[Any, Any, dict[str, Any]]:
    """构建生产环境传输并正向核验服务实际加载的模型身份。

    返回 ``(discovery_transport, deep_transport, identity_record)``；
    身份缺失、歧义或不一致时抛出 :class:`ProtocolControlModelIdentityError`。
    """

    discovery = protocol_control_discovery_transport_from_environment()
    deep = protocol_control_transport_from_environment()
    if on_event is not None:
        on_event(
            "核验发现传输实际模型身份：配置模型 "
            f"{discovery.model}，服务 {discovery.base_url}"
        )
    discovery_verified = discovery.verify_model_identity(force=True)
    if on_event is not None:
        on_event(
            "核验深析传输实际模型身份：配置模型 "
            f"{deep.model}，服务 {deep.base_url}"
        )
    deep_verified = deep.verify_model_identity(force=True)
    record = {
        "verified_before_run": True,
        "discovery": _transport_identity(
            discovery, stage="discovery", verified_model=discovery_verified
        ),
        "deep": _transport_identity(deep, stage="deep", verified_model=deep_verified),
    }
    return discovery, deep, record


# ---------------------------------------------------------------------------
# 来源链：真实生产路径驱动 + 冻结后显式停止
# ---------------------------------------------------------------------------


def _smoke_source_executor(production_executor: Callable[[Any], dict[str, Any]]):
    """包装生产解构执行器：冻结快照后以显式非重试边界停止来源任务。"""

    def execute(context: Any) -> dict[str, Any]:
        if context.step_id == STEP_GENERATE:
            raise StepFailure(
                retryable=False,
                error_code=SMOKE_STOP_ERROR_CODE,
                detail=(
                    "冒烟运行器已在冻结结构快照后停止来源任务；"
                    "本冒烟不驱动官方 IN/EX 语义解构路线。"
                ),
            )
        return production_executor(context)

    return execute


def _pick_study_phase(
    phase_candidates: list[dict[str, Any]],
    *,
    selected_phase: StudyPhase | None = None,
) -> StudyPhase:
    """Use the recognized phase set without silently preferring one phase."""

    available = {
        str(item.get("phase"))
        for item in phase_candidates
        if item.get("phase")
    }
    if selected_phase is not None:
        if selected_phase.value not in available:
            raise SmokeGuardError(
                "指定研究期别不在方案识别候选中；"
                f"指定：{selected_phase.value}，实际候选：{sorted(available) or '无'}。"
            )
        return selected_phase
    if len(available) == 1:
        return StudyPhase(next(iter(available)))
    raise SmokeGuardError(
        "方案未产生唯一研究期别；自动回放不得按固定优先级替用户选择。"
        f"实际候选：{sorted(available) or '无'}。请显式指定研究期别。"
    )


def _confirm_inputs_from_review(
    review: Any,
    *,
    selected_phase: StudyPhase | None = None,
    allow_synthetic_fallbacks: bool = True,
) -> dict[str, Any]:
    """从生产识别结果推导确认输入（与页面确认一致的值来源）。"""

    pending: ProtocolIdentityDecision = review.identity_decision
    official_date = pending.official_date
    if official_date is not None and official_date.value is not None:
        date_value = official_date.value.isoformat()
        date_precision = official_date.precision.value
    elif allow_synthetic_fallbacks:
        date_value = "2026-01-01"
        date_precision = "day"
    else:
        raise SmokeGuardError("真实方案未识别出版本日期，回放不会填入合成日期。")
    if not allow_synthetic_fallbacks:
        missing = [
            label
            for label, value in (
                ("方案编号", pending.protocol_code),
                ("项目名称", pending.project_name),
                ("方案版本", pending.official_version),
            )
            if not value
        ]
        if missing:
            raise SmokeGuardError(
                "真实方案身份信息不完整，回放不会填入合成值：" + "、".join(missing)
            )
    return {
        "protocol_code": pending.protocol_code or "SYN-SMOKE-001",
        "project_name": pending.project_name or "示例合成冒烟方案",
        "official_version": pending.official_version or "1.0",
        "official_date_value": date_value,
        "official_date_precision": date_precision,
        "study_phase": _pick_study_phase(
            list(review.phase_candidates),
            selected_phase=selected_phase,
        ),
    }


def _run_source_chain(
    workbench: ProtocolWorkbenchService,
    runner: JobRunner,
    session_factory: Any,
    *,
    now: Callable[[], datetime],
    fixture_path: Path,
    original_name: str,
    idempotency_key: str,
    selected_phase: StudyPhase | None = None,
    allow_synthetic_fallbacks: bool = True,
) -> dict[str, Any]:
    """以真实生产服务与真实 JobRunner 驱动来源链至冻结快照。"""

    started = now()
    start_result = workbench.start_first_deconstruction(
        upload_path=fixture_path,
        original_name=original_name,
        idempotency_key=idempotency_key,
        actor=RUNNER_LABEL,
    )
    source_job_id = start_result.job_id
    # 第一段：登记 -> 提取 -> 渲染对齐 -> 识别，止于用户确认边界。
    runner.run_job(source_job_id)
    review = workbench.get_identity_review(source_job_id)
    confirm_inputs = _confirm_inputs_from_review(
        review,
        selected_phase=selected_phase,
        allow_synthetic_fallbacks=allow_synthetic_fallbacks,
    )
    workbench.confirm_identity(
        source_job_id,
        **confirm_inputs,
        actor=RUNNER_LABEL,
    )
    # 第二段：冻结结构快照；随后在 generate_draft 处显式停止。
    runner.run_job(source_job_id)

    with session_factory() as session:
        store = JobStore(session, now=now)
        job = store.get_job(source_job_id)
        final_state = job.state
        extract_checkpoint = store.get_last_checkpoint(source_job_id, "extract_structure")
        confirm_checkpoint = store.get_last_checkpoint(
            source_job_id, "await_identity_confirm"
        )
        freeze_checkpoint = store.get_last_checkpoint(source_job_id, STEP_FREEZE)
        stop_step = next(
            (
                step
                for step in store.list_steps(source_job_id)
                if step.step_id == STEP_GENERATE
            ),
            None,
        )
    if extract_checkpoint is None or freeze_checkpoint is None:
        raise RuntimeError(
            "来源链未产出冻结结构快照检查点，冒烟运行器拒绝继续。"
        )
    snapshot = extract_checkpoint[1]
    identity_payload = confirm_checkpoint[1] if confirm_checkpoint else {}
    return {
        "job_id": source_job_id,
        "final_state": final_state,
        "stopped_after": SOURCE_STOP_STEP,
        "stop_error_code": (
            stop_step.error_code if stop_step is not None else None
        ),
        "snapshot_id": snapshot.get("snapshot_id"),
        "content_sha256": snapshot.get("content_sha256"),
        "block_count": snapshot.get("block_count"),
        "protocol_code": identity_payload.get("protocol_code"),
        "selected_phase": identity_payload.get("selected_phase"),
        "source_chain_started_at": started.isoformat(),
        "source_chain_finished_at": now().isoformat(),
    }


# ---------------------------------------------------------------------------
# 指标采集
# ---------------------------------------------------------------------------


def _attempt_metrics(attempts: list[Any]) -> dict[str, int]:
    outcomes = Counter(str(getattr(item, "outcome", "")) for item in attempts)
    return {
        "attempt_count": len(attempts),
        "schema_repair_count": int(outcomes.get("schema_invalid", 0)),
        "transport_retry_count": int(outcomes.get("transport_failed", 0)),
    }


def _control_metrics(
    session_factory: Any,
    job_id: str,
    closure: dict[str, Any],
    *,
    now: Callable[[], datetime],
) -> dict[str, Any]:
    """从持久检查点采集发现分布、深析范围、修复次数与最终门禁结果。"""

    distribution: Counter[str] = Counter()
    discovery_batches: list[dict[str, Any]] = []
    deep_batches: list[dict[str, Any]] = []
    totals = {"attempt_count": 0, "schema_repair_count": 0, "transport_retry_count": 0}
    gate: dict[str, Any] | None = None
    failure: dict[str, Any] | None = None
    final_state = None

    with session_factory() as session:
        store = JobStore(session, now=now)
        job = store.get_job(job_id)
        final_state = job.state
        steps = {step.step_id: step for step in store.list_steps(job_id)}
        for step_id in sorted(steps):
            step = steps[step_id]
            if step_id.startswith(_DISCOVERY_STEP_PREFIX):
                checkpoint = store.get_last_checkpoint(job_id, step_id)
                if checkpoint is None:
                    continue
                run_result = ProtocolControlDiscoveryAgentRunResult.model_validate(
                    checkpoint[1]["run_result"]
                )
                attempts = _attempt_metrics(list(run_result.attempts))
                for key in totals:
                    totals[key] += attempts[key]
                decisions = list(run_result.final_output or [])
                batch_distribution: Counter[str] = Counter(
                    decision.disposition.value for decision in decisions
                )
                distribution.update(batch_distribution)
                discovery_batches.append(
                    {
                        "step_id": step_id,
                        "discovery_batch_id": run_result.discovery_batch_id,
                        "decision_count": len(decisions),
                        "distribution": dict(sorted(batch_distribution.items())),
                        **attempts,
                    }
                )
            elif step_id.startswith("deep_"):
                checkpoint = store.get_last_checkpoint(job_id, step_id)
                if checkpoint is None:
                    continue
                run_result = ProtocolControlAgentRunResult.model_validate(
                    checkpoint[1]["run_result"]
                )
                attempts = _attempt_metrics(list(run_result.attempts))
                for key in totals:
                    totals[key] += attempts[key]
                output = run_result.final_output
                deep_batches.append(
                    {
                        "step_id": step_id,
                        "batch_id": output.batch_id if output else None,
                        "owned_structure_unit_ids": (
                            list(output.owned_structure_unit_ids) if output else []
                        ),
                        "candidate_count": len(output.candidates) if output else 0,
                        **attempts,
                    }
                )
            elif step_id == "gate":
                checkpoint = store.get_last_checkpoint(job_id, step_id)
                if checkpoint is not None:
                    gate = checkpoint[1]
            if step.state in {"failed_final", "failed"} and failure is None:
                failure = {
                    "step_id": step.step_id,
                    "error_code": step.error_code,
                    "detail": (step.detail or "")[:2000]
                    if hasattr(step, "detail")
                    else None,
                }
    candidate_ids = sorted(gate.get("candidate_ids", [])) if gate else []
    return {
        "job_id": job_id,
        "final_state": final_state,
        "deep_step_ids": list(closure.get("deep_step_ids", [])),
        "non_deep_structure_unit_ids": list(
            closure.get("non_deep_structure_unit_ids", [])
        ),
        "discovery": {
            "batch_count": len(discovery_batches),
            "distribution": dict(sorted(distribution.items())),
            "batches": discovery_batches,
            **totals,
        },
        "deep": {
            "batch_count": len(deep_batches),
            "batches": deep_batches,
            "candidate_count_total": sum(item["candidate_count"] for item in deep_batches),
            "owned_structure_unit_count": sum(
                len(item["owned_structure_unit_ids"]) for item in deep_batches
            ),
            **totals,
        },
        "result_kind": gate.get("result_kind") if gate else None,
        "formal_catalog_status": gate.get("formal_catalog_status") if gate else None,
        "gate_accepted": gate.get("accepted") if gate else None,
        "gate_version": gate.get("gate_version") if gate else None,
        "candidate_ids_count": len(candidate_ids),
        "failure": failure,
    }


def _candidate_gate_failure_codes(metrics: dict[str, Any]) -> list[str]:
    """Reject stale or malformed terminal checkpoints at the resume boundary."""

    failures: list[str] = []
    if metrics.get("final_state") != "completed":
        failures.append("CONTROL_JOB_NOT_COMPLETED")
    if metrics.get("result_kind") != CANDIDATE_CONTROL_PACKAGE_RESULT_KIND:
        failures.append("CANDIDATE_PACKAGE_RESULT_INVALID")
    if (
        metrics.get("formal_catalog_status")
        != FORMAL_CATALOG_STATUS_NOT_MATERIALIZED
    ):
        failures.append("FORMAL_CATALOG_BOUNDARY_INVALID")
    if metrics.get("gate_accepted") is not True:
        failures.append("CANDIDATE_GATE_NOT_ACCEPTED")
    return failures


def _read_closure_checkpoint(
    session_factory: Any,
    job_id: str,
    *,
    now: Callable[[], datetime],
) -> dict[str, Any]:
    with session_factory() as session:
        store = JobStore(session, now=now)
        checkpoint = store.get_last_checkpoint(job_id, "deterministic_closure")
    if checkpoint is None:
        return {}
    return checkpoint[1]


def _read_persisted_source_chain(
    session_factory: Any,
    *,
    source_job_id: str,
    source_snapshot_id: str,
    source_content_sha256: str,
    now: Callable[[], datetime],
) -> dict[str, Any]:
    """Read the frozen source chain from SQLite without reopening its file."""

    with session_factory() as session:
        store = JobStore(session, now=now)
        source_job = store.get_job(source_job_id)
        extract_checkpoint = store.get_last_checkpoint(
            source_job_id, "extract_structure"
        )
        confirm_checkpoint = store.get_last_checkpoint(
            source_job_id, "await_identity_confirm"
        )
        freeze_checkpoint = store.get_last_checkpoint(source_job_id, STEP_FREEZE)
        if extract_checkpoint is None or freeze_checkpoint is None:
            raise SmokeGuardError(
                "既有来源链未产出冻结结构快照检查点，恢复入口拒绝继续。"
            )
        extract_payload = extract_checkpoint[1]
        confirm_payload = confirm_checkpoint[1] if confirm_checkpoint else {}
        stop_step = next(
            (
                step
                for step in store.list_steps(source_job_id)
                if step.step_id == STEP_GENERATE
            ),
            None,
        )
        return {
            "job_id": source_job_id,
            "final_state": source_job.state,
            "stopped_after": SOURCE_STOP_STEP,
            "stop_error_code": stop_step.error_code if stop_step else None,
            "snapshot_id": source_snapshot_id,
            "content_sha256": source_content_sha256,
            "block_count": extract_payload.get("block_count"),
            "protocol_code": confirm_payload.get("protocol_code"),
            "selected_phase": confirm_payload.get("selected_phase"),
            "resumed_from_persisted_checkpoint": True,
            "source_file_read": False,
        }


def _resume_existing_control_job(
    *,
    out_dir: Path,
    job_id: str,
    retry_failed: bool = False,
    now: Callable[[], datetime] = utc_now,
) -> tuple[dict[str, Any], int]:
    """Resume one persisted control job without rebuilding the source chain.

    The database payload and existing checkpoints are the only source of
    execution inputs.  In particular, this path never accepts or opens a
    protocol file and never creates a second Job or idempotency record.
    """

    if not isinstance(job_id, str) or not job_id.strip():
        raise SmokeGuardError("恢复既有方案控制任务必须提供非空任务编号。")
    started = now()
    out_dir = Path(out_dir).resolve()
    data_root = out_dir / "data_v2"
    db_path = data_root / "enrollment-review-v2.sqlite3"
    if not db_path.is_file():
        raise SmokeGuardError(
            "恢复既有方案控制任务需要已有 SQLite 数据库："
            f"{db_path}"
        )
    safe_job_id = "".join(
        character if character.isalnum() or character in "-_" else "_"
        for character in job_id
    )
    resume_token = uuid.uuid4().hex[:8]
    record_filename = "run_record.json"
    if (out_dir / record_filename).exists():
        record_filename = (
            f"resume-{safe_job_id[:48]}-"
            f"{started.strftime('%Y%m%dT%H%M%S')}-{resume_token}.json"
        )
    record: dict[str, Any] = {
        "run_id": f"resume-{started.strftime('%Y%m%dT%H%M%S')}-{resume_token}",
        "claims_complete": False,
        "runner": "scripts/run_protocol_control_smoke.py",
        "started_at": started.isoformat(),
        "fixture": None,
        "model_identity": None,
        "source_chain": None,
        "control_execution": None,
        "semantic_acceptance": None,
        "resume": {
            "requested_job_id": job_id,
            "retry_failed_requested": retry_failed,
            "runner_claimed": False,
            "recovery": None,
            "completed_step_ids_before": [],
        },
        "boundaries": {
            "formal_catalog_materialized": False,
            "source_snapshot_bypassed": False,
            "fixture_is_synthetic_only": False,
            "read_only_protocol_replay": True,
            # Existing v1 jobs did not persist whether the source came from the
            # synthetic smoke fixture or an external read-only protocol.
            "real_protocol_content_used": None,
            "real_clinical_content_used": False,
            # No source-file read occurs in this path, so do not claim a
            # fresh mtime/hash comparison.
            "source_protocol_unchanged": None,
        },
        "notes": [
            "恢复路径只读取既有 SQLite 任务、冻结来源载荷和检查点；"
            "未重新读取原始方案文件。"
        ],
    }
    exit_code = EXIT_OK
    engine = None
    session_factory = None
    try:
        data_paths = resolve_data_paths(env_override=str(data_root))
        engine = build_engine(data_paths.db_path)
        session_factory = build_session_factory(engine)
        with session_factory() as session:
            store = JobStore(session, now=now)
            job = store.get_job(job_id)
            if job.job_type != PROTOCOL_CONTROL_EXECUTION_JOB_TYPE:
                raise SmokeGuardError(
                    "恢复入口只接受方案控制任务；实际任务类型："
                    f"{job.job_type}"
                )
            payload = verify_payload_sha256(job.payload_json, job.payload_sha256)
            source_job_id = payload.get("source_deconstruction_job_id")
            source_snapshot_id = payload.get("source_snapshot_id")
            source_content_sha256 = payload.get("source_content_sha256")
            if not all(
                isinstance(value, str) and value.strip()
                for value in (
                    source_job_id,
                    source_snapshot_id,
                    source_content_sha256,
                )
            ):
                raise SmokeGuardError(
                    "既有方案控制任务缺少可恢复的冻结来源身份或内容摘要。"
                )
            record["resume"]["before_state"] = job.state
            record["resume"]["before_progress"] = {
                "completed": job.progress_completed,
                "total": job.progress_total,
            }
            record["resume"]["completed_step_ids_before"] = [
                step.step_id
                for step in store.list_steps(job_id)
                if step.state == "completed"
            ]
        record["fixture"] = {
            "source": "persisted_frozen_snapshot",
            "path": None,
            "sha256": source_content_sha256,
        }
        record["source_chain"] = _read_persisted_source_chain(
            session_factory,
            source_job_id=source_job_id,
            source_snapshot_id=source_snapshot_id,
            source_content_sha256=source_content_sha256,
            now=now,
        )

        try:
            discovery_transport, deep_transport, identity_record = (
                verify_live_model_identity(
                    on_event=lambda text: record["notes"].append(text)
                )
            )
            record["model_identity"] = identity_record
        except ProtocolControlModelIdentityError as exc:
            record["model_identity"] = {
                "verified_before_run": False,
                "failure_class": "model_identity_mismatch",
                "reason": getattr(exc, "reason", "unknown"),
                "configured_model": getattr(exc, "configured_model", None),
                "served_model_ids": list(getattr(exc, "served_model_ids", ())),
                "diagnostic": str(exc),
            }
            record["notes"].append(
                "既有任务恢复前模型身份未通过正向匹配；"
                "未执行新的语义请求。"
            )
            return _finish(
                record,
                out_dir,
                started,
                now,
                EXIT_IDENTITY_FAIL_CLOSED,
                record_filename=record_filename,
            )

        if retry_failed:
            with session_factory() as session, session.begin():
                outcome = JobStore(session, now=now).retry_failed(job_id)
            record["resume"]["manual_retry"] = {
                "changed": outcome.changed,
                "state": outcome.state,
            }
            record["notes"].append(
                "已按显式请求仅重新排队失败步骤；已完成检查点保持不变。"
            )

        recovery = recover_expired_jobs(session_factory, now=now)
        record["resume"]["recovery"] = {
            "recovered_jobs": list(recovery.recovered_jobs),
            "cancelled_jobs": list(recovery.cancelled_jobs),
            "requeued_jobs": list(recovery.requeued_jobs),
            "failed_final_jobs": list(recovery.failed_final_jobs),
        }
        if job_id in recovery.requeued_jobs:
            record["notes"].append(
                "中断租约已按持久检查点恢复；仅未完成步骤重新进入队列。"
            )
        elif job_id not in recovery.recovered_jobs:
            record["notes"].append(
                "本次恢复未接管该任务租约；若仍有其他有效 worker，保持其执行权。"
            )

        control_executor = create_protocol_control_executor(
            ProtocolControlExecutorConfig(
                data_paths=data_paths,
                session_factory=session_factory,
                now=now,
                discovery_transport=discovery_transport,
                deep_transport=deep_transport,
            )
        )
        runner = JobRunner(
            session_factory,
            {PROTOCOL_CONTROL_EXECUTION_JOB_TYPE: control_executor},
            worker_id=f"{RUNNER_LABEL}-resume",
            now=now,
        )
        control_started = now()
        claimed = runner.run_job(job_id)
        record["resume"]["runner_claimed"] = claimed
        control_finished = now()

        closure = _read_closure_checkpoint(session_factory, job_id, now=now)
        metrics = _control_metrics(
            session_factory,
            job_id,
            closure,
            now=now,
        )
        metrics["latency_seconds"] = {
            "control_execution": round(
                (control_finished - control_started).total_seconds(), 3
            ),
        }
        record["control_execution"] = metrics
        semantic_failures = _candidate_gate_failure_codes(metrics)
        record["semantic_acceptance"] = {
            "mode": "external_observation_only",
            "incremental_candidate_required": False,
            "passed": not semantic_failures,
            "failure_codes": semantic_failures,
        }
        if semantic_failures:
            exit_code = EXIT_RUN_FAILED
        else:
            record["notes"].append(
                "既有方案控制任务已继续至技术终态；"
                "正式目录仍未物化，临床验收仍待独立核对。"
            )
    except SmokeGuardError:
        raise
    except Exception as exc:  # noqa: BLE001 - resume record must retain failure
        record["failure_class"] = "unexpected_error"
        record["failure_detail"] = f"{type(exc).__name__}: {exc}"
        record["failure_traceback"] = traceback.format_exc()[-4000:]
        exit_code = EXIT_RUN_FAILED
    finally:
        if engine is not None:
            engine.dispose()
    return _finish(
        record,
        out_dir,
        started,
        now,
        exit_code,
        record_filename=record_filename,
    )


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------


def run_smoke(
    *,
    out_dir: Path,
    fixture_path: Path | None = None,
    selected_phase: StudyPhase | None = None,
    allow_read_only_protocol: bool = False,
    resume_job_id: str | None = None,
    retry_failed: bool = False,
    now: Callable[[], datetime] = utc_now,
) -> tuple[dict[str, Any], int]:
    """Execute a full smoke run or resume one persisted control job."""

    global _NOW
    _NOW = now
    if resume_job_id is not None:
        if fixture_path is not None or selected_phase is not None or allow_read_only_protocol:
            raise SmokeGuardError(
                "恢复既有任务不得同时提供方案文件、研究期别或只读方案开关。"
            )
        return _resume_existing_control_job(
            out_dir=out_dir,
            job_id=resume_job_id,
            retry_failed=retry_failed,
            now=now,
        )
    if retry_failed:
        raise SmokeGuardError("--retry-failed 只能与 --resume-job-id 同时使用。")
    started = now()
    out_dir.mkdir(parents=True, exist_ok=True)
    record: dict[str, Any] = {
        "run_id": f"smoke-{started.strftime('%Y%m%dT%H%M%S')}-{uuid.uuid4().hex[:8]}",
        "claims_complete": False,
        "runner": "scripts/run_protocol_control_smoke.py",
        "started_at": started.isoformat(),
        "fixture": None,
        "model_identity": None,
        "source_chain": None,
        "control_execution": None,
        "semantic_acceptance": None,
        "boundaries": {
            "formal_catalog_materialized": False,
            "source_snapshot_bypassed": False,
            "fixture_is_synthetic_only": fixture_path is None,
            "read_only_protocol_replay": allow_read_only_protocol,
            "real_protocol_content_used": bool(
                fixture_path is not None and allow_read_only_protocol
            ),
            "real_clinical_content_used": False,
            "source_protocol_unchanged": None,
        },
        "notes": [],
    }
    exit_code = EXIT_OK
    discovery_transport = None
    deep_transport = None
    session_factory = None
    engine = None
    try:
        # 1) 先核验实际模型身份；失败即失效关闭，不建立任何任务。
        try:
            discovery_transport, deep_transport, identity_record = (
                verify_live_model_identity(
                    on_event=lambda text: record["notes"].append(text)
                )
            )
            record["model_identity"] = identity_record
        except ProtocolControlModelIdentityError as exc:
            record["model_identity"] = {
                "verified_before_run": False,
                "failure_class": "model_identity_mismatch",
                "reason": getattr(exc, "reason", "unknown"),
                "configured_model": getattr(exc, "configured_model", None),
                "served_model_ids": list(getattr(exc, "served_model_ids", ())),
                "diagnostic": str(exc),
            }
            record["notes"].append(
                "实际模型身份未通过正向匹配，已按边界失效关闭；未执行任何语义请求。"
            )
            return _finish(record, out_dir, started, now, EXIT_IDENTITY_FAIL_CLOSED)

        # 2) 准备词汇中立合成协议与受控临时数据目录。
        if fixture_path is not None:
            fixture = Path(fixture_path)
            if not fixture.is_file():
                raise SmokeGuardError(f"方案文件不存在：{fixture}")
            if allow_read_only_protocol:
                if selected_phase is None:
                    raise SmokeGuardError("真实方案只读回放必须显式指定研究期别。")
            else:
                _assert_fixture_allowed(fixture)
            fixture_sha256 = hashlib.sha256(fixture.read_bytes()).hexdigest()
            record["fixture"] = {
                "source": "explicit_external",
                "path": str(fixture),
                "sha256": fixture_sha256,
            }
        else:
            fixture = out_dir / "synthetic_protocol.docx"
            fixture_sha256 = build_synthetic_protocol_docx(fixture)
            record["fixture"] = {
                "source": "bundled_synthetic",
                "path": str(fixture),
                "sha256": fixture_sha256,
            }

        data_paths = resolve_data_paths(env_override=str(out_dir / "data_v2"))
        data_paths.ensure_directories()
        MigrationManager(data_paths).upgrade("head")
        engine = build_engine(data_paths.db_path)
        session_factory = build_session_factory(engine)
        workbench = ProtocolWorkbenchService(session_factory, data_paths=data_paths, now=now)
        production_source_executor = create_protocol_deconstruction_executor(
            ProtocolDeconstructionExecutorConfig(
                data_paths=data_paths,
                session_factory=session_factory,
                now=now,
            )
        )
        control_executor = create_protocol_control_executor(
            ProtocolControlExecutorConfig(
                data_paths=data_paths,
                session_factory=session_factory,
                now=now,
                discovery_transport=discovery_transport,
                deep_transport=deep_transport,
            )
        )
        runner = JobRunner(
            session_factory,
            {
                PROTOCOL_DECONSTRUCTION_JOB_TYPE: _smoke_source_executor(
                    production_source_executor
                ),
                PROTOCOL_CONTROL_EXECUTION_JOB_TYPE: control_executor,
            },
            worker_id=RUNNER_LABEL,
            now=now,
        )

        # 3) 真实生产来源链至冻结结构快照。
        record["notes"].append("开始真实生产来源链（登记->提取->渲染->识别->确认->冻结）。")
        source_metrics = _run_source_chain(
            workbench,
            runner,
            session_factory,
            now=now,
            fixture_path=fixture,
            original_name=fixture.name,
            idempotency_key=f"{record['run_id']}-source",
            selected_phase=selected_phase,
            allow_synthetic_fallbacks=not allow_read_only_protocol,
        )
        record["source_chain"] = source_metrics
        source_unchanged = hashlib.sha256(fixture.read_bytes()).hexdigest() == fixture_sha256
        record["boundaries"]["source_protocol_unchanged"] = source_unchanged
        if not source_unchanged:
            raise SmokeGuardError("来源方案在只读回放期间发生变化，已停止后续处理。")

        # 4) 生产控制任务 + 生产执行器 + 真实 JobRunner。
        control_service = ProtocolControlJobService(
            session_factory,
            data_paths=data_paths,
            actor=RUNNER_LABEL,
        )
        control_result = control_service.create_from_deconstruction(
            source_job_id=source_metrics["job_id"],
            idempotency_key=f"{record['run_id']}-control",
        )
        control_job_id = control_result.job_id
        record["notes"].append(
            f"控制任务已建立：{control_job_id}（来源快照 {control_result.snapshot_id}）。"
        )
        control_started = now()
        runner.run_job(control_job_id)
        control_finished = now()

        # 5) 采集持久化指标并写运行记录。
        closure = _read_closure_checkpoint(session_factory, control_job_id, now=now)
        metrics = _control_metrics(session_factory, control_job_id, closure, now=now)
        metrics["latency_seconds"] = {
            "control_execution": round(
                (control_finished - control_started).total_seconds(), 3
            ),
        }
        record["control_execution"] = metrics
        semantic_failures = _candidate_gate_failure_codes(metrics)
        if fixture_path is None and metrics["candidate_ids_count"] < 1:
            semantic_failures.append("SYNTHETIC_INCREMENTAL_CONTROL_MISSING")
        record["semantic_acceptance"] = {
            "mode": (
                "bundled_positive_control"
                if fixture_path is None
                else "external_observation_only"
            ),
            "incremental_candidate_required": fixture_path is None,
            "passed": not semantic_failures,
            "failure_codes": semantic_failures,
        }
        if semantic_failures:
            exit_code = EXIT_RUN_FAILED
        else:
            record["notes"].append(
                "控制任务完成：结果为水合候选控制点包，正式目录保持未物化。"
            )
    except ProtocolWorkbenchError as exc:
        record["failure_class"] = "workbench_error"
        record["failure_detail"] = f"{exc.code}: {exc.detail}"
        exit_code = EXIT_RUN_FAILED
    except ProtocolControlExecutionError as exc:
        record["failure_class"] = "control_job_creation_error"
        record["failure_detail"] = f"{exc.code}: {exc.detail}"
        exit_code = EXIT_RUN_FAILED
    except ProtocolControlModelIdentityError as exc:
        record["model_identity"] = record.get("model_identity") or {}
        record["model_identity"]["failure_class"] = "model_identity_mismatch"
        record["model_identity"]["diagnostic"] = str(exc)
        exit_code = EXIT_IDENTITY_FAIL_CLOSED
    except Exception as exc:  # noqa: BLE001 - 运行记录必须完整落盘
        record["failure_class"] = "unexpected_error"
        record["failure_detail"] = f"{type(exc).__name__}: {exc}"
        record["failure_traceback"] = traceback.format_exc()[-4000:]
        exit_code = EXIT_RUN_FAILED
    finally:
        if engine is not None:
            engine.dispose()
    return _finish(record, out_dir, started, now, exit_code)


def _finish(
    record: dict[str, Any],
    out_dir: Path,
    started: datetime,
    now: Callable[[], datetime],
    exit_code: int,
    *,
    record_filename: str = "run_record.json",
) -> tuple[dict[str, Any], int]:
    finished = now()
    record["finished_at"] = finished.isoformat()
    record["latency_seconds"] = {
        "total": round((finished - started).total_seconds(), 3),
    }
    record["claims_complete"] = False
    record_path = out_dir / record_filename
    record["run_record_path"] = str(record_path)
    record_path.write_text(
        json.dumps(record, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return record, exit_code


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="词汇中立的真实生产 harness 冒烟运行器（方案控制点两阶段生产链）。"
    )
    parser.add_argument(
        "--out-dir",
        help="运行记录与临时数据目录（默认 runs/protocol_control_smoke/<run_id>）",
    )
    parser.add_argument(
        "--fixture",
        help="可选的合成方案 DOCX；默认使用内置词汇中立合成协议。"
        "拒绝 projects/ 与真实数据目录下的文件。",
    )
    parser.add_argument(
        "--read-only-protocol",
        action="store_true",
        help="显式允许把外部原始 DOCX/PDF 作为只读异构回放语料。",
    )
    parser.add_argument(
        "--study-phase",
        choices=[item.value for item in StudyPhase],
        help="只读真实方案回放时必须明确指定的研究期别。",
    )
    parser.add_argument(
        "--resume-job-id",
        help="从既有 --out-dir/data_v2 SQLite 任务恢复；不重新读取方案文件。",
    )
    parser.add_argument(
        "--retry-failed",
        action="store_true",
        help="显式重新排队既有任务的失败步骤；已完成步骤与检查点保持不变。",
    )
    args = parser.parse_args(argv)
    if args.resume_job_id:
        if not args.out_dir:
            parser.error("--resume-job-id 必须同时提供 --out-dir")
        if args.fixture or args.read_only_protocol or args.study_phase:
            parser.error(
                "--resume-job-id 不得同时提供 --fixture、--read-only-protocol 或 --study-phase"
            )
    elif args.retry_failed:
        parser.error("--retry-failed 必须同时提供 --resume-job-id")

    if args.read_only_protocol and not args.fixture:
        parser.error("--read-only-protocol 必须同时提供 --fixture")
    if args.read_only_protocol and not args.study_phase:
        parser.error("--read-only-protocol 必须同时提供 --study-phase")
    if args.study_phase and not args.fixture:
        parser.error("--study-phase 只用于显式外部方案回放")

    if args.out_dir:
        out_dir = Path(args.out_dir).resolve()
    else:
        stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S")
        out_dir = REPO_ROOT / "runs" / "protocol_control_smoke" / stamp
    fixture_path = Path(args.fixture).resolve() if args.fixture else None
    selected_phase = StudyPhase(args.study_phase) if args.study_phase else None

    try:
        record, exit_code = run_smoke(
            out_dir=out_dir,
            fixture_path=fixture_path,
            selected_phase=selected_phase,
            allow_read_only_protocol=args.read_only_protocol,
            resume_job_id=args.resume_job_id,
            retry_failed=args.retry_failed,
        )
    except SmokeGuardError as exc:
        print(f"边界守卫：{exc}", file=sys.stderr)
        return EXIT_USAGE_OR_GUARD
    except Exception as exc:  # noqa: BLE001 - CLI 边界
        print(f"冒烟运行器异常：{type(exc).__name__}: {exc}", file=sys.stderr)
        return EXIT_USAGE_OR_GUARD

    print(f"运行记录：{record.get('run_record_path')}")
    print(f"claims_complete：{record['claims_complete']}")
    identity = record.get("model_identity") or {}
    if identity.get("verified_before_run"):
        print(
            "模型身份：discovery={0} deep={1}".format(
                (identity.get("discovery") or {}).get("verified_model"),
                (identity.get("deep") or {}).get("verified_model"),
            )
        )
    else:
        print(
            "模型身份：未通过核验（{0}）：{1}".format(
                identity.get("reason"), identity.get("diagnostic", "")[:300]
            )
        )
    control = record.get("control_execution") or {}
    if control:
        print(f"控制任务终态：{control.get('final_state')}")
        print(f"发现分布：{(control.get('discovery') or {}).get('distribution')}")
        print(f"深析批次：{(control.get('deep') or {}).get('batch_count')}")
        print(f"结果类型：{control.get('result_kind')}")
        print(f"正式目录状态：{control.get('formal_catalog_status')}")
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
