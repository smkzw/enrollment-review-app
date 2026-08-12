#!/usr/bin/env python3
"""Run D001 phase-II screening batch review and generate comparison reports."""

from __future__ import annotations

import argparse
import concurrent.futures
import html
import json
import mimetypes
import re
import shutil
import subprocess
import sys
import threading
import time
from dataclasses import asdict
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.batch_sources import BatchSourceFile, collect_review_files_from_folders, discover_subject_folders, safe_staged_filename
from app.markdown_export import generate_markdown_report
from app.pipeline.reviewer import (
    is_future_phase_verify_reasoning,
    is_source_traceability_verify_reasoning,
    parse_review_response,
    pass_verify_label,
)


PROJECT_CODE = "D001-02-II"
STUDY_STAGE = "Ⅱ期"
REVIEW_PHASE = "screening_run_in"
DEFAULT_BASE_URL = "http://127.0.0.1:8900"
SOURCE_ROOT = Path("/Users/smkzw/Documents/康哲项目资料/AI/入排/test-D001项目/全量-入组")
MANUAL_IE_XLSX = Path("/Users/smkzw/Documents/康哲项目资料/AI/入排/test-D001项目/EXCEL_D001-02-002_20260601084053732.xlsx")
PROJECT_DIR = ROOT / "projects" / PROJECT_CODE
REPORT_DIR = PROJECT_DIR / "reports"
STAGING_ROOT = ROOT / "output" / "d001_phase2_upload_staging"
RUN_LOG_DIR = PROJECT_DIR / "rerun_logs"

ISSUE_VERDICTS = {"fail", "needs_evidence", "insufficient", "investigator"}
VERDICT_LABELS = {
    "pass": "通过",
    "pass_verify": "通过（需验证）",
    "fail": "不符合",
    "needs_evidence": "待补证",
    "insufficient": "证据不足",
    "investigator": "需研究者判定",
    "na": "不适用",
    "": "未审核",
}


def now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def compact_ts() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def authenticate(session: requests.Session, base_url: str, username: str, password: str) -> dict:
    resp = session.post(f"{base_url}/api/auth/login", json={"username": username, "password": password}, timeout=120)
    if resp.status_code >= 400:
        raise RuntimeError(f"login failed: HTTP {resp.status_code}: {resp.text}")
    data = resp.json()
    session.headers.update({
        "X-Enrollment-User": data["username"],
        "X-Enrollment-Token": data["token"],
    })
    return data


def request_json(session: requests.Session, method: str, url: str, **kwargs) -> dict:
    resp = session.request(method, url, timeout=120, **kwargs)
    if resp.status_code == 404 and kwargs.pop("allow_404", False):
        return {}
    if resp.status_code >= 400:
        raise RuntimeError(f"{method} {url} failed: HTTP {resp.status_code}: {resp.text}")
    return resp.json() if resp.text else {}


def load_manual_ie() -> tuple[dict[str, dict], dict[str, str]]:
    manual: dict[str, dict] = {}
    center_names: dict[str, str] = {}
    wb = load_workbook(MANUAL_IE_XLSX, read_only=True, data_only=True)
    ws = wb["IE"]
    for row in ws.iter_rows(min_row=3, values_only=True):
        center_name = str(row[0] or "").strip()
        center_code = normalize_center_code(row[1])
        subject_id = str(row[4] or "").strip().upper()
        visit = str(row[9] or "").strip()
        if center_code and center_name:
            center_names[center_code] = center_name
        if not subject_id or visit != "筛选总结":
            continue
        manual[subject_id] = {
            "center_name": center_name,
            "center_code": center_code,
            "subject_status": str(row[6] or "").strip(),
            "ieyn": str(row[19] or "").strip(),
            "manual_item": str(row[21] or "").strip(),
            "manual_comment": str(row[23] or "").strip(),
        }
    return manual, center_names


def normalize_center_code(value: object) -> str:
    text = str(value or "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    if text.isdigit():
        return text.zfill(2)
    return text


def infer_center_from_subject(subject_id: str, manual: dict[str, dict], center_names: dict[str, str]) -> tuple[str, str]:
    if subject_id in manual:
        return manual[subject_id].get("center_code", ""), manual[subject_id].get("center_name", "")
    match = re.match(r"SA(\d{2})", subject_id)
    center_code = match.group(1) if match else ""
    return center_code, center_names.get(center_code, "")


def existing_project_subject_ids() -> set[str]:
    subjects_dir = PROJECT_DIR / "subjects"
    if not subjects_dir.exists():
        return set()
    return {path.name.upper() for path in subjects_dir.iterdir() if path.is_dir()}


def completed_project_subject_ids() -> set[str]:
    subjects_dir = PROJECT_DIR / "subjects"
    if not subjects_dir.exists():
        return set()
    completed: set[str] = set()
    for info_path in subjects_dir.glob("*/info.json"):
        try:
            info = json.loads(info_path.read_text(encoding="utf-8"))
        except Exception:
            continue
        sid = str(info.get("subject_id") or info_path.parent.name).strip().upper()
        phase_report = info_path.parent / "llm" / REVIEW_PHASE / "review_report.md"
        legacy_report = info_path.parent / "llm" / "review_report.md"
        if (
            sid
            and info.get("status") == "reviewed"
            and str(info.get("overall_verdict") or "").strip()
            and (phase_report.exists() or legacy_report.exists())
        ):
            completed.add(sid)
    return completed


def delete_subject_if_exists(session: requests.Session, base_url: str, subject_id: str) -> None:
    resp = session.delete(f"{base_url}/api/projects/{PROJECT_CODE}/subjects/{subject_id}", timeout=120)
    if resp.status_code not in {200, 404}:
        raise RuntimeError(f"delete subject {subject_id} failed: HTTP {resp.status_code}: {resp.text}")


def create_subject(session: requests.Session, base_url: str, subject_id: str, center_code: str, center_name: str) -> dict:
    return request_json(
        session,
        "POST",
        f"{base_url}/api/projects/{PROJECT_CODE}/subjects",
        json={"subject_id": subject_id, "center_code": center_code, "center_name": center_name},
    )


def stage_files(files: list[BatchSourceFile], staging_dir: Path) -> tuple[list[dict], list[dict]]:
    staging_dir.mkdir(parents=True, exist_ok=True)
    staged: list[dict] = []
    conversion_notes: list[dict] = []
    used_names: set[str] = set()
    for source in files:
        target_name = safe_staged_filename(source)
        if source.path.suffix.lower() == ".doc":
            target_name = str(Path(target_name).with_suffix(".docx"))
        stem = Path(target_name).stem
        suffix = Path(target_name).suffix
        candidate = target_name
        index = 2
        while candidate in used_names:
            candidate = f"{stem}_{index}{suffix}"
            index += 1
        used_names.add(candidate)
        target = staging_dir / candidate
        if source.path.suffix.lower() == ".doc":
            note = {"source": str(source.path), "target": str(target), "status": "pending"}
            result = subprocess.run(
                ["textutil", "-convert", "docx", str(source.path), "-output", str(target)],
                text=True,
                capture_output=True,
                timeout=90,
            )
            if result.returncode != 0 or not target.exists():
                note.update({"status": "failed", "stderr": result.stderr.strip()})
                conversion_notes.append(note)
                continue
            note["status"] = "converted"
            conversion_notes.append(note)
        else:
            shutil.copy2(source.path, target)
        staged.append({
            "path": str(target),
            "filename": target.name,
            "category": source.category,
            "source_path": str(source.path),
            "relative_path": source.relative_path,
        })
    return staged, conversion_notes


def upload_staged_files(
    session: requests.Session,
    base_url: str,
    subject_id: str,
    staged: list[dict],
    center_code: str,
    center_name: str,
) -> dict:
    multipart = []
    handles = []
    try:
        for item in staged:
            path = Path(item["path"])
            mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
            handle = path.open("rb")
            handles.append(handle)
            multipart.append(("files", (item["filename"], handle, mime)))
            multipart.append(("categories", (None, item["category"])))
        data = {"center_code": center_code, "center_name": center_name}
        resp = session.post(
            f"{base_url}/api/projects/{PROJECT_CODE}/subjects/{subject_id}/upload",
            files=multipart,
            data=data,
            timeout=900,
        )
        if resp.status_code >= 400:
            raise RuntimeError(f"upload {subject_id} failed: HTTP {resp.status_code}: {resp.text}")
        return resp.json()
    finally:
        for handle in handles:
            handle.close()


def process_subject(session: requests.Session, base_url: str, subject_id: str) -> dict:
    params = {"phase": REVIEW_PHASE, "study_stage": STUDY_STAGE}
    events = []
    start = time.time()
    with session.get(
        f"{base_url}/api/projects/{PROJECT_CODE}/subjects/{subject_id}/process",
        params=params,
        stream=True,
        timeout=None,
    ) as resp:
        if resp.status_code >= 400:
            raise RuntimeError(f"process {subject_id} failed: HTTP {resp.status_code}: {resp.text}")
        last_stage = ""
        for line in resp.iter_lines(decode_unicode=True):
            if not line or not line.startswith("data: "):
                continue
            event = json.loads(line[6:])
            events.append(event)
            stage = event.get("stage", "")
            message = event.get("message", "")
            if stage != last_stage or stage in {"done", "error"} or "审核完成" in message:
                print(f"[{subject_id}] {stage}: {message}", flush=True)
                last_stage = stage
            if stage == "error":
                raise RuntimeError(message)
    return {"elapsed_seconds": round(time.time() - start, 1), "events": events}


def read_subject_report(subject_id: str) -> dict:
    sd = PROJECT_DIR / "subjects" / subject_id
    raw_path = sd / "llm" / REVIEW_PHASE / "review_raw.md"
    report_path = sd / "llm" / REVIEW_PHASE / "review_report.md"
    text = ""
    if report_path.exists():
        text = report_path.read_text(encoding="utf-8")
    elif raw_path.exists():
        text = raw_path.read_text(encoding="utf-8")
    if not text:
        return {"subject_id": subject_id, "overall": "", "summary": "未生成审核报告", "rules": [], "error": "missing_report"}
    parsed = parse_review_response(text)
    return {
        "subject_id": subject_id,
        "overall": parsed.get("overall_verdict", ""),
        "summary": parsed.get("summary", ""),
        "rules": [
            {
                "rule_id": r.rule_id,
                "rule_name": r.rule_name,
                "rule_type": r.rule_type,
                "verdict": r.verdict,
                "verification_type": (
                    "future_phase" if r.verdict == "pass_verify" and is_future_phase_verify_reasoning(r.reasoning)
                    else "source_traceability" if r.verdict == "pass_verify" and is_source_traceability_verify_reasoning(r.reasoning)
                    else "other" if r.verdict == "pass_verify"
                    else ""
                ),
                "reasoning": r.reasoning,
            }
            for r in parsed.get("rule_results", [])
        ],
    }


def manual_rule_ids(text: str) -> list[str]:
    ids: list[str] = []
    for kind, num in re.findall(r"(入选标准|排除标准)\s*(\d+)", text or ""):
        prefix = "IN" if kind == "入选标准" else "EX"
        ids.append(f"{prefix}-{int(num):02d}")
    return ids


def compare_manual(system_report: dict, manual: dict | None) -> dict:
    if not manual:
        return {"status": "人工IE无记录", "expected_rules": [], "matched_rules": [], "missing_rules": []}
    manual_ie = str(manual.get("ieyn") or "").strip()
    issues = {r["rule_id"]: r for r in system_report.get("rules", []) if r.get("verdict") in ISSUE_VERDICTS}
    if manual_ie == "未判断":
        return {"status": "人工未判断", "expected_rules": [], "matched_rules": [], "missing_rules": []}
    if manual_ie in {"是", "Y", "y", "YES", "Yes", "yes", "通过", "符合"}:
        if issues:
            return {
                "status": "存在差异：人工通过但系统仍有关注条目",
                "expected_rules": [],
                "matched_rules": [],
                "missing_rules": sorted(issues),
            }
        return {"status": "人工通过且系统未见关注条目", "expected_rules": [], "matched_rules": [], "missing_rules": []}
    expected = manual_rule_ids(manual.get("manual_item", ""))
    matched = [rid for rid in expected if rid in issues]
    missing = [rid for rid in expected if rid not in issues]
    if not expected:
        return {"status": "人工未填具体条目", "expected_rules": [], "matched_rules": [], "missing_rules": []}
    if missing:
        return {"status": "存在差异：人工条目未被系统命中", "expected_rules": expected, "matched_rules": matched, "missing_rules": missing}
    return {"status": "人工条目被系统命中", "expected_rules": expected, "matched_rules": matched, "missing_rules": []}


def collect_project_markdown() -> str:
    cfg = json.loads((PROJECT_DIR / "config.json").read_text(encoding="utf-8"))
    return generate_markdown_report(
        project_path=PROJECT_DIR,
        project=cfg,
        scope="project",
        phase=REVIEW_PHASE,
    )


def clean_text(text: object) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def process_subject_item(
    args: argparse.Namespace,
    item,
    index: int,
    total: int,
    manual: dict[str, dict],
    center_names: dict[str, str],
    staging_root: Path,
) -> dict:
    center_code, center_name = infer_center_from_subject(item.subject_id, manual, center_names)
    subject_log = {
        "subject_id": item.subject_id,
        "folder": str(item.folder),
        "folders": [str(folder) for folder in (item.folders or (item.folder,))],
        "study_stage": item.study_stage,
        "center_code": center_code,
        "center_name": center_name,
        "retained_files": [],
        "skipped_files": [],
        "staged_files": [],
        "conversion_notes": [],
        "upload_response": {},
        "process": {},
        "error": "",
    }
    session = requests.Session()
    try:
        authenticate(session, args.base_url, args.username, args.password)
        print(f"\n=== {index}/{total} {item.subject_id} ===", flush=True)
        retained, skipped = collect_review_files_from_folders(item.folders or (item.folder,), item.subject_id)
        subject_log["retained_files"] = [asdict(f) | {"path": str(f.path)} for f in retained]
        subject_log["skipped_files"] = skipped
        staged, conversions = stage_files(retained, staging_root / item.subject_id)
        subject_log["staged_files"] = staged
        subject_log["conversion_notes"] = conversions
        if not staged:
            raise RuntimeError("没有可上传的审核文件")
        if args.recreate:
            delete_subject_if_exists(session, args.base_url, item.subject_id)
        create_subject(session, args.base_url, item.subject_id, center_code, center_name)
        subject_log["upload_response"] = upload_staged_files(
            session, args.base_url, item.subject_id, staged, center_code, center_name
        )
        subject_log["process"] = process_subject(session, args.base_url, item.subject_id)
    except Exception as exc:
        subject_log["error"] = str(exc)
        print(f"[{item.subject_id}] ERROR: {exc}", file=sys.stderr, flush=True)
    finally:
        session.close()
    return subject_log


def build_reports(run_log: dict, manual: dict[str, dict]) -> tuple[Path, Path, list[dict]]:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    rows = []
    for subject in run_log["subjects"]:
        sid = subject["subject_id"]
        report = read_subject_report(sid)
        issues = [r for r in report.get("rules", []) if r.get("verdict") in ISSUE_VERDICTS]
        reminders = [r for r in report.get("rules", []) if r.get("verdict") == "pass_verify"]
        comparison = compare_manual(report, manual.get(sid))
        rows.append({
            "subject_id": sid,
            "center_code": subject.get("center_code", ""),
            "center_name": subject.get("center_name", ""),
            "folder": subject.get("folder", ""),
            "overall": report.get("overall", ""),
            "summary": report.get("summary", ""),
            "issues": issues,
            "reminders": reminders,
            "manual": manual.get(sid, {}),
            "comparison": comparison,
            "process_error": subject.get("error", ""),
            "skipped_count": len(subject.get("skipped_files", [])),
            "uploaded_count": len(subject.get("staged_files", [])),
            "conversion_notes": subject.get("conversion_notes", []),
        })

    project_markdown = collect_project_markdown()
    markdown_path = REPORT_DIR / "d001_phase2_screening_full_batch_report.md"
    project_md_path = REPORT_DIR / "D001-02-II_screening_run_in_project_report.md"
    project_md_path.write_text(project_markdown, encoding="utf-8")
    markdown_path.write_text(render_markdown(rows, run_log), encoding="utf-8")
    html_path = REPORT_DIR / "d001_phase2_screening_full_batch_report.html"
    html_path.write_text(render_html(rows, run_log), encoding="utf-8")
    return markdown_path, html_path, rows


def render_markdown(rows: list[dict], run_log: dict) -> str:
    lines = [
        "# D001-02-II Ⅱ期筛选期批量入排审核与人工IE对照",
        "",
        f"- 生成时间：{now()}",
        f"- 源目录：{SOURCE_ROOT}",
        f"- 纳入规则：{run_log.get('inclusion_rule', '-')}",
        "- 排除上传：照片/图片/皮损照片、压缩包、隐藏系统文件、重复内容。",
        f"- 本次受试者数：{len(rows)}",
        "",
        "## 汇总",
        "",
        "| 中心 | 受试者 | 系统结论 | 人工IE | 人工条目 | 对照状态 | 未完全通过/关注条目 | 提醒条目 |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for row in rows:
        issues = "；".join(f"{r['rule_id']} {VERDICT_LABELS.get(r['verdict'], r['verdict'])}" for r in row["issues"]) or "-"
        reminders = "；".join(f"{r['rule_id']} {pass_verify_label(r.get('reasoning', ''))}" for r in row["reminders"]) or "-"
        manual = row["manual"]
        lines.append(
            "| "
            + " | ".join(
                md_cell(x)
                for x in [
                    f"{row['center_code']} {row['center_name']}".strip() or "-",
                    row["subject_id"],
                    VERDICT_LABELS.get(row["overall"], row["overall"] or "未审核"),
                    manual.get("ieyn", "无记录"),
                    manual.get("manual_item", ""),
                    row["comparison"]["status"],
                    issues,
                    reminders,
                ]
            )
            + " |"
        )
    lines.extend(["", "## 逐例明细", ""])
    for row in rows:
        lines.extend([
            f"### {row['subject_id']}",
            "",
            f"- 中心：{row['center_code']} {row['center_name']}".strip(),
            f"- 系统结论：{VERDICT_LABELS.get(row['overall'], row['overall'] or '未审核')}",
            f"- 人工IE：{row['manual'].get('ieyn', '无记录')}；{row['manual'].get('manual_item', '')}；{row['manual'].get('manual_comment', '')}",
            f"- 对照：{row['comparison']['status']}",
            f"- 上传文件数：{row['uploaded_count']}；剔除/去重文件数：{row['skipped_count']}",
            "",
        ])
        if row["issues"]:
            lines.extend(["| 规则 | 结论 | 依据/需处理 |", "|---|---|---|"])
            for rule in row["issues"]:
                lines.append(
                    f"| {md_cell(rule['rule_id'] + ' ' + rule['rule_name'])} | "
                    f"{md_cell(VERDICT_LABELS.get(rule['verdict'], rule['verdict']))} | "
                    f"{md_cell(rule['reasoning'])} |"
                )
            lines.append("")
        else:
            lines.extend(["未发现系统标记的未完全通过/关注条目。", ""])
        if row["reminders"]:
            lines.extend(["| 规则 | 提醒类型 | 依据 |", "|---|---|---|"])
            for rule in row["reminders"]:
                lines.append(
                    f"| {md_cell(rule['rule_id'] + ' ' + rule['rule_name'])} | "
                    f"{md_cell(pass_verify_label(rule.get('reasoning', '')))} | "
                    f"{md_cell(rule['reasoning'])} |"
                )
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def render_html(rows: list[dict], run_log: dict) -> str:
    counts: dict[str, int] = {}
    for row in rows:
        label = VERDICT_LABELS.get(row["overall"], row["overall"] or "未审核")
        counts[label] = counts.get(label, 0) + 1
    cards = "\n".join(
        f"<div class='metric'><strong>{esc(k)}</strong><span>{v}</span></div>"
        for k, v in sorted(counts.items())
    )
    subject_sections = "\n".join(render_subject_html(row) for row in rows)
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>D001-02-II Ⅱ期筛选期批量入排审核报告</title>
<style>
:root {{ --ink:#243040; --muted:#667085; --line:#d9e0e8; --bg:#f5f7fa; --card:#fff; --accent:#155e75; --warn:#b45309; --bad:#b42318; --ok:#15803d; }}
* {{ box-sizing:border-box; }}
body {{ margin:0; font-family:-apple-system,BlinkMacSystemFont,"Segoe UI","Microsoft YaHei",sans-serif; color:var(--ink); background:var(--bg); line-height:1.55; }}
header {{ padding:28px 32px 18px; background:var(--card); border-bottom:1px solid var(--line); position:sticky; top:0; z-index:10; }}
h1 {{ margin:0 0 10px; font-size:24px; letter-spacing:0; }}
.sub {{ color:var(--muted); font-size:14px; }}
main {{ width:100%; max-width:none; margin:0; padding:22px 32px 48px; }}
.metrics {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(160px,1fr)); gap:12px; margin:18px 0; }}
.metric {{ background:var(--card); border:1px solid var(--line); border-radius:8px; padding:14px; }}
.metric strong {{ display:block; color:var(--muted); font-size:13px; margin-bottom:6px; }}
.metric span {{ font-size:24px; font-weight:700; }}
.note {{ background:#ecfeff; border:1px solid #a5f3fc; border-radius:8px; padding:14px 16px; margin:16px 0; }}
table {{ width:100%; border-collapse:collapse; background:var(--card); border:1px solid var(--line); table-layout:fixed; }}
.table-wrap {{ width:100%; overflow-x:auto; background:var(--card); border:1px solid var(--line); }}
.table-wrap table {{ border:0; }}
th,td {{ border-bottom:1px solid var(--line); padding:9px 10px; vertical-align:top; text-align:left; overflow-wrap:anywhere; }}
th {{ background:#f8fafc; font-size:13px; color:#475467; }}
.summary-table th:nth-child(1){{width:9%}} .summary-table th:nth-child(2){{width:8%}} .summary-table th:nth-child(3){{width:8%}} .summary-table th:nth-child(4){{width:7%}} .summary-table th:nth-child(5){{width:9%}} .summary-table th:nth-child(6){{width:18%}}
.subject {{ margin:18px 0; background:var(--card); border:1px solid var(--line); border-radius:8px; overflow:hidden; }}
.subject-head {{ display:flex; justify-content:space-between; gap:14px; align-items:flex-start; padding:16px 18px; border-bottom:1px solid var(--line); }}
.subject h2 {{ margin:0; font-size:18px; }}
.meta {{ color:var(--muted); font-size:13px; }}
.badge {{ display:inline-block; padding:3px 8px; border-radius:999px; font-size:12px; font-weight:700; border:1px solid var(--line); background:#f8fafc; white-space:nowrap; }}
.badge.fail,.badge.needs_evidence,.badge.insufficient {{ color:var(--bad); background:#fff1f2; border-color:#fecdd3; }}
.badge.investigator {{ color:var(--warn); background:#fffbeb; border-color:#fde68a; }}
.badge.pass_verify {{ color:#3f6212; background:#f7fee7; border-color:#d9f99d; }}
.badge.pass {{ color:var(--ok); background:#f0fdf4; border-color:#bbf7d0; }}
.detail {{ padding:16px 18px; }}
.small {{ color:var(--muted); font-size:13px; }}
@media (max-width: 720px) {{ header{{padding:20px 16px}} main{{padding:16px 12px}} .subject-head{{display:block}} table{{font-size:13px; min-width:760px; table-layout:auto}} .summary-table{{min-width:960px}} th,td{{padding:8px 7px}} .summary-table th:nth-child(n){{width:auto}} }}
</style>
</head>
<body>
<header>
  <h1>D001-02-II Ⅱ期筛选期批量入排审核与人工IE对照</h1>
  <div class="sub">项目：D001-02-II ｜ 研究阶段：Ⅱ期 ｜ 审核阶段：筛选期 ｜ 受试者：{len(rows)}例</div>
</header>
<main>
  <div class="note">本报告汇总 D001-02-II Ⅱ期筛选期受试者入排审核结果，并与人工IE表进行逐例对照。</div>
  <div class="metrics"><div class="metric"><strong>纳入受试者</strong><span>{len(rows)}</span></div>{cards}</div>
  <h2>总体对照</h2>
  {render_summary_table(rows)}
  <h2>逐例审核问题</h2>
  {subject_sections}
</main>
</body>
</html>"""


def render_summary_table(rows: list[dict]) -> str:
    body = []
    for row in rows:
        issues = "<br>".join(
            f"{esc(r['rule_id'])} {esc(VERDICT_LABELS.get(r['verdict'], r['verdict']))}"
            for r in row["issues"]
        ) or "-"
        reminders = "<br>".join(
            f"{esc(r['rule_id'])} {esc(pass_verify_label(r.get('reasoning', '')))}"
            for r in row["reminders"]
        ) or "-"
        manual = row["manual"]
        body.append(
            "<tr>"
            f"<td>{esc(row['center_code'])}<br>{esc(row['center_name'])}</td>"
            f"<td>{esc(row['subject_id'])}</td>"
            f"<td>{badge(row['overall'])}</td>"
            f"<td>{esc(manual.get('ieyn', '无记录'))}</td>"
            f"<td>{esc(manual.get('manual_item', ''))}<br><span class='small'>{esc(manual.get('manual_comment', ''))}</span></td>"
            f"<td>{esc(row['comparison']['status'])}</td>"
            f"<td>{issues}</td>"
            f"<td>{reminders}</td>"
            "</tr>"
        )
    table = "<table class='summary-table'><thead><tr><th>中心</th><th>受试者</th><th>系统</th><th>人工IE</th><th>人工条目</th><th>对照</th><th>系统关注条目</th><th>提醒条目</th></tr></thead><tbody>" + "\n".join(body) + "</tbody></table>"
    return f"<div class='table-wrap'>{table}</div>"


def render_subject_html(row: dict) -> str:
    if row["issues"]:
        issue_rows = "\n".join(
            "<tr>"
            f"<td>{esc(rule['rule_id'])}<br><span class='small'>{esc(rule['rule_name'])}</span></td>"
            f"<td>{badge(rule['verdict'])}</td>"
            f"<td>{esc(rule['reasoning'])}</td>"
            "</tr>"
            for rule in row["issues"]
        )
        issue_table = "<div class='table-wrap'><table><thead><tr><th>规则</th><th>结论</th><th>依据/需处理</th></tr></thead><tbody>" + issue_rows + "</tbody></table></div>"
    else:
        issue_table = "<p>未见未完全通过或需关注条目。</p>"
    if row["reminders"]:
        reminder_rows = "\n".join(
            "<tr>"
            f"<td>{esc(rule['rule_id'])}<br><span class='small'>{esc(rule['rule_name'])}</span></td>"
            f"<td>{esc(pass_verify_label(rule.get('reasoning', '')))}</td>"
            f"<td>{esc(rule['reasoning'])}</td>"
            "</tr>"
            for rule in row["reminders"]
        )
        reminder_table = "<h3>提醒条目</h3><div class='table-wrap'><table><thead><tr><th>规则</th><th>提醒类型</th><th>依据</th></tr></thead><tbody>" + reminder_rows + "</tbody></table></div>"
    else:
        reminder_table = ""
    manual = row["manual"]
    return f"""<section class="subject">
  <div class="subject-head">
    <div>
      <h2>{esc(row['subject_id'])}</h2>
      <div class="meta">{esc(row['center_code'])}｜{esc(row['center_name'])} ｜ Ⅱ期筛选期审核</div>
    </div>
    <div>{badge(row['overall'])}</div>
  </div>
  <div class="detail">
    <p><strong>系统摘要：</strong>{esc(row['summary'])}</p>
    <p><strong>人工IE：</strong>{esc(manual.get('ieyn','无记录'))}；{esc(manual.get('manual_item',''))}；{esc(manual.get('manual_comment',''))}</p>
    <p><strong>对照状态：</strong>{esc(row['comparison']['status'])}</p>
    {issue_table}
    {reminder_table}
  </div>
</section>"""


def badge(verdict: str) -> str:
    label = VERDICT_LABELS.get(verdict, verdict or "未审核")
    cls = (verdict or "").replace("-", "_")
    return f"<span class='badge {esc(cls)}'>{esc(label)}</span>"


def esc(text: object) -> str:
    return html.escape(clean_text(text), quote=True)


def md_cell(text: object) -> str:
    return clean_text(text).replace("|", "\\|").replace("\n", "<br>")


def run(args: argparse.Namespace) -> int:
    run_id = compact_ts()
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_LOG_DIR.mkdir(parents=True, exist_ok=True)
    staging_root = STAGING_ROOT / run_id
    run_log_path = RUN_LOG_DIR / f"d001_phase2_screening_batch_{run_id}.json"
    manual, center_names = load_manual_ie()
    folder_keyword = None if args.all_folders else args.folder_keyword
    subject_folders = discover_subject_folders(SOURCE_ROOT, folder_keyword=folder_keyword, study_stage=STUDY_STAGE)
    if args.subjects:
        wanted = {sid.upper() for sid in args.subjects}
        subject_folders = [item for item in subject_folders if item.subject_id in wanted]
    if args.only_new:
        existing = existing_project_subject_ids()
        subject_folders = [item for item in subject_folders if item.subject_id not in existing]
    if args.only_unreviewed:
        completed = completed_project_subject_ids()
        subject_folders = [item for item in subject_folders if item.subject_id not in completed]
    if args.limit and args.limit > 0:
        subject_folders = subject_folders[: args.limit]

    if folder_keyword:
        inclusion_rule = f"源目录下文件夹名包含“{folder_keyword}”且能识别SA编号者纳入；按项目固定为Ⅱ期。"
    else:
        inclusion_rule = "源目录下能识别SA编号的受试者根目录全部纳入；同一受试者的多个非嵌套资料根目录合并；按项目固定为Ⅱ期。"
    if args.only_new:
        inclusion_rule += " 已跳过项目中既有受试者。"
    if args.only_unreviewed:
        inclusion_rule += " 已跳过项目中已有完整审核报告且状态为已审核的受试者。"

    run_log = {
        "run_id": run_id,
        "started_at": now(),
        "project_code": PROJECT_CODE,
        "study_stage": STUDY_STAGE,
        "review_phase": REVIEW_PHASE,
        "source_root": str(SOURCE_ROOT),
        "subject_count": len(subject_folders),
        "folder_keyword": folder_keyword or "",
        "only_new": bool(args.only_new),
        "only_unreviewed": bool(args.only_unreviewed),
        "limit": args.limit or 0,
        "subject_workers": max(1, args.subject_workers),
        "manual_ie_usage": "人工IE仅用于审核后对照，不进入DeepSeek提示词。",
        "inclusion_rule": inclusion_rule,
        "subjects": [],
        "errors": [],
    }

    session = requests.Session()
    login = authenticate(session, args.base_url, args.username, args.password)
    run_log["login"] = {"username": login.get("username"), "role": login.get("role")}
    run_log["health"] = request_json(session, "GET", f"{args.base_url}/api/health")
    run_log["project"] = request_json(session, "GET", f"{args.base_url}/api/projects/{PROJECT_CODE}")
    run_log["workflow"] = request_json(session, "GET", f"{args.base_url}/api/projects/{PROJECT_CODE}/phases")

    print(f"Discovered {len(subject_folders)} subjects for {PROJECT_CODE}/{STUDY_STAGE}: {[s.subject_id for s in subject_folders]}", flush=True)
    worker_count = max(1, args.subject_workers)
    subject_logs: list[dict | None] = [None] * len(subject_folders)
    write_lock = threading.Lock()

    def persist_progress() -> None:
        with write_lock:
            run_log["subjects"] = [log for log in subject_logs if log is not None]
            run_log["errors"] = [
                {"subject_id": log["subject_id"], "error": log["error"]}
                for log in run_log["subjects"]
                if log.get("error")
            ]
            run_log["updated_at"] = now()
            run_log_path.write_text(json.dumps(run_log, ensure_ascii=False, indent=2), encoding="utf-8")

    if worker_count == 1:
        for index, item in enumerate(subject_folders, start=1):
            subject_logs[index - 1] = process_subject_item(
                args, item, index, len(subject_folders), manual, center_names, staging_root
            )
            persist_progress()
            if subject_logs[index - 1].get("error") and not args.continue_on_error:
                break
    else:
        print(f"Subject workers: {worker_count} (OCR/LLM overlap enabled)", flush=True)
        with concurrent.futures.ThreadPoolExecutor(max_workers=worker_count) as executor:
            future_map = {
                executor.submit(
                    process_subject_item,
                    args,
                    item,
                    index,
                    len(subject_folders),
                    manual,
                    center_names,
                    staging_root,
                ): index
                for index, item in enumerate(subject_folders, start=1)
            }
            for future in concurrent.futures.as_completed(future_map):
                index = future_map[future]
                try:
                    subject_logs[index - 1] = future.result()
                except Exception as exc:
                    item = subject_folders[index - 1]
                    subject_logs[index - 1] = {"subject_id": item.subject_id, "error": str(exc)}
                    print(f"[{item.subject_id}] ERROR: {exc}", file=sys.stderr, flush=True)
                persist_progress()
                if subject_logs[index - 1].get("error") and not args.continue_on_error:
                    executor.shutdown(cancel_futures=True)
                    break

    markdown_path, html_path, rows = build_reports(run_log, manual)
    run_log["finished_at"] = now()
    run_log["markdown_report"] = str(markdown_path)
    run_log["html_report"] = str(html_path)
    run_log["comparison_rows"] = rows
    run_log_path.write_text(json.dumps(run_log, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nRun log: {run_log_path}", flush=True)
    print(f"Markdown report: {markdown_path}", flush=True)
    print(f"HTML report: {html_path}", flush=True)
    return 1 if run_log["errors"] else 0


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default=DEFAULT_BASE_URL)
    parser.add_argument("--username", default="smkzw")
    parser.add_argument("--password", default="")
    parser.add_argument("--subjects", nargs="*")
    parser.add_argument("--folder-keyword", default="筛败")
    parser.add_argument("--all-folders", action="store_true")
    parser.add_argument("--only-new", action="store_true")
    parser.add_argument("--only-unreviewed", action="store_true")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--subject-workers", type=int, default=1)
    parser.add_argument("--no-recreate", dest="recreate", action="store_false")
    parser.add_argument("--continue-on-error", action="store_true", default=True)
    args = parser.parse_args()
    return run(args)


if __name__ == "__main__":
    raise SystemExit(main())
